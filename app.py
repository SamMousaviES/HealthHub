import copy
import functools
import hashlib
import hmac
import io
import json
import os
import base64
import mimetypes
import re
import secrets
import shlex
import shutil
import sqlite3
import subprocess
import threading
import time
import urllib.error
import urllib.request
import uuid
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

import paramiko
from flask import Flask, abort, flash, jsonify, redirect, render_template, request, send_file, session, url_for
from waitress import serve

from src.dashboard.auth import (
    authenticate_dashboard_user,
    compose_display_name,
    dashboard_accounts,
    dashboard_password,
    dashboard_sections,
    dashboard_username,
    display_name_for_username,
    ensure_dashboard_seed_accounts,
    ensure_csrf_token,
    host_allows_registration,
    login_required,
    normalize_username,
    register_dashboard_user,
    request_wants_json,
    resolve_dashboard_tab,
    verify_csrf,
    viewer_app_title,
    viewer_username,
)
from src.dashboard.coach import (
    build_assistant_prompt as build_assistant_prompt_from_context,
    empty_assistant_response,
    normalize_assistant_action,
    parse_assistant_response_text,
)
from src.dashboard.config import (
    AGENT_DISPLAY_NAMES,
    AGENT_OUTPUT_MAX_CHARS,
    AGENT_REMOTE_TIMEOUT_SECONDS,
    APP_ROOT,
    APP_TITLE,
    ASSISTANT_ALLOWED_ACTIONS,
    CLIENT_POLL_SECONDS,
    COACH_DIET_MEAL_FOOD_OPTIONS,
    COACH_SCOPE_OPTIONS,
    COACH_ATTACHMENTS_DIR,
    COACH_ATTACHMENT_MAX_FILES_PER_MESSAGE,
    COACH_ATTACHMENT_MAX_FILE_BYTES,
    COACH_ATTACHMENT_MAX_TOTAL_BYTES_PER_USER,
    COACH_ATTACHMENT_RETENTION_SECONDS,
    DEFAULT_DIET_MEALS_PER_DAY,
    DEFAULT_GYM_DAYS_PER_WEEK,
    DEFAULT_GYM_PREFERRED_MINUTES,
    DIET_MEAL_COUNT_OPTIONS,
    DIET_MEAL_CUSTOMIZATION_MODES,
    DIET_PLANNER_VERSION,
    DIET_PREFERENCE_OPTIONS,
    EXERCISE_DB_URL,
    EXERCISE_IMAGE_BASE_URL,
    EXERCISE_MEDIA_CACHE_FILE,
    EXERCISE_MEDIA_CACHE_SECONDS,
    GRAPH_GAP_SECONDS,
    GRAPH_HEIGHT,
    GRAPH_PADDING,
    GRAPH_WIDTH,
    GYM_KNOWLEDGE_DB_FILE,
    GYM_MISC_DIR,
    GYM_PROFILE_GOALS,
    GYM_USER_DB_FILE,
    HEALTH_GENDER_OPTIONS,
    HISTORY_FILE,
    HISTORY_METRICS,
    HISTORY_SECONDS,
    OPENAI_API_BASE_URL,
    OPENAI_API_KEY,
    OPENAI_VISION_DETAIL,
    OPENAI_VISION_MAX_IMAGES_PER_MESSAGE,
    OPENAI_VISION_MODEL,
    OPENAI_VISION_TIMEOUT_SECONDS,
    REFRESH_SECONDS,
    REMOTE118,
    REMOTE_SNAPSHOT_SCRIPT,
    SAMPLE_SECONDS,
    SUPERAGENT_NAME,
)
from src.dashboard.diet import build_diet_agent_prompt
from src.dashboard.gym import build_gym_agent_prompt
from src.dashboard.remote import (
    collect_remote_snapshot,
    compose_shell,
    container_shell,
    remote_codex_exec,
    run_local,
    run_remote,
)

STATE_CACHE = {"updated_at": None, "snapshots": {}}
HISTORY_CACHE = {}
CACHE_LOCK = threading.Lock()
SAMPLER_STARTED = False
GYM_DB_READY = False
GYM_DB_LOCK = threading.Lock()
EXERCISE_MEDIA_LOCK = threading.Lock()
EXERCISE_MEDIA_CACHE = None
DIET_SLOT_SEQUENCES = {
    3: [
        {"key": "breakfast", "label": "Breakfast", "share": 0.28},
        {"key": "lunch", "label": "Lunch", "share": 0.38},
        {"key": "dinner", "label": "Dinner", "share": 0.34},
    ],
    5: [
        {"key": "breakfast", "label": "Breakfast", "share": 0.22},
        {"key": "snack_am", "label": "Snack 1", "share": 0.12},
        {"key": "lunch", "label": "Lunch", "share": 0.27},
        {"key": "snack_pm", "label": "Snack 2", "share": 0.12},
        {"key": "dinner", "label": "Dinner", "share": 0.27},
    ],
}
DIET_MEAL_TIME_WINDOWS = {
    3: {
        "breakfast": "09:00 - 11:00",
        "lunch": "14:00 - 16:00",
        "dinner": "19:00 - 21:00",
    },
    5: {
        "breakfast": "08:30 - 09:30",
        "snack_am": "11:00 - 12:00",
        "lunch": "13:30 - 14:30",
        "snack_pm": "16:30 - 17:30",
        "dinner": "19:30 - 20:30",
    },
    "default": {
        "breakfast": "09:00 - 11:00",
        "lunch": "14:00 - 16:00",
        "dinner": "19:00 - 21:00",
        "snack": "16:00 - 17:00",
    },
}
DIET_FALLBACK_MESSAGES = [
    "Small repeatable meals beat perfect plans that never get cooked.",
    "Your target moves with consistent meals, not with one big day of motivation.",
    "Keep today's food simple enough that tomorrow still feels easy.",
    "Protein first, vegetables next, and let consistency do the hard work.",
    "When prep is easy, discipline costs less.",
    "A calm meal plan is easier to follow than a clever one.",
    "Treat each meal like a vote for the body you want to keep.",
    "The best diet day is the one you can repeat all week.",
    "Simple ingredients and clear portions remove friction from progress.",
    "Eat for the target in front of you, not for the mood of the hour.",
    "You do not need fancy meals. You need meals you will actually make.",
    "A steady routine wins even when today's energy is average.",
]
DIET_FALLBACK_MEAL_TEMPLATES = [
    {
        "template_key": "protein_oats_bowl",
        "slot_key": "breakfast",
        "title": "Protein oats bowl",
        "goal_bias": "balanced",
        "summary": "High-protein oats with yogurt and fruit.",
        "prep_text": "Cook the oats, then top with yogurt, banana, berries, and chia seeds.",
        "base_calories": 670,
        "protein_g": 41,
        "carbs_g": 82,
        "fat_g": 16,
        "fiber_g": 14,
        "ingredients": [
            {"qty": 80, "unit": "g", "item": "rolled oats", "step": 5},
            {"qty": 250, "unit": "g", "item": "Greek yogurt", "step": 25},
            {"qty": 1, "unit": "", "item": "banana", "step": 0.5},
            {"qty": 100, "unit": "g", "item": "berries", "step": 10},
            {"qty": 15, "unit": "g", "item": "chia seeds", "step": 5},
        ],
        "sort_order": 1,
    },
    {
        "template_key": "eggs_toast_plate",
        "slot_key": "breakfast",
        "title": "Eggs, toast, and fruit",
        "goal_bias": "gain",
        "summary": "Fast savory breakfast with a little extra energy.",
        "prep_text": "Cook the eggs, toast the bread, and eat with cottage cheese and fruit.",
        "base_calories": 620,
        "protein_g": 39,
        "carbs_g": 53,
        "fat_g": 27,
        "fiber_g": 9,
        "ingredients": [
            {"qty": 3, "unit": "", "item": "eggs", "step": 1},
            {"qty": 2, "unit": "slices", "item": "whole-grain toast", "step": 1},
            {"qty": 150, "unit": "g", "item": "cottage cheese", "step": 25},
            {"qty": 1, "unit": "", "item": "apple", "step": 0.5},
        ],
        "sort_order": 2,
    },
    {
        "template_key": "skyr_muesli_bowl",
        "slot_key": "breakfast",
        "title": "Skyr muesli bowl",
        "goal_bias": "lean",
        "summary": "Cold breakfast with strong protein and easy prep.",
        "prep_text": "Add everything to one bowl and mix before eating.",
        "base_calories": 560,
        "protein_g": 36,
        "carbs_g": 58,
        "fat_g": 18,
        "fiber_g": 11,
        "ingredients": [
            {"qty": 300, "unit": "g", "item": "skyr", "step": 25},
            {"qty": 60, "unit": "g", "item": "muesli", "step": 5},
            {"qty": 1, "unit": "", "item": "banana", "step": 0.5},
            {"qty": 20, "unit": "g", "item": "walnuts", "step": 5},
            {"qty": 10, "unit": "g", "item": "honey", "step": 5},
        ],
        "sort_order": 3,
    },
    {
        "template_key": "chicken_rice_bowl",
        "slot_key": "lunch",
        "title": "Chicken rice bowl",
        "goal_bias": "lean",
        "summary": "Simple batch-cook lunch with high protein.",
        "prep_text": "Cook the chicken and rice, then serve with vegetables and olive oil.",
        "base_calories": 760,
        "protein_g": 63,
        "carbs_g": 74,
        "fat_g": 20,
        "fiber_g": 8,
        "ingredients": [
            {"qty": 220, "unit": "g", "item": "chicken breast", "step": 10},
            {"qty": 250, "unit": "g", "item": "cooked rice", "step": 10},
            {"qty": 200, "unit": "g", "item": "mixed vegetables", "step": 10},
            {"qty": 12, "unit": "g", "item": "olive oil", "step": 2},
        ],
        "sort_order": 4,
    },
    {
        "template_key": "tuna_pasta_bowl",
        "slot_key": "lunch",
        "title": "Tuna pasta bowl",
        "goal_bias": "balanced",
        "summary": "Pasta meal with pantry ingredients and good protein.",
        "prep_text": "Warm the pasta with passata and spinach, then fold in tuna, oil, and parmesan.",
        "base_calories": 730,
        "protein_g": 52,
        "carbs_g": 76,
        "fat_g": 24,
        "fiber_g": 10,
        "ingredients": [
            {"qty": 1, "unit": "can", "item": "tuna in water, drained", "step": 0.5},
            {"qty": 250, "unit": "g", "item": "cooked pasta", "step": 10},
            {"qty": 120, "unit": "g", "item": "passata", "step": 10},
            {"qty": 80, "unit": "g", "item": "spinach", "step": 10},
            {"qty": 15, "unit": "g", "item": "olive oil", "step": 2},
            {"qty": 25, "unit": "g", "item": "parmesan", "step": 5},
        ],
        "sort_order": 5,
    },
    {
        "template_key": "beef_potato_skillet",
        "slot_key": "lunch",
        "title": "Beef and potato skillet",
        "goal_bias": "gain",
        "summary": "One-pan lunch with more energy for heavier days.",
        "prep_text": "Brown the beef, cook the potatoes until tender, and finish with vegetables and yogurt sauce.",
        "base_calories": 790,
        "protein_g": 49,
        "carbs_g": 63,
        "fat_g": 31,
        "fiber_g": 8,
        "ingredients": [
            {"qty": 180, "unit": "g", "item": "lean beef mince", "step": 10},
            {"qty": 350, "unit": "g", "item": "potatoes", "step": 10},
            {"qty": 150, "unit": "g", "item": "mixed vegetables", "step": 10},
            {"qty": 100, "unit": "g", "item": "Greek yogurt", "step": 10},
            {"qty": 15, "unit": "g", "item": "olive oil", "step": 2},
        ],
        "sort_order": 6,
    },
    {
        "template_key": "salmon_potato_tray",
        "slot_key": "dinner",
        "title": "Salmon potato tray",
        "goal_bias": "balanced",
        "summary": "Easy tray-bake dinner with steady fats and protein.",
        "prep_text": "Roast the salmon, potatoes, and vegetables together until cooked through.",
        "base_calories": 760,
        "protein_g": 46,
        "carbs_g": 56,
        "fat_g": 33,
        "fiber_g": 7,
        "ingredients": [
            {"qty": 180, "unit": "g", "item": "salmon", "step": 10},
            {"qty": 300, "unit": "g", "item": "potatoes", "step": 10},
            {"qty": 200, "unit": "g", "item": "mixed vegetables", "step": 10},
            {"qty": 10, "unit": "g", "item": "olive oil", "step": 2},
        ],
        "sort_order": 7,
    },
    {
        "template_key": "turkey_wrap_plate",
        "slot_key": "dinner",
        "title": "Turkey wrap plate",
        "goal_bias": "lean",
        "summary": "Quick dinner with lean protein and easy carbs.",
        "prep_text": "Cook the turkey, warm the wraps, and serve with vegetables and yogurt sauce.",
        "base_calories": 720,
        "protein_g": 55,
        "carbs_g": 70,
        "fat_g": 22,
        "fiber_g": 8,
        "ingredients": [
            {"qty": 200, "unit": "g", "item": "turkey mince", "step": 10},
            {"qty": 3, "unit": "", "item": "medium wraps", "step": 1},
            {"qty": 120, "unit": "g", "item": "Greek yogurt", "step": 10},
            {"qty": 150, "unit": "g", "item": "lettuce and tomato", "step": 10},
        ],
        "sort_order": 8,
    },
    {
        "template_key": "chicken_lentil_soup",
        "slot_key": "dinner",
        "title": "Chicken lentil soup",
        "goal_bias": "lean",
        "summary": "Warm one-pot dinner that stays simple.",
        "prep_text": "Simmer the chicken, lentils, and vegetables together and serve with bread.",
        "base_calories": 690,
        "protein_g": 52,
        "carbs_g": 61,
        "fat_g": 20,
        "fiber_g": 12,
        "ingredients": [
            {"qty": 180, "unit": "g", "item": "chicken breast", "step": 10},
            {"qty": 200, "unit": "g", "item": "cooked lentils", "step": 10},
            {"qty": 250, "unit": "g", "item": "soup vegetables", "step": 10},
            {"qty": 2, "unit": "slices", "item": "whole-grain bread", "step": 1},
            {"qty": 10, "unit": "g", "item": "olive oil", "step": 2},
        ],
        "sort_order": 9,
    },
    {
        "template_key": "skyr_berries_almonds",
        "slot_key": "snack",
        "title": "Skyr, berries, and almonds",
        "goal_bias": "lean",
        "summary": "Fast snack with protein and a small crunch.",
        "prep_text": "Put everything in one bowl and eat cold.",
        "base_calories": 320,
        "protein_g": 28,
        "carbs_g": 22,
        "fat_g": 12,
        "fiber_g": 8,
        "ingredients": [
            {"qty": 250, "unit": "g", "item": "skyr", "step": 25},
            {"qty": 100, "unit": "g", "item": "berries", "step": 10},
            {"qty": 20, "unit": "g", "item": "almonds", "step": 5},
            {"qty": 10, "unit": "g", "item": "honey", "step": 5},
        ],
        "sort_order": 10,
    },
    {
        "template_key": "cottage_cheese_fruit",
        "slot_key": "snack",
        "title": "Cottage cheese fruit cup",
        "goal_bias": "balanced",
        "summary": "Easy snack that travels well.",
        "prep_text": "Serve the cottage cheese with sliced fruit and peanut butter on the side.",
        "base_calories": 330,
        "protein_g": 27,
        "carbs_g": 24,
        "fat_g": 13,
        "fiber_g": 6,
        "ingredients": [
            {"qty": 200, "unit": "g", "item": "cottage cheese", "step": 25},
            {"qty": 1, "unit": "", "item": "apple", "step": 0.5},
            {"qty": 20, "unit": "g", "item": "peanut butter", "step": 5},
        ],
        "sort_order": 11,
    },
    {
        "template_key": "whey_banana_shake",
        "slot_key": "snack",
        "title": "Whey banana shake",
        "goal_bias": "gain",
        "summary": "Very fast shake for busy afternoons.",
        "prep_text": "Blend everything until smooth and drink cold.",
        "base_calories": 360,
        "protein_g": 31,
        "carbs_g": 39,
        "fat_g": 9,
        "fiber_g": 7,
        "ingredients": [
            {"qty": 30, "unit": "g", "item": "whey protein", "step": 5},
            {"qty": 1, "unit": "", "item": "banana", "step": 0.5},
            {"qty": 35, "unit": "g", "item": "rolled oats", "step": 5},
            {"qty": 300, "unit": "ml", "item": "milk", "step": 25},
        ],
        "sort_order": 12,
    },
]
DIET_NUTRIENT_DEFINITIONS = (
    {"code": "protein_g", "label": "Protein", "unit": "g", "category": "macro", "sort_order": 1, "description": "Protein intake for recovery and muscle retention."},
    {"code": "carbs_g", "label": "Carbs", "unit": "g", "category": "macro", "sort_order": 2, "description": "Carbohydrate intake for training fuel and recovery."},
    {"code": "fat_g", "label": "Fat", "unit": "g", "category": "macro", "sort_order": 3, "description": "Dietary fat intake for hormones and satiety."},
    {"code": "fiber_g", "label": "Fiber", "unit": "g", "category": "quality", "sort_order": 4, "description": "Fiber intake for digestion, fullness, and food quality."},
)
DIET_TRACKED_NUTRIENT_CODES = tuple(item["code"] for item in DIET_NUTRIENT_DEFINITIONS)
HEALTH_MEASUREMENT_FIELDS = (
    {"key": "weight_kg", "label": "Weight", "unit": "kg"},
    {"key": "arm_cm", "label": "Arm", "unit": "cm"},
    {"key": "chest_cm", "label": "Chest", "unit": "cm"},
    {"key": "waist_cm", "label": "Waist", "unit": "cm"},
    {"key": "thigh_cm", "label": "Thigh", "unit": "cm"},
    {"key": "calf_cm", "label": "Calf", "unit": "cm"},
)
GYM_SPLIT_LABELS = {
    "push": "Push",
    "pull": "Pull",
    "legs": "Legs",
    "lower_body": "Lower Body",
    "upper_body": "Upper Body",
    "full_body": "Full Body",
    "recovery": "Recovery",
}
GYM_FALLBACK_TEMPLATES = [
    {
        "source_name": "fallback",
        "sheet_name": "Day 1 - Push",
        "split_key": "push",
        "title": "Push",
        "next_split_key": "legs",
        "exercise_name": "Dumbbell Floor Press",
        "sets_text": "4",
        "reps_text": "8-10",
        "rest_text": "2 min",
        "suggested_weight_text": "14-16 kg each",
        "equipment_type": "dumbbell",
        "sort_order": 1,
    },
    {
        "source_name": "fallback",
        "sheet_name": "Day 1 - Push",
        "split_key": "push",
        "title": "Push",
        "next_split_key": "legs",
        "exercise_name": "Seated Arnold Press",
        "sets_text": "4",
        "reps_text": "8-10",
        "rest_text": "2 min",
        "suggested_weight_text": "10-12 kg each",
        "equipment_type": "dumbbell",
        "sort_order": 2,
    },
    {
        "source_name": "fallback",
        "sheet_name": "Day 2 - Pull",
        "split_key": "pull",
        "title": "Pull",
        "next_split_key": "push",
        "exercise_name": "Close-Grip Lat Pulldown",
        "sets_text": "4",
        "reps_text": "8-10",
        "rest_text": "2 min",
        "suggested_weight_text": "40-45 kg",
        "equipment_type": "machine",
        "sort_order": 1,
    },
    {
        "source_name": "fallback",
        "sheet_name": "Day 2 - Pull",
        "split_key": "pull",
        "title": "Pull",
        "next_split_key": "push",
        "exercise_name": "Single-Arm Cable Row",
        "sets_text": "3",
        "reps_text": "10-12",
        "rest_text": "90 sec",
        "suggested_weight_text": "20-25 kg",
        "equipment_type": "cable",
        "sort_order": 2,
    },
    {
        "source_name": "fallback",
        "sheet_name": "Day 3 - Lower Body",
        "split_key": "legs",
        "title": "Lower Body",
        "next_split_key": "pull",
        "exercise_name": "Front Squat",
        "sets_text": "4",
        "reps_text": "6-8",
        "rest_text": "2-3 min",
        "suggested_weight_text": "35-45 kg",
        "equipment_type": "barbell",
        "sort_order": 1,
    },
    {
        "source_name": "fallback",
        "sheet_name": "Day 3 - Lower Body",
        "split_key": "legs",
        "title": "Lower Body",
        "next_split_key": "pull",
        "exercise_name": "Walking Lunge",
        "sets_text": "3",
        "reps_text": "10 each leg",
        "rest_text": "90 sec",
        "suggested_weight_text": "10-12 kg each",
        "equipment_type": "dumbbell",
        "sort_order": 2,
    },
]
GYM_EXERCISE_IMAGE_ALIASES = {
    "front squat": "Front Squat (Clean Grip)",
    "stiff leg deadlift barbell": "Stiff-Legged Barbell Deadlift",
    "walking lunge": "Bodyweight Walking Lunge",
    "cable pull through": "Pull Through",
    "seated calf raise toes out": "Seated Calf Raise",
    "weighted sit up": "Weighted Sit-Ups - With Bands",
    "flutter kicks": "Flutter Kicks",
    "close grip lat pulldown": "Close-Grip Front Lat Pulldown",
    "seated cable row": "Seated Cable Rows",
    "single arm cable row": "Kneeling Single-Arm High Pulley Row",
    "reverse pec deck fly": "Reverse Machine Flyes",
    "zottman curl": "Zottman Curl",
    "cable crunch": "Cable Crunch",
    "side plank": "Side Bridge",
    "dumbbell floor press": "Dumbbell Floor Press",
    "seated arnold press": "Arnold Dumbbell Press",
    "incline cable fly low to high": "Incline Cable Flye",
    "low to high cable fly": "Low Cable Crossover",
    "leaning lateral raise": "Side Lateral Raise",
    "overhead rope extension": "Cable Rope Overhead Triceps Extension",
    "rope overhead triceps extension": "Cable Rope Overhead Triceps Extension",
    "hanging knee raise": "Hanging Leg Raise",
    "ab wheel rollout": "Barbell Ab Rollout - On Knees",
    "plank with reach": "Plank",
}

app = Flask(__name__)
app.secret_key = os.getenv("DASHBOARD_SECRET_KEY", secrets.token_hex(32))

REGISTRATION_USERNAME_PATTERN = re.compile(r"^[A-Za-z0-9._-]{3,32}$")


SERVERS = {
    "106": {
        "id": "106",
        "name": "Server 106",
        "label": "Raspberry Pi / Nextcloud",
        "host": "192.168.1.106",
        "accent": "amber",
        "local": True,
        "host_actions": [
            {"id": "reboot", "label": "Reboot", "button_class": "warn"},
        ],
        "links": [
            {"label": "Nextcloud", "url": "https://nc.sam-mousavi.com"},
            {"label": "Paperless", "url": "https://paperless.sam-mousavi.com"},
            {"label": "n8n", "url": "https://n8n.sam-mousavi.com"},
            {"label": "Shell", "url": "https://ssh.sam-mousavi.com"},
        ],
    },
    "118": {
        "id": "118",
        "name": "Server 118",
        "label": "MSI / AI and Edge Services",
        "host": os.getenv("REMOTE118_HOST", "192.168.1.118"),
        "accent": "teal",
        "local": False,
        "host_actions": [
            {"id": "reboot", "label": "Reboot", "button_class": "warn"},
        ],
        "links": [
            {"label": "LLM", "url": "https://llm.sam-mousavi.com"},
            {"label": "Whisper", "url": "https://whisper.sam-mousavi.com"},
            {"label": "Jupyter", "url": "https://jupyter.sam-mousavi.com"},
        ],
    },
}

SERVICES = {
    "106": [
        {
            "id": "nextcloud",
            "name": "Nextcloud AIO",
            "icon": "cloud",
            "mode": "compose",
            "control_mode": "read_only",
            "path": "/home/sam/Docker/nextcloud",
            "containers": [
                "nextcloud-aio-mastercontainer",
                "nextcloud-aio-apache",
                "nextcloud-aio-nextcloud",
                "nextcloud-aio-database",
                "nextcloud-aio-redis",
                "nextcloud-aio-talk",
                "nextcloud-aio-whiteboard",
                "nextcloud-aio-notify-push",
                "nextcloud-aio-collabora",
            ],
            "url": "https://nc.sam-mousavi.com",
            "description": "Primary Nextcloud stack. Monitored here, but not controllable from Health Hub.",
        },
        {
            "id": "paperless",
            "name": "Paperless",
            "icon": "files",
            "mode": "compose",
            "control_mode": "compose",
            "path": "/home/sam/Docker/paperless",
            "containers": ["paperless-webserver-1", "paperless-db-1", "paperless-broker-1"],
            "url": "https://paperless.sam-mousavi.com",
            "description": "Document archive and OCR pipeline.",
        },
        {
            "id": "n8n",
            "name": "n8n",
            "icon": "workflow",
            "mode": "compose",
            "control_mode": "compose",
            "path": "/home/sam/Docker/n8n",
            "containers": ["n8n-n8n-1", "n8n-postgres-1"],
            "url": "https://n8n.sam-mousavi.com",
            "description": "Automation workflows and integrations.",
        },
        {
            "id": "ttyd",
            "name": "Shell (SSH)",
            "icon": "terminal",
            "mode": "compose",
            "control_mode": "compose",
            "path": "/home/sam/Docker/ttyd",
            "containers": ["ttyd"],
            "url": "https://ssh.sam-mousavi.com",
            "description": "Browser shell access.",
        },
        {
            "id": "amnezia",
            "name": "Amnezia WireGuard",
            "icon": "shield",
            "mode": "container",
            "control_mode": "container",
            "containers": ["amnezia-wireguard"],
            "url": None,
            "description": "Standalone VPN container on the Pi host.",
        },
    ],
    "118": [
        {
            "id": "open_webui",
            "name": "Open WebUI",
            "icon": "layout",
            "mode": "compose",
            "control_mode": "compose",
            "path": "/home/sam/Docker/open-webui",
            "containers": ["open-webui"],
            "url": "https://llm.sam-mousavi.com",
            "description": "Main LLM front-end backed by Ollama.",
        },
        {
            "id": "ollama_msi",
            "name": "Ollama",
            "icon": "bot",
            "mode": "compose",
            "control_mode": "compose",
            "path": "/home/sam/Docker/ollama",
            "containers": ["ollama"],
            "url": "https://llm.sam-mousavi.com",
            "description": "Model runtime on server 118.",
        },
        {
            "id": "whisper",
            "name": "Whisper",
            "icon": "audio",
            "mode": "container",
            "control_mode": "container",
            "containers": ["whisper"],
            "url": "https://whisper.sam-mousavi.com",
            "description": "Speech-to-text API service.",
        },
        {
            "id": "jupyter",
            "name": "Jupyter",
            "icon": "notebook",
            "mode": "compose",
            "control_mode": "compose",
            "path": "/home/sam/Docker/jupyter",
            "containers": ["jupyter"],
            "url": "https://jupyter.sam-mousavi.com",
            "description": "Notebook environment on server 118.",
        },
        {
            "id": "rustdesk",
            "name": "RustDesk",
            "icon": "monitor",
            "mode": "compose",
            "control_mode": "compose",
            "path": "/home/sam/Docker/rustdesk",
            "containers": ["hbbs", "hbbr"],
            "url": None,
            "description": "Self-hosted remote desktop relay and rendezvous.",
            "details": [
                {"label": "Server", "value": "rust.sam-mousavi.com"},
                {"label": "Relay", "value": "rust.sam-mousavi.com"},
                {"label": "Key", "value": "m3Tp3MieZLeRUK8eZzDulIAiOwqRbVwzro0lm9L7020="},
            ],
        },
        {
            "id": "coturn",
            "name": "Coturn",
            "icon": "arrows",
            "mode": "compose",
            "control_mode": "compose",
            "path": "/home/sam/Docker/coturn",
            "containers": ["coturn-server"],
            "url": None,
            "description": "TURN/STUN relay service.",
        },
        {
            "id": "talk_hpb",
            "name": "Talk HPB Stack",
            "icon": "talk",
            "mode": "compose",
            "control_mode": "compose",
            "path": "/home/sam/Docker/talk-hpb",
            "containers": ["talk-signaling", "talk-janus", "talk-nats", "nextcloud-talk-hpb"],
            "url": "https://talk.sam-mousavi.com",
            "description": "Standalone signaling and media stack on server 118.",
        },
    ],
}

_LEGACY_REMOTE118 = {
    "host": SERVERS["118"]["host"],
    "user": os.getenv("REMOTE118_USER", "sam"),
    "port": int(os.getenv("REMOTE118_PORT", "22")),
    "key_path": os.getenv("REMOTE118_KEY_PATH", "/home/sam/.ssh/dashboard_118"),
}


_LEGACY_REMOTE_SNAPSHOT_SCRIPT = r"""
import json
import os
import shutil
import subprocess
import time


def cpu_percent(delay=0.15):
    def read():
        with open('/proc/stat', 'r', encoding='utf-8') as handle:
            parts = handle.readline().split()[1:]
        values = [int(item) for item in parts]
        idle = values[3] + values[4]
        total = sum(values)
        return idle, total

    idle1, total1 = read()
    time.sleep(delay)
    idle2, total2 = read()
    idle_delta = idle2 - idle1
    total_delta = total2 - total1
    if total_delta <= 0:
        return 0.0
    return round(100.0 * (1.0 - (idle_delta / total_delta)), 1)


def meminfo():
    values = {}
    with open('/proc/meminfo', 'r', encoding='utf-8') as handle:
        for line in handle:
            key, value = line.split(':', 1)
            values[key] = int(value.strip().split()[0])
    total = values.get('MemTotal', 0)
    available = values.get('MemAvailable', 0)
    used = max(total - available, 0)
    percent = round((used / total) * 100, 1) if total else 0.0
    return {
        'total_gb': round(total / 1024 / 1024, 2),
        'used_gb': round(used / 1024 / 1024, 2),
        'percent': percent,
    }


def temperature():
    probes = [
        '/sys/class/thermal/thermal_zone0/temp',
        '/sys/class/thermal/thermal_zone1/temp',
    ]
    for probe in probes:
        if os.path.exists(probe):
            with open(probe, 'r', encoding='utf-8') as handle:
                raw = handle.read().strip()
            try:
                value = float(raw)
            except ValueError:
                continue
            return round(value / 1000.0, 1) if value > 1000 else round(value, 1)
    return None


def containers():
    result = subprocess.run(
        ['docker', 'ps', '-a', '--format', '{{json .}}'],
        capture_output=True,
        text=True,
        check=False,
    )
    items = []
    if result.returncode != 0:
        return items, result.stderr.strip()
    for line in result.stdout.splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            payload = json.loads(line)
        except json.JSONDecodeError:
            continue
        items.append(
            {
                'name': payload.get('Names', ''),
                'image': payload.get('Image', ''),
                'status_text': payload.get('Status', ''),
            }
        )
    return items, ''


disk = shutil.disk_usage('/')
container_items, container_error = containers()

payload = {
    'reachable': True,
    'hostname': os.uname().nodename,
    'cpu_percent': cpu_percent(),
    'load_avg': [round(value, 2) for value in os.getloadavg()],
    'memory': meminfo(),
    'disk': {
        'total_gb': round(disk.total / 1024 / 1024 / 1024, 1),
        'used_gb': round((disk.total - disk.free) / 1024 / 1024 / 1024, 1),
        'percent': round(((disk.total - disk.free) / disk.total) * 100, 1) if disk.total else 0.0,
    },
    'temperature_c': temperature(),
    'uptime_seconds': int(float(open('/proc/uptime', 'r', encoding='utf-8').read().split()[0])),
    'containers': container_items,
    'container_error': container_error,
}

print(json.dumps(payload))
"""


def now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


COACH_ATTACHMENT_TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".log"}
COACH_ATTACHMENT_SPREADSHEET_EXTENSIONS = {".xlsx"}
COACH_ATTACHMENT_DOCUMENT_EXTENSIONS = {".pdf"}
COACH_ATTACHMENT_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
COACH_ATTACHMENT_ALLOWED_EXTENSIONS = (
    COACH_ATTACHMENT_TEXT_EXTENSIONS
    | COACH_ATTACHMENT_SPREADSHEET_EXTENSIONS
    | COACH_ATTACHMENT_DOCUMENT_EXTENSIONS
    | COACH_ATTACHMENT_IMAGE_EXTENSIONS
)


def coach_attachment_retention_days():
    return max(1, int(round(COACH_ATTACHMENT_RETENTION_SECONDS / 86400.0)))


def coach_attachment_allowed_extensions():
    allowed = (
        COACH_ATTACHMENT_TEXT_EXTENSIONS
        | COACH_ATTACHMENT_SPREADSHEET_EXTENSIONS
        | COACH_ATTACHMENT_DOCUMENT_EXTENSIONS
    )
    if assistant_image_analysis_enabled():
        allowed |= COACH_ATTACHMENT_IMAGE_EXTENSIONS
    return allowed


def coach_attachment_accept_text():
    ordered = [".txt", ".md", ".csv", ".json", ".log", ".xlsx", ".pdf"]
    if assistant_image_analysis_enabled():
        ordered.extend([".png", ".jpg", ".jpeg", ".webp", ".gif"])
    return ",".join(ordered)


def parse_iso_utc(value: str | None):
    normalized = str(value or "").strip()
    if not normalized:
        return None
    try:
        if normalized.endswith("Z"):
            return datetime.fromisoformat(normalized[:-1] + "+00:00")
        parsed = datetime.fromisoformat(normalized)
        return parsed if parsed.tzinfo else parsed.replace(tzinfo=timezone.utc)
    except ValueError:
        return None


def display_timestamp(value: str | None) -> str:
    if not value:
        return "Not generated yet"
    normalized = str(value).strip()
    parsed = parse_iso_utc(normalized)
    if parsed is None:
        return normalized
    return parsed.strftime("%Y-%m-%d %H:%M UTC")


def truncate_text(value: str | None, limit: int):
    text = str(value or "").strip()
    if len(text) <= limit:
        return text
    return text[: max(0, limit - 1)].rstrip() + "…"


def sanitize_attachment_filename(value: str | None):
    raw = str(value or "").strip()
    if not raw:
        raw = "attachment"
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", raw).strip("._")
    return safe or "attachment"


def format_file_size(byte_size: int | float | None):
    size = max(0.0, float(byte_size or 0))
    units = ["B", "KB", "MB", "GB"]
    unit_index = 0
    while size >= 1024.0 and unit_index < len(units) - 1:
        size /= 1024.0
        unit_index += 1
    precision = 0 if unit_index == 0 else 1
    return f"{size:.{precision}f} {units[unit_index]}"


def coach_attachment_kind_label(file_ext: str | None):
    ext = str(file_ext or "").strip().lower()
    if ext in COACH_ATTACHMENT_TEXT_EXTENSIONS:
        return "Text"
    if ext in COACH_ATTACHMENT_SPREADSHEET_EXTENSIONS:
        return "Spreadsheet"
    if ext in COACH_ATTACHMENT_DOCUMENT_EXTENSIONS:
        return "PDF"
    if ext in COACH_ATTACHMENT_IMAGE_EXTENSIONS:
        return "Image"
    return "File"


def coach_attachment_is_image(file_ext: str | None, mime_type: str | None = None):
    ext = str(file_ext or "").strip().lower()
    mime = str(mime_type or "").strip().lower()
    return ext in COACH_ATTACHMENT_IMAGE_EXTENSIONS or mime.startswith("image/")


def coach_attachment_user_dir(username: str):
    current = normalize_username(username) or dashboard_username()
    return COACH_ATTACHMENTS_DIR / current


def attachment_preview_text(text: str | None, limit: int = 180):
    compact = re.sub(r"\s+", " ", str(text or "").strip())
    return truncate_text(compact, limit) if compact else ""


def extract_text_from_xlsx_bytes(data: bytes):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return ""
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        return ""
    lines = []
    try:
        for sheet in workbook.worksheets[:3]:
            lines.append(f"[Sheet] {sheet.title}")
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_index > 60:
                    lines.append("…")
                    break
                values = [str(cell).strip() for cell in row[:12] if cell not in {None, ""}]
                if values:
                    lines.append("\t".join(values))
            if len(lines) >= 220:
                break
    finally:
        workbook.close()
    return "\n".join(lines)


def extract_text_from_pdf_bytes(data: bytes):
    try:
        from pypdf import PdfReader
    except ImportError:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            return ""
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception:
        return ""
    parts = []
    for page in reader.pages[:8]:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(part.strip() for part in parts if part and part.strip())


def extract_attachment_text_bytes(file_ext: str, data: bytes):
    ext = str(file_ext or "").strip().lower()
    if ext in COACH_ATTACHMENT_TEXT_EXTENSIONS:
        for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", errors="ignore")
    if ext in COACH_ATTACHMENT_SPREADSHEET_EXTENSIONS:
        return extract_text_from_xlsx_bytes(data)
    if ext in COACH_ATTACHMENT_DOCUMENT_EXTENSIONS:
        return extract_text_from_pdf_bytes(data)
    return ""


def prepare_uploaded_assistant_attachments(files):
    active_files = [item for item in (files or []) if getattr(item, "filename", None)]
    if len(active_files) > COACH_ATTACHMENT_MAX_FILES_PER_MESSAGE:
        raise ValueError(f"Attach up to {COACH_ATTACHMENT_MAX_FILES_PER_MESSAGE} files per coach message.")
    allowed_extensions = coach_attachment_allowed_extensions()
    prepared = []
    for upload in active_files:
        original_name = sanitize_attachment_filename(upload.filename)
        file_ext = Path(original_name).suffix.lower()
        if file_ext not in allowed_extensions:
            if assistant_image_analysis_enabled():
                raise ValueError(
                    "Unsupported attachment type. Use text, CSV, JSON, XLSX, PDF, or common image files."
                )
            raise ValueError(
                "Unsupported attachment type. Use text, CSV, JSON, XLSX, or PDF files."
            )
        data = upload.read()
        byte_size = len(data)
        if byte_size <= 0:
            continue
        if byte_size > COACH_ATTACHMENT_MAX_FILE_BYTES:
            raise ValueError(
                f"{original_name} is too large. Keep each file under {format_file_size(COACH_ATTACHMENT_MAX_FILE_BYTES)}."
            )
        mime_type = (upload.mimetype or mimetypes.guess_type(original_name)[0] or "application/octet-stream").strip()
        analysis_text = truncate_text(extract_attachment_text_bytes(file_ext, data), 6000)
        preview_text = attachment_preview_text(analysis_text)
        if not preview_text and file_ext in COACH_ATTACHMENT_IMAGE_EXTENSIONS:
            preview_text = "Image attached."
        elif not preview_text and file_ext in (COACH_ATTACHMENT_DOCUMENT_EXTENSIONS | COACH_ATTACHMENT_SPREADSHEET_EXTENSIONS):
            preview_text = "Text preview unavailable in this environment."
        prepared.append(
            {
                "original_name": original_name,
                "file_ext": file_ext,
                "mime_type": mime_type,
                "byte_size": byte_size,
                "sha256": hashlib.sha256(data).hexdigest(),
                "storage_key": f"{uuid.uuid4().hex}{file_ext}",
                "data": data,
                "is_image": coach_attachment_is_image(file_ext, mime_type),
                "analysis_text": analysis_text,
                "analysis_json": "",
                "preview_text": preview_text,
                "kind_label": coach_attachment_kind_label(file_ext),
            }
        )
    return prepared


def normalize_split_key(value: str | None) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", (value or "").strip().lower()).strip("_")
    if slug in {"lower_body", "lower"}:
        return "legs"
    if slug in {"upper", "upper_body"}:
        return "upper_body"
    return slug or "full_body"


def split_label(split_key: str) -> str:
    return GYM_SPLIT_LABELS.get(split_key, split_key.replace("_", " ").title())


def normalize_exercise_name(value: str | None) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (value or "").lower()).strip()


def infer_equipment_type(exercise_name: str) -> str:
    lowered = exercise_name.lower()
    if "barbell" in lowered or "squat" in lowered or "deadlift" in lowered:
        return "barbell"
    if "dumbbell" in lowered or "goblet" in lowered or "arnold" in lowered or "curl" in lowered or "lunge" in lowered:
        return "dumbbell"
    if "cable" in lowered or "rope" in lowered or "pulldown" in lowered:
        return "cable"
    if "machine" in lowered or "pec deck" in lowered or "leg press" in lowered:
        return "machine"
    if "plank" in lowered or "bodyweight" in lowered or "dead bug" in lowered or "flutter" in lowered:
        return "bodyweight"
    return "mixed"


def parse_weight_kg_from_text(value: str | None):
    text = (value or "").strip().replace(",", ".")
    if not text:
        return None
    lowered = text.lower()
    if "bodyweight" in lowered or "body weight" in lowered:
        return 0.0
    matches = re.findall(r"\d+(?:\.\d+)?", text)
    if not matches:
        return None
    numbers = [float(item) for item in matches[:2]]
    if not numbers:
        return None
    if len(numbers) == 1:
        return round(numbers[0], 1)
    return round(sum(numbers) / len(numbers), 1)


def format_weight_kg(value):
    if value is None:
        return ""
    if float(value).is_integer():
        return str(int(value))
    return f"{float(value):.1f}"


def latest_program_workbook():
    candidates = [
        path
        for path in GYM_MISC_DIR.glob("*.xlsx")
        if path.name.lower() != "weightvssizes.xlsx"
    ]
    if not candidates:
        return None

    def sort_key(path: Path):
        match = re.search(r"(20\d{6})", path.stem)
        dated_value = int(match.group(1)) if match else 0
        return (dated_value, path.stat().st_mtime, path.name.lower())

    return max(candidates, key=sort_key)


def excel_serial_to_iso_date(value: str) -> str | None:
    try:
        serial = float(value)
    except (TypeError, ValueError):
        return None
    base = datetime(1899, 12, 30)
    return (base + timedelta(days=serial)).date().isoformat()


def load_workbook_rows(path: Path):
    ns_main = {"a": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    ns_rel = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
    with zipfile.ZipFile(path) as archive:
        shared = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for string_item in shared_root.findall("a:si", ns_main):
                shared.append("".join(node.text or "" for node in string_item.findall(".//a:t", ns_main)))

        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        rel_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_map = {rel.attrib["Id"]: rel.attrib["Target"] for rel in rel_root.findall("r:Relationship", ns_rel)}
        sheet_payloads = []
        for sheet in workbook.find("a:sheets", ns_main):
            name = sheet.attrib.get("name", "Sheet")
            rel_id = sheet.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "")
            target = rel_map.get(rel_id, "")
            if not target:
                continue
            sheet_path = "xl/" + target.replace("\\", "/").lstrip("/")
            sheet_root = ET.fromstring(archive.read(sheet_path))
            rows = []
            for row in sheet_root.findall(".//a:sheetData/a:row", ns_main):
                values = {}
                for cell in row.findall("a:c", ns_main):
                    ref = cell.attrib.get("r", "")
                    match = re.match(r"[A-Z]+", ref)
                    if not match:
                        continue
                    column = match.group(0)
                    cell_type = cell.attrib.get("t")
                    value_node = cell.find("a:v", ns_main)
                    text = ""
                    if cell_type == "inlineStr":
                        text = "".join(node.text or "" for node in cell.findall(".//a:t", ns_main))
                    elif value_node is not None:
                        raw = value_node.text or ""
                        if cell_type == "s":
                            try:
                                text = shared[int(raw)]
                            except (ValueError, IndexError):
                                text = raw
                        else:
                            text = raw
                    values[column] = text.strip()
                if values:
                    rows.append(values)
            sheet_payloads.append({"name": name, "rows": rows})
        return sheet_payloads


def program_rows_from_examples():
    workbook = latest_program_workbook()
    if not workbook or not workbook.exists():
        return []
    payloads = []
    for sheet in load_workbook_rows(workbook):
        rows = sheet["rows"]
        if len(rows) < 2:
            continue
        title = sheet["name"]
        split_key = normalize_split_key(title.split("-", 1)[-1] if "-" in title else title)
        next_split_key = {"push": "legs", "legs": "pull", "pull": "push"}.get(split_key, "pull")
        for index, row in enumerate(rows[1:], start=1):
            exercise_name = row.get("A", "").strip()
            if not exercise_name:
                continue
            payloads.append(
                {
                    "source_name": workbook.name,
                    "sheet_name": title,
                    "split_key": split_key,
                    "title": split_label(split_key),
                    "next_split_key": next_split_key,
                    "exercise_name": exercise_name,
                    "sets_text": row.get("B", ""),
                    "reps_text": row.get("C", ""),
                    "rest_text": row.get("D", ""),
                    "suggested_weight_text": row.get("E", ""),
                    "equipment_type": infer_equipment_type(exercise_name),
                    "sort_order": index,
                }
            )
    return payloads


def measurement_rows_from_examples():
    workbook = GYM_MISC_DIR / "WeightvsSizes.xlsx"
    if not workbook.exists():
        return []
    rows = load_workbook_rows(workbook)
    if not rows:
        return []
    measurements = []
    for row in rows[0]["rows"][1:]:
        measured_on = excel_serial_to_iso_date(row.get("A", ""))
        if not measured_on:
            continue
        measurements.append(
            {
                "measured_on": measured_on,
                "weight_kg": row.get("B", ""),
                "arm_cm": row.get("C", ""),
                "chest_cm": row.get("D", ""),
                "waist_cm": row.get("E", ""),
                "thigh_cm": row.get("F", ""),
                "calf_cm": row.get("G", ""),
                "note": row.get("H", ""),
                "source": workbook.name,
            }
        )
    return measurements


def load_exercise_media_cache_from_disk():
    if not EXERCISE_MEDIA_CACHE_FILE.exists():
        return None
    try:
        return json.loads(EXERCISE_MEDIA_CACHE_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None


def save_exercise_media_cache_to_disk(payload: dict):
    EXERCISE_MEDIA_CACHE_FILE.parent.mkdir(parents=True, exist_ok=True)
    EXERCISE_MEDIA_CACHE_FILE.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def build_exercise_media_index():
    request = urllib.request.Request(
        EXERCISE_DB_URL,
        headers={
            "User-Agent": "Mozilla/5.0",
            "Accept-Language": "en-US,en;q=0.8",
        },
    )
    with urllib.request.urlopen(request, timeout=30) as response:
        charset = response.headers.get_content_charset() or "utf-8"
        data = json.loads(response.read().decode(charset, errors="replace"))
    index = {}
    for item in data:
        name = item.get("name", "").strip()
        if not name:
            continue
        normalized = normalize_exercise_name(name)
        image_paths = item.get("images") or []
        image_urls = [EXERCISE_IMAGE_BASE_URL + path for path in image_paths[:2]]
        index[normalized] = {
            "name": name,
            "image_urls": image_urls,
            "instructions": item.get("instructions") or [],
            "source_url": item.get("url") or "",
        }
    payload = {
        "updated_at": time.time(),
        "items": index,
    }
    save_exercise_media_cache_to_disk(payload)
    return payload


def current_exercise_media_index():
    global EXERCISE_MEDIA_CACHE
    with EXERCISE_MEDIA_LOCK:
        if EXERCISE_MEDIA_CACHE is None:
            EXERCISE_MEDIA_CACHE = load_exercise_media_cache_from_disk()
        cached = EXERCISE_MEDIA_CACHE
        if cached and time.time() - float(cached.get("updated_at", 0.0)) < EXERCISE_MEDIA_CACHE_SECONDS:
            return cached.get("items", {})
        try:
            EXERCISE_MEDIA_CACHE = build_exercise_media_index()
        except Exception:
            if cached:
                return cached.get("items", {})
            return {}
        return EXERCISE_MEDIA_CACHE.get("items", {})


def resolve_exercise_media(exercise_name: str):
    index = current_exercise_media_index()
    normalized = normalize_exercise_name(exercise_name)
    normalized_without_prefix = normalize_exercise_name(re.sub(r"^abs finisher\s+[—-]\s+", "", exercise_name, flags=re.I))
    candidates = [normalized]
    alias = GYM_EXERCISE_IMAGE_ALIASES.get(normalized)
    if alias:
        candidates.append(normalize_exercise_name(alias))
    if normalized_without_prefix and normalized_without_prefix != normalized:
        candidates.append(normalized_without_prefix)
        stripped_alias = GYM_EXERCISE_IMAGE_ALIASES.get(normalized_without_prefix)
        if stripped_alias:
            candidates.append(normalize_exercise_name(stripped_alias))
    candidates.append(normalize_exercise_name(re.sub(r"\([^)]*\)", "", exercise_name)))
    for key in candidates:
        if not key:
            continue
        payload = index.get(key)
        if payload and (payload.get("image_urls") or payload.get("instructions")):
            return payload
    words = [word for word in normalized.split() if len(word) > 2]
    for key, payload in index.items():
        if not (payload.get("image_urls") or payload.get("instructions")):
            continue
        if all(word in key for word in words[:2]) and words:
            return payload
    return {"name": exercise_name, "image_urls": [], "instructions": [], "source_url": ""}


def gym_user_connection():
    GYM_USER_DB_FILE.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(GYM_USER_DB_FILE)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def gym_knowledge_connection():
    GYM_KNOWLEDGE_DB_FILE.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(GYM_KNOWLEDGE_DB_FILE)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    return connection


def ensure_sqlite_column(connection, table_name: str, column_name: str, definition: str):
    columns = {row["name"] for row in connection.execute(f"PRAGMA table_info({table_name})").fetchall()}
    if column_name not in columns:
        try:
            connection.execute(f"ALTER TABLE {table_name} ADD COLUMN {column_name} {definition}")
        except sqlite3.OperationalError as exc:
            if "duplicate column name" not in str(exc).lower():
                raise


def ensure_gym_databases():
    global GYM_DB_READY
    if GYM_DB_READY:
        return
    with GYM_DB_LOCK:
        if GYM_DB_READY:
            return
        with gym_user_connection() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    username TEXT NOT NULL UNIQUE,
                    display_name TEXT,
                    first_name TEXT,
                    last_name TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS auth_accounts (
                    user_id INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE,
                    password_hash TEXT NOT NULL,
                    role TEXT NOT NULL DEFAULT 'member',
                    is_active INTEGER NOT NULL DEFAULT 1,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS gym_profiles (
                    user_id INTEGER PRIMARY KEY REFERENCES users(id) ON DELETE CASCADE,
                    height_cm REAL,
                    age_years REAL,
                    current_weight_kg REAL,
                    desired_weight_kg REAL,
                    gender TEXT NOT NULL DEFAULT 'unspecified',
                    goal TEXT NOT NULL DEFAULT 'recomp',
                    preferred_session_minutes INTEGER NOT NULL DEFAULT 60,
                    training_days_per_week INTEGER NOT NULL DEFAULT 4,
                    preferred_meals_per_day INTEGER NOT NULL DEFAULT 3,
                    notes TEXT,
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS body_measurements (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    measured_on TEXT NOT NULL,
                    weight_kg REAL,
                    arm_cm REAL,
                    chest_cm REAL,
                    waist_cm REAL,
                    thigh_cm REAL,
                    calf_cm REAL,
                    note TEXT,
                    source TEXT,
                    created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS gym_sessions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    performed_on TEXT NOT NULL,
                    split_key TEXT NOT NULL,
                    duration_minutes INTEGER,
                    body_weight_kg REAL,
                    notes TEXT,
                    created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS gym_session_exercises (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    session_id INTEGER NOT NULL REFERENCES gym_sessions(id) ON DELETE CASCADE,
                    exercise_name TEXT NOT NULL,
                    weight_kg REAL,
                    weight_text TEXT,
                    completed INTEGER NOT NULL DEFAULT 1,
                    status TEXT NOT NULL DEFAULT 'done',
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS user_exercise_weights (
                    user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    exercise_name TEXT NOT NULL,
                    weight_kg REAL,
                    weight_text TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY (user_id, exercise_name)
                );
                CREATE TABLE IF NOT EXISTS diet_daily_plans (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    plan_date TEXT NOT NULL,
                    meal_count INTEGER NOT NULL,
                    target_calories INTEGER NOT NULL,
                    target_protein_g INTEGER NOT NULL,
                    target_carbs_g INTEGER NOT NULL,
                    target_fat_g INTEGER NOT NULL,
                    current_weight_kg REAL,
                    desired_weight_kg REAL,
                    goal TEXT NOT NULL,
                    message_text TEXT NOT NULL,
                    rationale_text TEXT NOT NULL,
                    source_signature TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    UNIQUE (user_id, plan_date)
                );
                CREATE TABLE IF NOT EXISTS diet_daily_meals (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    plan_id INTEGER NOT NULL REFERENCES diet_daily_plans(id) ON DELETE CASCADE,
                    meal_order INTEGER NOT NULL,
                    slot_key TEXT NOT NULL,
                    meal_label TEXT NOT NULL,
                    template_key TEXT NOT NULL,
                    meal_title TEXT NOT NULL,
                    meal_summary TEXT,
                    prep_text TEXT NOT NULL,
                    portion_multiplier REAL NOT NULL DEFAULT 1.0,
                    calories INTEGER NOT NULL,
                    protein_g INTEGER NOT NULL,
                    carbs_g INTEGER NOT NULL,
                    fat_g INTEGER NOT NULL,
                    ingredients_json TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS diet_daily_meal_items (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    meal_id INTEGER NOT NULL REFERENCES diet_daily_meals(id) ON DELETE CASCADE,
                    item_order INTEGER NOT NULL,
                    item_text TEXT NOT NULL,
                    status TEXT NOT NULL DEFAULT 'pending',
                    share_ratio REAL NOT NULL DEFAULT 0,
                    calories REAL NOT NULL DEFAULT 0
                );
                CREATE TABLE IF NOT EXISTS diet_daily_plan_targets (
                    plan_id INTEGER NOT NULL REFERENCES diet_daily_plans(id) ON DELETE CASCADE,
                    nutrient_code TEXT NOT NULL,
                    target_amount REAL NOT NULL,
                    PRIMARY KEY (plan_id, nutrient_code)
                );
                CREATE TABLE IF NOT EXISTS diet_daily_meal_nutrients (
                    meal_id INTEGER NOT NULL REFERENCES diet_daily_meals(id) ON DELETE CASCADE,
                    nutrient_code TEXT NOT NULL,
                    amount REAL NOT NULL,
                    PRIMARY KEY (meal_id, nutrient_code)
                );
                CREATE TABLE IF NOT EXISTS diet_daily_meal_item_nutrients (
                    meal_item_id INTEGER NOT NULL REFERENCES diet_daily_meal_items(id) ON DELETE CASCADE,
                    nutrient_code TEXT NOT NULL,
                    amount REAL NOT NULL,
                    PRIMARY KEY (meal_item_id, nutrient_code)
                );
                CREATE TABLE IF NOT EXISTS coach_insights (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    agent_key TEXT NOT NULL,
                    context_key TEXT NOT NULL,
                    payload_json TEXT NOT NULL,
                    source_label TEXT NOT NULL,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    UNIQUE (user_id, agent_key, context_key)
                );
                CREATE TABLE IF NOT EXISTS assistant_messages (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    role TEXT NOT NULL,
                    content_text TEXT NOT NULL,
                    metadata_json TEXT,
                    created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS assistant_actions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    message_id INTEGER NOT NULL REFERENCES assistant_messages(id) ON DELETE CASCADE,
                    action_type TEXT NOT NULL,
                    summary_text TEXT NOT NULL,
                    payload_json TEXT NOT NULL,
                    status TEXT NOT NULL DEFAULT 'pending',
                    resolution_note TEXT,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS assistant_attachments (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    message_id INTEGER NOT NULL REFERENCES assistant_messages(id) ON DELETE CASCADE,
                    original_name TEXT NOT NULL,
                    storage_path TEXT NOT NULL,
                    mime_type TEXT,
                    file_ext TEXT,
                    byte_size INTEGER NOT NULL DEFAULT 0,
                    sha256 TEXT NOT NULL,
                    analysis_text TEXT,
                    preview_text TEXT,
                    created_at TEXT NOT NULL,
                    expires_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS plan_overrides (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    domain TEXT NOT NULL,
                    scope TEXT NOT NULL,
                    override_type TEXT NOT NULL,
                    effective_date TEXT,
                    payload_json TEXT NOT NULL,
                    source TEXT,
                    created_at TEXT NOT NULL
                );
                """
            )
            ensure_sqlite_column(connection, "users", "first_name", "TEXT")
            ensure_sqlite_column(connection, "users", "last_name", "TEXT")
            ensure_sqlite_column(connection, "auth_accounts", "role", "TEXT NOT NULL DEFAULT 'member'")
            ensure_sqlite_column(connection, "auth_accounts", "is_active", "INTEGER NOT NULL DEFAULT 1")
            ensure_sqlite_column(connection, "gym_session_exercises", "weight_kg", "REAL")
            ensure_sqlite_column(connection, "gym_session_exercises", "status", "TEXT NOT NULL DEFAULT 'done'")
            ensure_sqlite_column(connection, "user_exercise_weights", "weight_kg", "REAL")
            ensure_sqlite_column(connection, "gym_profiles", "gender", "TEXT NOT NULL DEFAULT 'unspecified'")
            ensure_sqlite_column(connection, "gym_profiles", "age_years", "REAL")
            ensure_sqlite_column(connection, "gym_profiles", "preferred_meals_per_day", f"INTEGER NOT NULL DEFAULT {DEFAULT_DIET_MEALS_PER_DAY}")
            ensure_sqlite_column(connection, "diet_daily_meals", "status", "TEXT NOT NULL DEFAULT 'pending'")
            ensure_sqlite_column(connection, "assistant_attachments", "analysis_json", "TEXT")
            connection.execute(
                """
                UPDATE gym_profiles
                SET gender = 'unspecified'
                WHERE gender IS NULL OR TRIM(gender) = ''
                """
            )
            connection.execute(
                """
                UPDATE gym_session_exercises
                SET status = CASE
                    WHEN COALESCE(completed, 1) = 1 THEN 'done'
                    ELSE 'pending'
                END
                WHERE status IS NULL OR TRIM(status) = ''
                """
            )
            connection.execute(
                """
                UPDATE diet_daily_meals
                SET status = 'pending'
                WHERE status IS NULL OR TRIM(status) = ''
                """
            )
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_assistant_attachments_user_message ON assistant_attachments(user_id, message_id)"
            )
            connection.execute(
                "CREATE INDEX IF NOT EXISTS idx_assistant_attachments_expires_at ON assistant_attachments(expires_at)"
            )
            for table_name in ("gym_session_exercises", "user_exercise_weights"):
                rows = connection.execute(
                    f"SELECT rowid AS row_id, weight_text FROM {table_name} WHERE weight_kg IS NULL AND weight_text IS NOT NULL AND TRIM(weight_text) != ''"
                ).fetchall()
                for row in rows:
                    parsed = parse_weight_kg_from_text(row["weight_text"])
                    if parsed is not None:
                        connection.execute(
                            f"UPDATE {table_name} SET weight_kg = ? WHERE rowid = ?",
                            (parsed, row["row_id"]),
                        )
        COACH_ATTACHMENTS_DIR.mkdir(parents=True, exist_ok=True)
        with gym_knowledge_connection() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS exercise_templates (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    source_name TEXT NOT NULL,
                    sheet_name TEXT NOT NULL,
                    split_key TEXT NOT NULL,
                    title TEXT NOT NULL,
                    next_split_key TEXT,
                    exercise_name TEXT NOT NULL,
                    sets_text TEXT,
                    reps_text TEXT,
                    rest_text TEXT,
                    suggested_weight_text TEXT,
                    equipment_type TEXT,
                    sort_order INTEGER NOT NULL DEFAULT 0
                );
                CREATE TABLE IF NOT EXISTS coaching_rules (
                    rule_key TEXT PRIMARY KEY,
                    rule_value TEXT NOT NULL,
                    description TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS diet_meal_templates (
                    template_key TEXT PRIMARY KEY,
                    slot_key TEXT NOT NULL,
                    title TEXT NOT NULL,
                    goal_bias TEXT NOT NULL,
                    summary TEXT NOT NULL,
                    prep_text TEXT NOT NULL,
                    base_calories INTEGER NOT NULL,
                    protein_g INTEGER NOT NULL,
                    carbs_g INTEGER NOT NULL,
                    fat_g INTEGER NOT NULL,
                    ingredients_json TEXT NOT NULL,
                    sort_order INTEGER NOT NULL DEFAULT 0
                );
                CREATE TABLE IF NOT EXISTS nutrient_definitions (
                    nutrient_code TEXT PRIMARY KEY,
                    label TEXT NOT NULL,
                    unit TEXT NOT NULL,
                    category TEXT NOT NULL,
                    sort_order INTEGER NOT NULL DEFAULT 0,
                    description TEXT
                );
                CREATE TABLE IF NOT EXISTS diet_food_library (
                    food_key TEXT PRIMARY KEY,
                    normalized_name TEXT NOT NULL UNIQUE,
                    food_name TEXT NOT NULL,
                    serving_text TEXT NOT NULL,
                    calories REAL NOT NULL,
                    source_label TEXT,
                    source_kind TEXT NOT NULL DEFAULT 'seed',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS diet_food_library_aliases (
                    alias_normalized TEXT PRIMARY KEY,
                    food_key TEXT NOT NULL REFERENCES diet_food_library(food_key) ON DELETE CASCADE
                );
                CREATE TABLE IF NOT EXISTS diet_food_library_nutrients (
                    food_key TEXT NOT NULL REFERENCES diet_food_library(food_key) ON DELETE CASCADE,
                    nutrient_code TEXT NOT NULL,
                    amount REAL NOT NULL,
                    PRIMARY KEY (food_key, nutrient_code)
                );
                CREATE TABLE IF NOT EXISTS diet_meal_template_nutrients (
                    template_key TEXT NOT NULL,
                    nutrient_code TEXT NOT NULL,
                    amount REAL NOT NULL,
                    PRIMARY KEY (template_key, nutrient_code)
                );
                CREATE TABLE IF NOT EXISTS diet_messages (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    message_text TEXT NOT NULL,
                    sort_order INTEGER NOT NULL DEFAULT 0
                );
                """
            )
            template_count = connection.execute("SELECT COUNT(*) AS count FROM exercise_templates").fetchone()["count"]
            if template_count == 0:
                seed_rows = program_rows_from_examples() or GYM_FALLBACK_TEMPLATES
                connection.executemany(
                    """
                    INSERT INTO exercise_templates (
                        source_name, sheet_name, split_key, title, next_split_key,
                        exercise_name, sets_text, reps_text, rest_text,
                        suggested_weight_text, equipment_type, sort_order
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            row["source_name"],
                            row["sheet_name"],
                            row["split_key"],
                            row["title"],
                            row["next_split_key"],
                            row["exercise_name"],
                            row["sets_text"],
                            row["reps_text"],
                            row["rest_text"],
                            row["suggested_weight_text"],
                            row["equipment_type"],
                            row["sort_order"],
                        )
                        for row in seed_rows
                    ],
                )
            rule_count = connection.execute("SELECT COUNT(*) AS count FROM coaching_rules").fetchone()["count"]
            if rule_count == 0:
                connection.executemany(
                    "INSERT INTO coaching_rules (rule_key, rule_value, description) VALUES (?, ?, ?)",
                    [
                        ("default_start_split", "pull", "Default first training day for a new cycle."),
                        ("reentry_after_days", "4", "If training gap reaches this many days, use a reentry session."),
                        ("reentry_split", "full_body", "Suggested split after a long break from training."),
                        ("legs_next_split", "pull", "Keep upper body emphasis after a leg-heavy day."),
                    ],
                )
            diet_template_count = connection.execute("SELECT COUNT(*) AS count FROM diet_meal_templates").fetchone()["count"]
            if diet_template_count == 0:
                connection.executemany(
                    """
                    INSERT INTO diet_meal_templates (
                        template_key, slot_key, title, goal_bias, summary, prep_text,
                        base_calories, protein_g, carbs_g, fat_g, ingredients_json, sort_order
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    [
                        (
                            item["template_key"],
                            item["slot_key"],
                            item["title"],
                            item["goal_bias"],
                            item["summary"],
                            item["prep_text"],
                            item["base_calories"],
                            item["protein_g"],
                            item["carbs_g"],
                            item["fat_g"],
                            json.dumps(item["ingredients"]),
                            item["sort_order"],
                        )
                        for item in DIET_FALLBACK_MEAL_TEMPLATES
                    ],
                )
            connection.executemany(
                """
                INSERT INTO nutrient_definitions (
                    nutrient_code, label, unit, category, sort_order, description
                ) VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(nutrient_code) DO UPDATE SET
                    label = excluded.label,
                    unit = excluded.unit,
                    category = excluded.category,
                    sort_order = excluded.sort_order,
                    description = excluded.description
                """,
                [
                    (
                        item["code"],
                        item["label"],
                        item["unit"],
                        item["category"],
                        item["sort_order"],
                        item["description"],
                    )
                    for item in DIET_NUTRIENT_DEFINITIONS
                ],
            )
            for item in COACH_DIET_MEAL_FOOD_OPTIONS:
                upsert_diet_food_library_item(
                    item,
                    connection=connection,
                    source_label="Seed food library",
                    source_kind="seed",
                )
            template_rows = connection.execute(
                """
                SELECT template_key, protein_g, carbs_g, fat_g
                FROM diet_meal_templates
                ORDER BY template_key
                """
            ).fetchall()
            fallback_map = {item["template_key"]: item for item in DIET_FALLBACK_MEAL_TEMPLATES}
            connection.executemany(
                """
                INSERT INTO diet_meal_template_nutrients (
                    template_key, nutrient_code, amount
                ) VALUES (?, ?, ?)
                ON CONFLICT(template_key, nutrient_code) DO UPDATE SET amount = excluded.amount
                """,
                [
                    (row["template_key"], nutrient_code, float(amount))
                    for row in template_rows
                    for nutrient_code, amount in {
                        "protein_g": row["protein_g"],
                        "carbs_g": row["carbs_g"],
                        "fat_g": row["fat_g"],
                        "fiber_g": (fallback_map.get(row["template_key"], {}) or {}).get("fiber_g", 0),
                    }.items()
                ],
            )
            diet_message_count = connection.execute("SELECT COUNT(*) AS count FROM diet_messages").fetchone()["count"]
            if diet_message_count == 0:
                connection.executemany(
                    "INSERT INTO diet_messages (message_text, sort_order) VALUES (?, ?)",
                    [(message, index + 1) for index, message in enumerate(DIET_FALLBACK_MESSAGES)],
                )
        ensure_dashboard_seed_accounts()
        GYM_DB_READY = True


def ensure_app_user(username: str):
    ensure_gym_databases()
    current = normalize_username(username) or dashboard_username()
    display_name = display_name_for_username(current)
    timestamp = now_iso()
    with gym_user_connection() as connection:
        connection.execute(
            """
            INSERT INTO users (username, display_name, created_at, updated_at)
            VALUES (?, ?, ?, ?)
            ON CONFLICT(username) DO UPDATE SET display_name = excluded.display_name, updated_at = excluded.updated_at
            """,
            (current, display_name, timestamp, timestamp),
        )
        user_row = connection.execute("SELECT * FROM users WHERE username = ?", (current,)).fetchone()
        connection.execute(
            """
            INSERT INTO gym_profiles (
                user_id, goal, preferred_session_minutes, training_days_per_week, preferred_meals_per_day, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?)
            ON CONFLICT(user_id) DO NOTHING
            """,
            (user_row["id"], "recomp", DEFAULT_GYM_PREFERRED_MINUTES, DEFAULT_GYM_DAYS_PER_WEEK, DEFAULT_DIET_MEALS_PER_DAY, timestamp),
        )
        imported = connection.execute(
            "SELECT COUNT(*) AS count FROM body_measurements WHERE user_id = ?",
            (user_row["id"],),
        ).fetchone()["count"]
        if imported == 0 and current == dashboard_username():
            for item in measurement_rows_from_examples():
                connection.execute(
                    """
                    INSERT INTO body_measurements (
                        user_id, measured_on, weight_kg, arm_cm, chest_cm, waist_cm,
                        thigh_cm, calf_cm, note, source, created_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        user_row["id"],
                        item["measured_on"],
                        safe_float(item["weight_kg"]),
                        safe_float(item["arm_cm"]),
                        safe_float(item["chest_cm"]),
                        safe_float(item["waist_cm"]),
                        safe_float(item["thigh_cm"]),
                        safe_float(item["calf_cm"]),
                        item["note"],
                        item["source"],
                        timestamp,
                    ),
                )
            latest_weight = connection.execute(
                """
                SELECT weight_kg FROM body_measurements
                WHERE user_id = ? AND weight_kg IS NOT NULL
                ORDER BY measured_on DESC, id DESC LIMIT 1
                """,
                (user_row["id"],),
            ).fetchone()
            if latest_weight and latest_weight["weight_kg"] is not None:
                connection.execute(
                    "UPDATE gym_profiles SET current_weight_kg = ?, updated_at = ? WHERE user_id = ?",
                    (latest_weight["weight_kg"], timestamp, user_row["id"]),
                )
        return user_row["id"]


def user_display_name_from_row(user_row: sqlite3.Row | dict | None):
    if not user_row:
        return ""
    username = user_row["username"] if "username" in user_row.keys() else user_row.get("username")
    display_name = user_row["display_name"] if "display_name" in user_row.keys() else user_row.get("display_name")
    first_name = user_row["first_name"] if "first_name" in user_row.keys() else user_row.get("first_name")
    last_name = user_row["last_name"] if "last_name" in user_row.keys() else user_row.get("last_name")
    return compose_display_name(first_name, last_name, username, display_name)


def safe_float(value: str | None):
    if value in {None, ""}:
        return None
    try:
        return round(float(str(value).strip().replace(",", ".")), 2)
    except (TypeError, ValueError):
        return None


def safe_int(value: str | None, default: int | None = None):
    if value in {None, ""}:
        return default
    try:
        return int(float(str(value).strip().replace(",", ".")))
    except (TypeError, ValueError):
        return default


def _legacy_normalize_username(value: str | None):
    return str(value or "").strip().lower()


def normalize_gender(value: str | None):
    current = str(value or "").strip().lower()
    return current if current in HEALTH_GENDER_OPTIONS else "unspecified"


def gender_profile_config(value: str | None):
    current = normalize_gender(value)
    profiles = {
        "male": {
            "label": HEALTH_GENDER_OPTIONS["male"],
            "target_bmi_min": 22.0,
            "target_bmi_max": 26.0,
            "maintenance_factor": 25.0,
            "training_bonus": 75.0,
        },
        "female": {
            "label": HEALTH_GENDER_OPTIONS["female"],
            "target_bmi_min": 20.5,
            "target_bmi_max": 24.0,
            "maintenance_factor": 22.8,
            "training_bonus": 60.0,
        },
        "unspecified": {
            "label": HEALTH_GENDER_OPTIONS["unspecified"],
            "target_bmi_min": 21.0,
            "target_bmi_max": 25.0,
            "maintenance_factor": 24.0,
            "training_bonus": 68.0,
        },
    }
    return {"key": current, **profiles[current]}


def normalize_meals_per_day(value: str | int | None, default: int = DEFAULT_DIET_MEALS_PER_DAY):
    current = safe_int(str(value), default) if value is not None else default
    if current not in DIET_MEAL_COUNT_OPTIONS:
        return 5 if current and current >= 5 else 3
    return current


def normalize_age_years(value: str | int | float | None):
    age_years = safe_float(value)
    if age_years is None:
        return None
    if age_years < 10 or age_years > 120:
        raise ValueError("Enter a valid age in years.")
    return round(float(age_years), 1)


def normalize_gym_exercise_status(value: str | None):
    current = str(value or "").strip().lower()
    if current in {"done", "completed", "complete"}:
        return "done"
    if current in {"skip", "skipped"}:
        return "skipped"
    return "pending"


def normalize_diet_meal_status(value: str | None):
    current = str(value or "").strip().lower()
    if current == "partial":
        return "partial"
    return normalize_gym_exercise_status(value)


def summarize_diet_item_statuses(items: list[dict]):
    total_count = len(items)
    done_count = sum(1 for item in items if normalize_diet_meal_status(item.get("status")) == "done")
    skipped_count = sum(1 for item in items if normalize_diet_meal_status(item.get("status")) == "skipped")
    pending_count = max(0, total_count - done_count - skipped_count)
    if total_count == 0 or pending_count == total_count:
        status = "pending"
    elif done_count == total_count:
        status = "done"
    elif skipped_count == total_count:
        status = "skipped"
    else:
        status = "partial"
    return {
        "status": status,
        "total_count": total_count,
        "done_count": done_count,
        "skipped_count": skipped_count,
        "pending_count": pending_count,
        "is_done": status == "done",
        "is_skipped": status == "skipped",
        "is_partial": status == "partial",
        "is_pending": status == "pending",
    }


def empty_agent_payload(agent_key: str):
    return {
        "headline": f"{AGENT_DISPLAY_NAMES.get(agent_key, 'AI Coach')} has not generated notes for this context yet.",
        "bullets": [],
        "watchout": "",
    }


def normalize_agent_payload(agent_key: str, payload):
    if isinstance(payload, str):
        payload = {"headline": payload}
    if not isinstance(payload, dict):
        payload = {}
    headline = truncate_text(payload.get("headline") or empty_agent_payload(agent_key)["headline"], 180)
    raw_bullets = payload.get("bullets") if isinstance(payload.get("bullets"), list) else []
    bullets = [truncate_text(item, 220) for item in raw_bullets if str(item or "").strip()][:4]
    watchout = truncate_text(payload.get("watchout"), 220)
    return {
        "headline": headline,
        "bullets": bullets,
        "watchout": watchout,
    }


def parse_agent_payload_text(agent_key: str, text: str):
    raw = str(text or "").strip()
    if not raw:
        return empty_agent_payload(agent_key)
    candidate = raw
    match = re.search(r"\{.*\}", raw, flags=re.DOTALL)
    if match:
        candidate = match.group(0)
    try:
        parsed = json.loads(candidate)
    except json.JSONDecodeError:
        lines = [line.strip(" -*\t") for line in raw.splitlines() if line.strip()]
        return normalize_agent_payload(
            agent_key,
            {
                "headline": lines[0] if lines else empty_agent_payload(agent_key)["headline"],
                "bullets": lines[1:4],
                "watchout": lines[4] if len(lines) > 4 else "",
            },
        )
    return normalize_agent_payload(agent_key, parsed)


def _legacy_normalize_assistant_action(raw_action):
    if not isinstance(raw_action, dict):
        return None
    action_type = str(raw_action.get("type") or "").strip().lower()
    if action_type not in ASSISTANT_ALLOWED_ACTIONS:
        return None
    payload = raw_action.get("payload")
    if not isinstance(payload, dict):
        payload = {}
    summary = truncate_text(raw_action.get("summary"), 180)
    if action_type == "record_weight_checkin":
        weight_kg = safe_float(payload.get("weight_kg"))
        if weight_kg is None or weight_kg <= 0:
            return None
        measured_on = str(payload.get("measured_on") or datetime.now().date().isoformat()).strip()
        if not re.fullmatch(r"20\d{2}-\d{2}-\d{2}", measured_on):
            measured_on = datetime.now().date().isoformat()
        note = truncate_text(payload.get("note"), 200)
        payload = {
            "weight_kg": weight_kg,
            "measured_on": measured_on,
            "note": note,
        }
        summary = summary or f"Record weight {weight_kg:.1f} kg for {measured_on}."
    elif action_type == "update_gender":
        gender = normalize_gender(payload.get("gender"))
        if gender not in HEALTH_GENDER_OPTIONS:
            return None
        payload = {"gender": gender}
        summary = summary or f"Set gender to {HEALTH_GENDER_OPTIONS[gender]}."
    elif action_type == "update_goal":
        goal = str(payload.get("goal") or "").strip()
        if goal not in GYM_PROFILE_GOALS:
            return None
        payload = {"goal": goal}
        summary = summary or f"Change goal to {GYM_PROFILE_GOALS[goal]}."
    elif action_type == "update_training_days_per_week":
        training_days = safe_int(payload.get("training_days_per_week"), DEFAULT_GYM_DAYS_PER_WEEK) or DEFAULT_GYM_DAYS_PER_WEEK
        training_days = max(1, min(7, training_days))
        payload = {"training_days_per_week": training_days}
        summary = summary or f"Set training days per week to {training_days}."
    elif action_type == "update_session_minutes":
        session_minutes = safe_int(payload.get("preferred_session_minutes"), DEFAULT_GYM_PREFERRED_MINUTES) or DEFAULT_GYM_PREFERRED_MINUTES
        session_minutes = max(20, min(180, session_minutes))
        payload = {"preferred_session_minutes": session_minutes}
        summary = summary or f"Set session length to {session_minutes} minutes."
    elif action_type == "update_meals_per_day":
        meals_per_day = normalize_meals_per_day(payload.get("meals_per_day"), DEFAULT_DIET_MEALS_PER_DAY)
        payload = {"meals_per_day": meals_per_day}
        summary = summary or f"Set meals per day to {meals_per_day}."
    elif action_type == "refresh_diet_plan":
        plan_date = str(payload.get("plan_date") or datetime.now().date().isoformat()).strip()
        if not re.fullmatch(r"20\d{2}-\d{2}-\d{2}", plan_date):
            plan_date = datetime.now().date().isoformat()
        payload = {"plan_date": plan_date}
        summary = summary or f"Refresh the diet plan for {plan_date}."
    return {
        "type": action_type,
        "summary": summary,
        "payload": payload,
    }


def _legacy_empty_assistant_response():
    return {
        "reply": "I can help with your health, diet, and gym data and ask for confirmation before making changes.",
        "actions": [],
    }


def _legacy_parse_assistant_response_text(text: str):
    raw = str(text or "").strip()
    if not raw:
        return empty_assistant_response()
    candidate = raw
    match = re.search(r"\{.*\}", raw, flags=re.DOTALL)
    if match:
        candidate = match.group(0)
    try:
        parsed = json.loads(candidate)
    except json.JSONDecodeError:
        lines = [line.strip() for line in raw.splitlines() if line.strip()]
        return {
            "reply": truncate_text(lines[0] if lines else empty_assistant_response()["reply"], 900),
            "actions": [],
        }
    reply = truncate_text(parsed.get("reply"), 900) or empty_assistant_response()["reply"]
    raw_actions = parsed.get("actions") if isinstance(parsed.get("actions"), list) else []
    actions = [item for item in (normalize_assistant_action(action) for action in raw_actions[:4]) if item]
    return {
        "reply": reply,
        "actions": actions,
    }


def load_gym_profile(username: str):
    user_id = ensure_app_user(username)
    with gym_user_connection() as connection:
        user_row = connection.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        profile_row = connection.execute("SELECT * FROM gym_profiles WHERE user_id = ?", (user_id,)).fetchone()
        latest_measurement = connection.execute(
            """
            SELECT * FROM body_measurements
            WHERE user_id = ? ORDER BY measured_on DESC, id DESC LIMIT 1
            """,
            (user_id,),
        ).fetchone()
    return {
        "user_id": user_id,
        "username": user_row["username"],
        "display_name": user_display_name_from_row(user_row),
        "first_name": (user_row["first_name"] or "").strip(),
        "last_name": (user_row["last_name"] or "").strip(),
        "height_cm": profile_row["height_cm"],
        "age_years": profile_row["age_years"],
        "current_weight_kg": profile_row["current_weight_kg"],
        "desired_weight_kg": profile_row["desired_weight_kg"],
        "gender": normalize_gender(profile_row["gender"]),
        "gender_label": HEALTH_GENDER_OPTIONS.get(normalize_gender(profile_row["gender"]), HEALTH_GENDER_OPTIONS["unspecified"]),
        "goal": profile_row["goal"],
        "goal_label": GYM_PROFILE_GOALS.get(profile_row["goal"], profile_row["goal"]),
        "preferred_session_minutes": profile_row["preferred_session_minutes"],
        "training_days_per_week": profile_row["training_days_per_week"],
        "preferred_meals_per_day": normalize_meals_per_day(profile_row["preferred_meals_per_day"], DEFAULT_DIET_MEALS_PER_DAY),
        "notes": profile_row["notes"] or "",
        "updated_at": profile_row["updated_at"],
        "latest_measurement": dict(latest_measurement) if latest_measurement else None,
    }


def list_gym_templates():
    ensure_gym_databases()
    with gym_knowledge_connection() as connection:
        rows = connection.execute(
            """
            SELECT split_key, title, next_split_key, equipment_type, exercise_name,
                   sets_text, reps_text, rest_text, suggested_weight_text, source_name,
                   sort_order
            FROM exercise_templates
            ORDER BY split_key, sort_order, id
            """
        ).fetchall()
    grouped = {}
    for row in rows:
        grouped.setdefault(row["split_key"], {"split_key": row["split_key"], "title": row["title"], "next_split_key": row["next_split_key"], "source_name": row["source_name"], "exercises": []})
        grouped[row["split_key"]]["exercises"].append(dict(row))
    return list(grouped.values())


def list_gym_rules():
    ensure_gym_databases()
    with gym_knowledge_connection() as connection:
        rows = connection.execute("SELECT * FROM coaching_rules ORDER BY rule_key").fetchall()
    return {row["rule_key"]: row["rule_value"] for row in rows}


def user_exercise_weights(username: str):
    user_id = ensure_app_user(username)
    with gym_user_connection() as connection:
        rows = connection.execute(
            """
            SELECT exercise_name, weight_kg, weight_text, updated_at
            FROM user_exercise_weights
            WHERE user_id = ?
            """,
            (user_id,),
        ).fetchall()
    return {
        row["exercise_name"]: {
            "weight_kg": row["weight_kg"] if row["weight_kg"] is not None else parse_weight_kg_from_text(row["weight_text"]),
            "updated_at": row["updated_at"],
        }
        for row in rows
    }


def exercise_weight_history(username: str, exercise_names: list[str], limit_per_exercise: int = 10):
    user_id = ensure_app_user(username)
    normalized_names = [item.strip() for item in exercise_names if str(item or "").strip()]
    if not normalized_names:
        return {}
    placeholders = ", ".join("?" for _ in normalized_names)
    with gym_user_connection() as connection:
        rows = connection.execute(
            f"""
            SELECT gse.exercise_name, gs.performed_on, gse.weight_kg, gse.id
            FROM gym_session_exercises AS gse
            JOIN gym_sessions AS gs ON gs.id = gse.session_id
            WHERE gs.user_id = ?
              AND gse.status = 'done'
              AND gse.weight_kg IS NOT NULL
              AND gse.exercise_name IN ({placeholders})
            ORDER BY gse.exercise_name, gs.performed_on DESC, gse.id DESC
            """,
            (user_id, *normalized_names),
        ).fetchall()
    grouped = {name: [] for name in normalized_names}
    for row in rows:
        bucket = grouped.setdefault(row["exercise_name"], [])
        if len(bucket) >= limit_per_exercise:
            continue
        bucket.append(
            {
                "date": row["performed_on"],
                "value": round(float(row["weight_kg"]), 1),
            }
        )
    return {
        name: list(reversed(points))
        for name, points in grouped.items()
    }


def current_session_logged_weights(username: str, split_key: str, performed_on: str):
    user_id = ensure_app_user(username)
    with gym_user_connection() as connection:
        rows = connection.execute(
            """
            SELECT gse.exercise_name, gse.weight_kg, gse.weight_text, gse.updated_at, gse.status, gse.completed
            FROM gym_session_exercises AS gse
            JOIN gym_sessions AS gs ON gs.id = gse.session_id
            WHERE gs.user_id = ? AND gs.split_key = ? AND gs.performed_on = ?
            """,
            (user_id, split_key, performed_on),
        ).fetchall()
    return {
        row["exercise_name"]: {
            "weight_kg": row["weight_kg"] if row["weight_kg"] is not None else parse_weight_kg_from_text(row["weight_text"]),
            "updated_at": row["updated_at"],
            "status": normalize_gym_exercise_status(
                row["status"] if row["status"] is not None else ("done" if row["completed"] else "pending")
            ),
        }
        for row in rows
    }


def gym_session_progress(exercise_names: list[str], session_entries: dict):
    normalized_names = [item for item in exercise_names if item]
    total_count = len(normalized_names)
    done_count = 0
    skipped_count = 0
    for name in normalized_names:
        status = normalize_gym_exercise_status((session_entries.get(name) or {}).get("status"))
        if status == "done":
            done_count += 1
        elif status == "skipped":
            skipped_count += 1
    resolved_count = done_count + skipped_count
    return {
        "total_count": total_count,
        "done_count": done_count,
        "skipped_count": skipped_count,
        "resolved_count": resolved_count,
        "is_complete": total_count > 0 and resolved_count >= total_count,
        "has_activity": bool(session_entries),
    }


def average_numbers_from_text(value: str | None, fallback: float | None = None):
    matches = [float(item) for item in re.findall(r"\d+(?:\.\d+)?", str(value or "").replace(",", "."))]
    if not matches:
        return fallback
    return sum(matches) / len(matches)


# Approximate resistance-training MET anchors from the Compendium of Physical Activities.
def estimate_gym_standard_met(exercise: dict):
    name = str(exercise.get("exercise_name") or "").strip().lower()
    equipment = str(exercise.get("equipment_type") or "").strip().lower()
    light_bodyweight_markers = ("plank", "dead bug", "flutter", "crunch", "hollow", "bird dog", "side plank")
    vigorous_markers = ("burpee", "jump", "thruster", "clean", "snatch", "kettlebell swing")
    heavy_resistance_markers = ("squat", "deadlift", "leg press", "walking lunge", "hip thrust", "step up")
    if any(marker in name for marker in light_bodyweight_markers):
        return 2.8
    if any(marker in name for marker in vigorous_markers):
        return 6.0 if equipment != "bodyweight" else 6.5
    if any(marker in name for marker in heavy_resistance_markers):
        return 5.0
    if equipment == "bodyweight":
        return 3.8
    return 3.5


def predicted_profile_rmr_ml_kg_min(profile: dict):
    gender = str(profile.get("gender") or "").strip().lower()
    weight_kg = safe_float(profile.get("current_weight_kg"))
    height_cm = safe_float(profile.get("height_cm"))
    age_years = safe_float(profile.get("age_years"))
    if gender not in {"male", "female"} or weight_kg is None or height_cm is None or age_years is None or weight_kg <= 0:
        return None
    if gender == "male":
        kcal_day = 66.4730 + (5.0033 * height_cm) + (13.7516 * weight_kg) - (6.7550 * age_years)
    else:
        kcal_day = 655.0955 + (1.8496 * height_cm) + (9.5634 * weight_kg) - (4.6756 * age_years)
    kcal_min = kcal_day / 1440.0
    liters_per_min = kcal_min / 5.0
    return (liters_per_min / weight_kg) * 1000.0


def predicted_profile_bmr_kcal_day(profile: dict):
    gender = normalize_gender(profile.get("gender"))
    weight_kg = safe_float(profile.get("current_weight_kg"))
    height_cm = safe_float(profile.get("height_cm"))
    age_years = safe_float(profile.get("age_years"))
    if gender not in {"male", "female"} or weight_kg is None or height_cm is None or age_years is None or weight_kg <= 0:
        return None
    base_kcal_day = (10.0 * weight_kg) + (6.25 * height_cm) - (5.0 * age_years)
    return base_kcal_day + (5.0 if gender == "male" else -161.0)


def diet_activity_multiplier(training_days_per_week):
    training_days = safe_int(training_days_per_week, DEFAULT_GYM_DAYS_PER_WEEK) or DEFAULT_GYM_DAYS_PER_WEEK
    if training_days <= 1:
        return 1.35
    if training_days <= 3:
        return 1.45
    if training_days <= 5:
        return 1.55
    return 1.65


def estimate_gym_corrected_met(profile: dict, exercise: dict):
    standard_met = estimate_gym_standard_met(exercise)
    predicted_rmr = predicted_profile_rmr_ml_kg_min(profile)
    if predicted_rmr is None or predicted_rmr <= 0:
        return standard_met
    return round(standard_met * (3.5 / predicted_rmr), 2)


def estimate_gym_exercise_effort_score(exercise: dict):
    sets = average_numbers_from_text(exercise.get("sets_text"), 3.0) or 3.0
    reps_text = str(exercise.get("reps_text") or "")
    reps_lower = reps_text.lower()
    if "sec" in reps_lower or "hold" in reps_lower:
        active_seconds_per_set = max(20.0, min(90.0, average_numbers_from_text(reps_text, 30.0) or 30.0))
    else:
        reps = average_numbers_from_text(reps_text, 10.0) or 10.0
        active_seconds_per_set = max(18.0, min(70.0, reps * 4.0))
    rest_seconds = max(20.0, min(180.0, average_numbers_from_text(exercise.get("rest_text"), 75.0) or 75.0))
    met = estimate_gym_standard_met(exercise)
    intensity_multiplier = 0.85 + (met / 10.0)
    transition_seconds = 18.0
    return max(1.0, sets * (active_seconds_per_set + rest_seconds + transition_seconds) * intensity_multiplier)


def estimate_gym_session_burn(
    profile: dict,
    exercise_rows: list[dict],
    duration_minutes: float | int | None,
    *,
    body_weight_kg: float | int | None = None,
):
    duration_minutes = safe_float(duration_minutes)
    weight_kg = safe_float(body_weight_kg)
    if weight_kg is None or weight_kg <= 0:
        weight_kg = safe_float(profile.get("current_weight_kg"))
    if not exercise_rows or duration_minutes is None or duration_minutes <= 0:
        return {
            "available": False,
            "actual_kcal": 0.0,
            "planned_kcal": 0.0,
            "actual_text": "n/a",
            "planned_text": "n/a",
            "summary_text": "n/a",
            "note": "No active session is available to estimate yet.",
        }
    if weight_kg is None or weight_kg <= 0:
        return {
            "available": False,
            "actual_kcal": 0.0,
            "planned_kcal": 0.0,
            "actual_text": "Add weight in Health",
            "planned_text": "n/a",
            "summary_text": "Add weight in Health",
            "note": "Current body weight is required for the calorie estimate.",
        }
    weighted_rows = []
    total_score = 0.0
    for exercise in exercise_rows:
        met = estimate_gym_corrected_met(profile, exercise)
        score = estimate_gym_exercise_effort_score(exercise)
        weighted_rows.append((exercise, met, score))
        total_score += score
    if total_score <= 0:
        total_score = float(len(weighted_rows) or 1)
    actual_kcal = 0.0
    planned_kcal = 0.0
    for exercise, met, score in weighted_rows:
        allocated_minutes = float(duration_minutes) * (score / total_score)
        estimated_kcal = met * 3.5 * weight_kg / 200.0 * allocated_minutes
        planned_kcal += estimated_kcal
        if normalize_gym_exercise_status(exercise.get("status")) == "done":
            actual_kcal += estimated_kcal
    actual_kcal = round(actual_kcal, 1)
    planned_kcal = round(planned_kcal, 1)
    return {
        "available": True,
        "actual_kcal": actual_kcal,
        "planned_kcal": planned_kcal,
        "actual_text": f"{actual_kcal:.0f} kcal",
        "planned_text": f"{planned_kcal:.0f} kcal",
        "summary_text": f"{actual_kcal:.0f} / {planned_kcal:.0f} kcal",
        "note": "Estimated from Compendium MET values, your current body weight, and the planned session length.",
    }


def apply_gym_exercise_override_events(exercise_rows: list[dict], events: list[dict], *, note_prefix: str):
    applied_notes = []
    for row in events:
        payload = row.get("payload") or {}
        split_key = normalize_split_key(payload.get("split_key"))
        if split_key and exercise_rows and split_key != normalize_split_key(exercise_rows[0].get("split_key")):
            continue
        exercise_name = str(payload.get("exercise_name") or "").strip().lower()
        if not exercise_name:
            continue
        target = next((item for item in exercise_rows if str(item.get("exercise_name") or "").strip().lower() == exercise_name), None)
        if not target:
            continue
        override_type = row.get("override_type")
        if override_type == "replace_gym_exercise":
            replacement_name = str(payload.get("replacement_exercise_name") or "").strip()
            if replacement_name:
                original_name = target["exercise_name"]
                target["exercise_name"] = replacement_name
                applied_notes.append(f"{note_prefix}: {original_name} -> {replacement_name}.")
        elif override_type == "update_gym_exercise_scheme":
            changed = False
            for key in ("sets_text", "reps_text", "rest_text"):
                value = str(payload.get(key) or "").strip()
                if not value:
                    continue
                target[key] = value
                changed = True
            if changed:
                applied_notes.append(f"{note_prefix}: {target['exercise_name']} scheme updated.")
        elif override_type == "update_gym_exercise_load":
            suggested_weight_kg = safe_float(payload.get("suggested_weight_kg"))
            if suggested_weight_kg is not None and suggested_weight_kg >= 0:
                target["suggested_weight_kg"] = suggested_weight_kg
                target["suggested_weight_text"] = format_weight_kg(suggested_weight_kg)
                applied_notes.append(f"{note_prefix}: {target['exercise_name']} suggested load {suggested_weight_kg:.1f} kg.")
    return applied_notes


def ensure_gym_session(connection, user_id: int, performed_on: str, split_key: str, duration_minutes: int | None, body_weight_kg: float | None):
    row = connection.execute(
        """
        SELECT id FROM gym_sessions
        WHERE user_id = ? AND performed_on = ? AND split_key = ?
        ORDER BY id DESC LIMIT 1
        """,
        (user_id, performed_on, split_key),
    ).fetchone()
    if row:
        return row["id"]
    cursor = connection.execute(
        """
        INSERT INTO gym_sessions (
            user_id, performed_on, split_key, duration_minutes, body_weight_kg, notes, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?)
        """,
        (user_id, performed_on, split_key, duration_minutes, body_weight_kg, "", now_iso()),
    )
    return cursor.lastrowid


def planned_gym_exercise_names_for_session(
    connection: sqlite3.Connection,
    user_id: int,
    split_key: str,
    effective_date: str,
    template_map: dict,
):
    template = template_map.get(split_key) or next(iter(template_map.values()), None)
    if not template:
        return []
    exercise_rows = [
        {
            **row,
            "suggested_weight_kg": parse_weight_kg_from_text(row.get("suggested_weight_text")),
        }
        for row in template.get("exercises", [])[:6]
    ]
    override_state = collect_gym_override_state_for_user(connection, user_id, effective_date)
    apply_gym_exercise_override_events(exercise_rows, override_state.get("long_term_events", []), note_prefix="Regular plan")
    apply_gym_exercise_override_events(exercise_rows, override_state.get("today_events", []), note_prefix="Today only")
    return [row.get("exercise_name") for row in exercise_rows if row.get("exercise_name")]


def recent_gym_sessions(username: str, limit: int = 6):
    user_id = ensure_app_user(username)
    with gym_user_connection() as connection:
        rows = connection.execute(
            """
            SELECT
                gs.*,
                COALESCE(SUM(CASE WHEN gse.status = 'done' THEN 1 ELSE 0 END), 0) AS done_count,
                COALESCE(SUM(CASE WHEN gse.status = 'skipped' THEN 1 ELSE 0 END), 0) AS skipped_count,
                COALESCE(SUM(CASE WHEN gse.status = 'pending' THEN 1 ELSE 0 END), 0) AS pending_count,
                COUNT(gse.id) AS total_count
            FROM gym_sessions AS gs
            LEFT JOIN gym_session_exercises AS gse ON gse.session_id = gs.id
            WHERE gs.user_id = ?
            GROUP BY gs.id
            ORDER BY
                gs.performed_on DESC,
                CASE
                    WHEN COALESCE(SUM(CASE WHEN gse.status IN ('done', 'skipped') THEN 1 ELSE 0 END), 0) > 0 THEN 1
                    ELSE 0
                END DESC,
                gs.id DESC
            LIMIT ?
            """,
            (user_id, limit),
        ).fetchall()
    sessions = []
    for row in rows:
        item = dict(row)
        item["split_label"] = split_label(item["split_key"])
        sessions.append(item)
    return sessions


def normalize_gym_day(value: str | None):
    raw = str(value or "").strip().lower()
    if not raw or raw == "today":
        return "today"
    try:
        return datetime.fromisoformat(raw).date().isoformat()
    except ValueError:
        return "today"


def load_latest_gym_session_for_day(connection: sqlite3.Connection, user_id: int, performed_on: str):
    row = connection.execute(
        """
        SELECT
            gs.*,
            COALESCE(SUM(CASE WHEN gse.status = 'done' THEN 1 ELSE 0 END), 0) AS done_count,
            COALESCE(SUM(CASE WHEN gse.status = 'skipped' THEN 1 ELSE 0 END), 0) AS skipped_count,
            COALESCE(SUM(CASE WHEN gse.status = 'pending' THEN 1 ELSE 0 END), 0) AS pending_count,
            COUNT(gse.id) AS total_count
        FROM gym_sessions AS gs
        LEFT JOIN gym_session_exercises AS gse ON gse.session_id = gs.id
        WHERE gs.user_id = ? AND gs.performed_on = ?
        GROUP BY gs.id
        ORDER BY
            CASE
                WHEN COALESCE(SUM(CASE WHEN gse.status IN ('done', 'skipped') THEN 1 ELSE 0 END), 0) > 0 THEN 1
                ELSE 0
            END DESC,
            gs.id DESC
        LIMIT 1
        """,
        (user_id, performed_on),
    ).fetchone()
    if not row:
        return None
    item = dict(row)
    item["split_label"] = split_label(item["split_key"])
    return item


def recent_body_measurements(username: str, limit: int = 8):
    user_id = ensure_app_user(username)
    with gym_user_connection() as connection:
        rows = connection.execute(
            """
            SELECT *
            FROM body_measurements
            WHERE user_id = ?
            ORDER BY measured_on DESC, id DESC
            LIMIT ?
            """,
            (user_id, limit),
        ).fetchall()
    return [dict(row) for row in rows]


def build_body_measurement_chart(field: dict, measurements: list[dict]):
    points = []
    for item in measurements:
        value = item.get(field["key"])
        if value is None:
            continue
        points.append(
            {
                "date": item.get("measured_on") or "",
                "value": round(float(value), 1),
            }
        )
    latest_value = points[-1]["value"] if points else None
    return {
        **field,
        "points": points,
        "point_count": len(points),
        "latest_value_text": f"{latest_value:.1f} {field['unit']}" if latest_value is not None else "No data yet",
    }


def load_coach_insight(username: str, agent_key: str, context_key: str):
    user_id = ensure_app_user(username)
    with gym_user_connection() as connection:
        row = connection.execute(
            """
            SELECT *
            FROM coach_insights
            WHERE user_id = ? AND agent_key = ? AND context_key = ?
            ORDER BY id DESC LIMIT 1
            """,
            (user_id, agent_key, context_key),
        ).fetchone()
    if not row:
        return None
    try:
        payload = json.loads(row["payload_json"] or "{}")
    except json.JSONDecodeError:
        payload = empty_agent_payload(agent_key)
    return {
        "agent_key": row["agent_key"],
        "context_key": row["context_key"],
        "payload": normalize_agent_payload(agent_key, payload),
        "source_label": row["source_label"],
        "updated_at": row["updated_at"],
        "updated_at_text": display_timestamp(row["updated_at"]),
    }


def save_coach_insight(username: str, agent_key: str, context_key: str, payload: dict, source_label: str = "Codex on MSI"):
    user_id = ensure_app_user(username)
    timestamp = now_iso()
    normalized_payload = normalize_agent_payload(agent_key, payload)
    serialized = json.dumps(normalized_payload, sort_keys=True)
    with gym_user_connection() as connection:
        existing = connection.execute(
            """
            SELECT id
            FROM coach_insights
            WHERE user_id = ? AND agent_key = ? AND context_key = ?
            ORDER BY id DESC LIMIT 1
            """,
            (user_id, agent_key, context_key),
        ).fetchone()
        if existing:
            connection.execute(
                """
                UPDATE coach_insights
                SET payload_json = ?, source_label = ?, updated_at = ?
                WHERE id = ?
                """,
                (serialized, source_label, timestamp, existing["id"]),
            )
        else:
            connection.execute(
                """
                INSERT INTO coach_insights (
                    user_id, agent_key, context_key, payload_json, source_label, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (user_id, agent_key, context_key, serialized, source_label, timestamp, timestamp),
            )
    return load_coach_insight(username, agent_key, context_key)


def save_assistant_message(username: str, role: str, content_text: str, metadata: dict | None = None):
    user_id = ensure_app_user(username)
    timestamp = now_iso()
    with gym_user_connection() as connection:
        cursor = connection.execute(
            """
            INSERT INTO assistant_messages (
                user_id, role, content_text, metadata_json, created_at
            ) VALUES (?, ?, ?, ?, ?)
            """,
            (
                user_id,
                str(role or "assistant").strip().lower(),
                truncate_text(content_text, 4000),
                json.dumps(metadata or {}, sort_keys=True),
                timestamp,
            ),
        )
    return cursor.lastrowid


def save_assistant_actions(username: str, message_id: int, actions: list[dict]):
    user_id = ensure_app_user(username)
    timestamp = now_iso()
    saved_ids = []
    if not actions:
        return saved_ids
    with gym_user_connection() as connection:
        for action in actions:
            cursor = connection.execute(
                """
                INSERT INTO assistant_actions (
                    user_id, message_id, action_type, summary_text, payload_json, status, resolution_note, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, 'pending', '', ?, ?)
                """,
                (
                    user_id,
                    message_id,
                    action["type"],
                    action["summary"],
                    json.dumps(action["payload"], sort_keys=True),
                    timestamp,
                    timestamp,
                ),
            )
            saved_ids.append(cursor.lastrowid)
    return saved_ids


def delete_assistant_attachment_files(rows):
    for row in rows:
        storage_path = str((row or {}).get("storage_path") or "").strip()
        if not storage_path:
            continue
        file_path = APP_ROOT / storage_path
        try:
            if file_path.exists():
                file_path.unlink()
        except OSError:
            pass


def prune_assistant_attachments(username: str):
    user_id = ensure_app_user(username)
    now_value = datetime.now(timezone.utc)
    rows_to_delete = []
    with gym_user_connection() as connection:
        rows = connection.execute(
            """
            SELECT *
            FROM assistant_attachments
            WHERE user_id = ?
            ORDER BY created_at ASC, id ASC
            """,
            (user_id,),
        ).fetchall()
        kept_rows = []
        total_bytes = 0
        for row in rows:
            storage_path = str(row["storage_path"] or "").strip()
            file_path = APP_ROOT / storage_path if storage_path else None
            expires_at = parse_iso_utc(row["expires_at"])
            if not storage_path or not file_path or not file_path.exists() or (expires_at and expires_at <= now_value):
                rows_to_delete.append(dict(row))
                continue
            kept_rows.append(dict(row))
            total_bytes += max(0, int(row["byte_size"] or 0))
        overflow_rows = []
        while total_bytes > COACH_ATTACHMENT_MAX_TOTAL_BYTES_PER_USER and kept_rows:
            oldest = kept_rows.pop(0)
            overflow_rows.append(oldest)
            total_bytes -= max(0, int(oldest.get("byte_size") or 0))
        rows_to_delete.extend(overflow_rows)
        if rows_to_delete:
            placeholders = ", ".join("?" for _ in rows_to_delete)
            connection.execute(
                f"DELETE FROM assistant_attachments WHERE id IN ({placeholders})",
                [row["id"] for row in rows_to_delete],
            )
    delete_assistant_attachment_files(rows_to_delete)


def save_assistant_attachments(username: str, message_id: int, prepared_attachments: list[dict]):
    if not prepared_attachments:
        return []
    ensure_gym_databases()
    prune_assistant_attachments(username)
    user_id = ensure_app_user(username)
    user_dir = coach_attachment_user_dir(username)
    user_dir.mkdir(parents=True, exist_ok=True)
    rows_to_delete = []
    current_total_bytes = 0
    with gym_user_connection() as connection:
        total_row = connection.execute(
            "SELECT COALESCE(SUM(byte_size), 0) AS total_bytes FROM assistant_attachments WHERE user_id = ?",
            (user_id,),
        ).fetchone()
        current_total_bytes = int((total_row or {})["total_bytes"] if total_row else 0)
        incoming_total = sum(int(item.get("byte_size") or 0) for item in prepared_attachments)
        if current_total_bytes + incoming_total > COACH_ATTACHMENT_MAX_TOTAL_BYTES_PER_USER:
            existing_rows = connection.execute(
                """
                SELECT *
                FROM assistant_attachments
                WHERE user_id = ?
                ORDER BY created_at ASC, id ASC
                """,
                (user_id,),
            ).fetchall()
            remaining_total = current_total_bytes
            for row in existing_rows:
                if remaining_total + incoming_total <= COACH_ATTACHMENT_MAX_TOTAL_BYTES_PER_USER:
                    break
                rows_to_delete.append(dict(row))
                remaining_total -= max(0, int(row["byte_size"] or 0))
            if rows_to_delete:
                placeholders = ", ".join("?" for _ in rows_to_delete)
                connection.execute(
                    f"DELETE FROM assistant_attachments WHERE id IN ({placeholders})",
                    [row["id"] for row in rows_to_delete],
                )
                current_total_bytes = max(0, remaining_total)
    delete_assistant_attachment_files(rows_to_delete)
    incoming_total = sum(int(item.get("byte_size") or 0) for item in prepared_attachments)
    if current_total_bytes + incoming_total > COACH_ATTACHMENT_MAX_TOTAL_BYTES_PER_USER:
        raise ValueError(
            f"Coach file storage is full for this user. Keep uploads under {format_file_size(COACH_ATTACHMENT_MAX_TOTAL_BYTES_PER_USER)} total or wait for older files to expire."
        )
    timestamp = now_iso()
    expires_at = (datetime.now(timezone.utc) + timedelta(seconds=COACH_ATTACHMENT_RETENTION_SECONDS)).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    saved_rows = []
    with gym_user_connection() as connection:
        for item in prepared_attachments:
            storage_key = str(item.get("storage_key") or f"{uuid.uuid4().hex}{item.get('file_ext') or ''}")
            file_path = user_dir / storage_key
            file_path.write_bytes(item["data"])
            relative_path = file_path.relative_to(APP_ROOT).as_posix()
            cursor = connection.execute(
                """
                INSERT INTO assistant_attachments (
                    user_id, message_id, original_name, storage_path, mime_type, file_ext, byte_size,
                    sha256, analysis_text, analysis_json, preview_text, created_at, expires_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    user_id,
                    message_id,
                    item["original_name"],
                    relative_path,
                    item.get("mime_type"),
                    item.get("file_ext"),
                    int(item.get("byte_size") or 0),
                    item.get("sha256") or "",
                    item.get("analysis_text") or "",
                    item.get("analysis_json") or "",
                    item.get("preview_text") or "",
                    timestamp,
                    expires_at,
                ),
            )
            saved_rows.append(
                {
                    "id": cursor.lastrowid,
                    "original_name": item["original_name"],
                    "storage_path": relative_path,
                    "mime_type": item.get("mime_type"),
                    "file_ext": item.get("file_ext"),
                    "byte_size": int(item.get("byte_size") or 0),
                    "analysis_text": item.get("analysis_text") or "",
                    "analysis_json": item.get("analysis_json") or "",
                    "preview_text": item.get("preview_text") or "",
                    "created_at": timestamp,
                    "expires_at": expires_at,
                }
            )
    return saved_rows


def load_assistant_attachments_for_messages(username: str, message_ids: list[int], include_analysis_text: bool = False):
    if not message_ids:
        return {}
    user_id = ensure_app_user(username)
    placeholders = ", ".join("?" for _ in message_ids)
    with gym_user_connection() as connection:
        rows = connection.execute(
            f"""
            SELECT *
            FROM assistant_attachments
            WHERE user_id = ? AND message_id IN ({placeholders})
            ORDER BY id ASC
            """,
            (user_id, *message_ids),
        ).fetchall()
    attachments_by_message = {message_id: [] for message_id in message_ids}
    for row in rows:
        try:
            analysis_payload = json.loads(row["analysis_json"] or "{}")
        except json.JSONDecodeError:
            analysis_payload = {}
        if not isinstance(analysis_payload, dict):
            analysis_payload = {}
        is_image = coach_attachment_is_image(row["file_ext"], row["mime_type"])
        attachment = {
            "id": row["id"],
            "original_name": row["original_name"],
            "mime_type": row["mime_type"] or "",
            "file_ext": row["file_ext"] or "",
            "byte_size": int(row["byte_size"] or 0),
            "size_text": format_file_size(row["byte_size"]),
            "kind_label": coach_attachment_kind_label(row["file_ext"]),
            "preview_text": row["preview_text"] or "",
            "created_at_text": display_timestamp(row["created_at"]),
            "expires_at_text": display_timestamp(row["expires_at"]),
            "download_url": f"/assistant/attachment/{row['id']}",
            "thumbnail_url": f"/assistant/attachment/{row['id']}?inline=1" if is_image else "",
            "is_image": is_image,
            "analysis_summary": truncate_text(analysis_payload.get("summary"), 240),
            "analysis_items": analysis_payload.get("items") if isinstance(analysis_payload.get("items"), list) else [],
        }
        if include_analysis_text:
            attachment["analysis_text"] = truncate_text(row["analysis_text"], 2000)
        attachments_by_message.setdefault(row["message_id"], []).append(attachment)
    return attachments_by_message


def sync_profile_current_weight_from_measurements(connection: sqlite3.Connection, user_id: int):
    latest_row = connection.execute(
        """
        SELECT weight_kg
        FROM body_measurements
        WHERE user_id = ? AND weight_kg IS NOT NULL
        ORDER BY measured_on DESC, id DESC
        LIMIT 1
        """,
        (user_id,),
    ).fetchone()
    if not latest_row or latest_row["weight_kg"] is None:
        return None
    latest_weight = float(latest_row["weight_kg"])
    connection.execute(
        "UPDATE gym_profiles SET current_weight_kg = ?, updated_at = ? WHERE user_id = ?",
        (latest_weight, now_iso(), user_id),
    )
    return latest_weight


def clear_coach_insights(connection: sqlite3.Connection, user_id: int, agent_keys: tuple[str, ...] | None = None, context_key: str | None = None):
    conditions = ["user_id = ?"]
    parameters: list[object] = [user_id]
    if agent_keys:
        placeholders = ", ".join("?" for _ in agent_keys)
        conditions.append(f"agent_key IN ({placeholders})")
        parameters.extend(agent_keys)
    if context_key:
        conditions.append("context_key = ?")
        parameters.append(context_key)
    connection.execute(
        f"DELETE FROM coach_insights WHERE {' AND '.join(conditions)}",
        parameters,
    )


def normalize_coach_scope(value: str | None):
    current = str(value or "").strip().lower()
    return current if current in COACH_SCOPE_OPTIONS else None


def normalize_diet_meal_customization_mode(value: str | None):
    current = str(value or "").strip().lower()
    return current if current in DIET_MEAL_CUSTOMIZATION_MODES else None


def normalize_diet_preference_keys(values):
    raw_items = values if isinstance(values, list) else [values]
    normalized = []
    for item in raw_items:
        key = re.sub(r"[^a-z0-9]+", "_", str(item or "").strip().lower()).strip("_")
        if key in DIET_PREFERENCE_OPTIONS:
            normalized.append(key)
    if not normalized:
        return []
    if "balanced" in normalized:
        return ["balanced"]
    macro_preferences = [key for key in normalized if key in {"lower_carb", "higher_carb"}]
    if len(macro_preferences) > 1:
        latest_macro = macro_preferences[-1]
        normalized = [key for key in normalized if key not in {"lower_carb", "higher_carb"}]
        normalized.append(latest_macro)
    ordered = []
    seen = set()
    for key in normalized:
        if key in seen:
            continue
        seen.add(key)
        ordered.append(key)
    return ordered


def effective_diet_preference_keys(preference_keys: list[str] | None):
    keys = normalize_diet_preference_keys(preference_keys or [])
    return [] if keys == ["balanced"] else keys


def diet_preference_labels(preference_keys: list[str] | None):
    keys = normalize_diet_preference_keys(preference_keys or [])
    if not keys:
        return []
    if keys == ["balanced"]:
        return [DIET_PREFERENCE_OPTIONS["balanced"]]
    return [DIET_PREFERENCE_OPTIONS[key] for key in keys]


def measurement_summary_text(payload: dict):
    labels = {
        "weight_kg": ("Weight", "kg"),
        "arm_cm": ("Arm", "cm"),
        "chest_cm": ("Chest", "cm"),
        "waist_cm": ("Waist", "cm"),
        "thigh_cm": ("Thigh", "cm"),
        "calf_cm": ("Calf", "cm"),
    }
    parts = []
    for key, (label, unit) in labels.items():
        value = safe_float(payload.get(key))
        if value is None:
            continue
        parts.append(f"{label} {value:.1f} {unit}")
    return ", ".join(parts)


def list_plan_overrides_for_user(
    connection: sqlite3.Connection,
    user_id: int,
    domain: str | None = None,
):
    conditions = ["user_id = ?"]
    parameters: list[object] = [user_id]
    if domain:
        conditions.append("domain = ?")
        parameters.append(domain)
    rows = connection.execute(
        f"""
        SELECT *
        FROM plan_overrides
        WHERE {' AND '.join(conditions)}
        ORDER BY created_at ASC, id ASC
        """,
        parameters,
    ).fetchall()
    payloads = []
    for row in rows:
        try:
            payload = json.loads(row["payload_json"] or "{}")
        except json.JSONDecodeError:
            payload = {}
        payloads.append(
            {
                **dict(row),
                "payload": payload,
            }
        )
    return payloads


def append_plan_override(
    connection: sqlite3.Connection,
    user_id: int,
    domain: str,
    scope: str,
    override_type: str,
    payload: dict,
    *,
    effective_date: str | None = None,
    source: str = "assistant",
):
    connection.execute(
        """
        INSERT INTO plan_overrides (
            user_id, domain, scope, override_type, effective_date, payload_json, source, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            user_id,
            domain,
            scope,
            override_type,
            effective_date,
            json.dumps(payload, sort_keys=True),
            source,
            now_iso(),
        ),
    )


def collect_diet_override_state_for_user(connection: sqlite3.Connection, user_id: int, plan_date: str):
    rows = list_plan_overrides_for_user(connection, user_id, domain="diet")
    long_term_preferences = []
    today_preferences = []
    today_meals_per_day = None
    long_term_meal_events = []
    today_meal_events = []
    notes = []
    for row in rows:
        scope = normalize_coach_scope(row.get("scope"))
        override_type = row.get("override_type")
        payload = row.get("payload") or {}
        if scope == "long_term":
            if override_type == "set_diet_preferences":
                long_term_preferences = normalize_diet_preference_keys(payload.get("preference_keys"))
            elif override_type == "customize_diet_meal":
                slot_key = normalize_diet_slot_key(payload.get("slot_key"))
                mode = normalize_diet_meal_customization_mode(payload.get("mode")) or "append"
                items = normalize_diet_meal_override_items(payload)
                if slot_key and items:
                    long_term_meal_events.append(
                        {
                            "slot_key": slot_key,
                            "mode": mode,
                            "items": items,
                        }
                    )
        elif scope == "today" and (row.get("effective_date") or "") == plan_date:
            if override_type == "set_diet_preferences":
                today_preferences = normalize_diet_preference_keys(payload.get("preference_keys"))
            elif override_type == "update_meals_per_day":
                today_meals_per_day = normalize_meals_per_day(payload.get("meals_per_day"), DEFAULT_DIET_MEALS_PER_DAY)
            elif override_type == "customize_diet_meal":
                slot_key = normalize_diet_slot_key(payload.get("slot_key"))
                mode = normalize_diet_meal_customization_mode(payload.get("mode")) or "append"
                items = normalize_diet_meal_override_items(payload)
                if slot_key and items:
                    today_meal_events.append(
                        {
                            "slot_key": slot_key,
                            "mode": mode,
                            "items": items,
                        }
                    )

    active_raw_preferences = today_preferences if today_preferences else long_term_preferences
    active_preferences = effective_diet_preference_keys(active_raw_preferences)
    if long_term_preferences:
        notes.append(f"Regular food preference: {', '.join(diet_preference_labels(long_term_preferences))}.")
    if today_preferences:
        notes.append(f"Today only: {', '.join(diet_preference_labels(today_preferences))}.")
    if today_meals_per_day is not None:
        notes.append(f"Today only: {today_meals_per_day} meals.")
    if long_term_meal_events:
        notes.extend(
            (
                f"Regular {diet_slot_label(item['slot_key']).lower()} is replaced with "
                f"{', '.join(str(entry.get('food_name')).lower() for entry in item.get('items', []) if entry.get('food_name'))}."
                if item.get("mode") == "replace"
                else f"Regular {diet_slot_label(item['slot_key']).lower()} includes "
                f"{', '.join(str(entry.get('food_name')).lower() for entry in item.get('items', []) if entry.get('food_name'))}."
            )
            for item in long_term_meal_events
        )
    if today_meal_events:
        notes.extend(
            (
                f"Today only: replace {diet_slot_label(item['slot_key']).lower()} with "
                f"{', '.join(str(entry.get('food_name')).lower() for entry in item.get('items', []) if entry.get('food_name'))}."
                if item.get("mode") == "replace"
                else f"Today only: add {', '.join(str(entry.get('food_name')).lower() for entry in item.get('items', []) if entry.get('food_name'))} to {diet_slot_label(item['slot_key']).lower()}."
            )
            for item in today_meal_events
        )
    return {
        "active_preference_keys": active_preferences,
        "active_preference_labels": diet_preference_labels(active_raw_preferences),
        "today_meals_per_day": today_meals_per_day,
        "meal_events": [*long_term_meal_events, *today_meal_events],
        "notes": notes,
    }


def collect_gym_override_state_for_user(connection: sqlite3.Connection, user_id: int, plan_date: str):
    rows = list_plan_overrides_for_user(connection, user_id, domain="gym")
    long_term_split = None
    today_split = None
    today_session_minutes = None
    long_term_events = []
    today_events = []
    notes = []
    for row in rows:
        scope = normalize_coach_scope(row.get("scope"))
        override_type = row.get("override_type")
        payload = row.get("payload") or {}
        if scope == "today" and (row.get("effective_date") or "") != plan_date:
            continue
        if override_type == "set_gym_split":
            split_key = normalize_split_key(payload.get("split_key"))
            if split_key:
                if scope == "today":
                    today_split = split_key
                elif scope == "long_term":
                    long_term_split = split_key
        elif override_type == "update_session_minutes":
            session_minutes = safe_int(payload.get("preferred_session_minutes"), DEFAULT_GYM_PREFERRED_MINUTES)
            if scope == "today" and session_minutes:
                today_session_minutes = max(20, min(180, session_minutes))
        elif override_type in {"replace_gym_exercise", "update_gym_exercise_scheme", "update_gym_exercise_load"}:
            if scope == "today":
                today_events.append(row)
            elif scope == "long_term":
                long_term_events.append(row)
    if long_term_split:
        notes.append(f"Regular split override: {split_label(long_term_split)}.")
    if today_split:
        notes.append(f"Today only split: {split_label(today_split)}.")
    if today_session_minutes:
        notes.append(f"Today only session length: {today_session_minutes} min.")
    if long_term_events:
        notes.append(f"Regular exercise edits: {len(long_term_events)} active change{'s' if len(long_term_events) != 1 else ''}.")
    if today_events:
        notes.append(f"Today-only exercise edits: {len(today_events)} active change{'s' if len(today_events) != 1 else ''}.")
    return {
        "long_term_split": long_term_split,
        "today_split": today_split,
        "today_session_minutes": today_session_minutes,
        "long_term_events": long_term_events,
        "today_events": today_events,
        "force_today_session": bool(today_split or today_session_minutes or today_events),
        "notes": notes,
    }


def load_assistant_messages(username: str, limit: int = 20, include_attachment_text: bool = False):
    user_id = ensure_app_user(username)
    prune_assistant_attachments(username)
    with gym_user_connection() as connection:
        messages = connection.execute(
            """
            SELECT *
            FROM assistant_messages
            WHERE user_id = ?
            ORDER BY id DESC
            LIMIT ?
            """,
            (user_id, limit),
        ).fetchall()
    ordered_messages = list(reversed(messages))
    message_ids = [row["id"] for row in ordered_messages]
    actions_by_message = {message_id: [] for message_id in message_ids}
    attachments_by_message = load_assistant_attachments_for_messages(
        username,
        message_ids,
        include_analysis_text=include_attachment_text,
    )
    if message_ids:
        placeholders = ", ".join("?" for _ in message_ids)
        with gym_user_connection() as connection:
            action_rows = connection.execute(
                f"""
                SELECT *
                FROM assistant_actions
                WHERE user_id = ? AND message_id IN ({placeholders})
                ORDER BY id ASC
                """,
                (user_id, *message_ids),
            ).fetchall()
        for row in action_rows:
            try:
                payload = json.loads(row["payload_json"] or "{}")
            except json.JSONDecodeError:
                payload = {}
            actions_by_message.setdefault(row["message_id"], []).append(
                {
                    "id": row["id"],
                    "type": row["action_type"],
                    "label": ASSISTANT_ALLOWED_ACTIONS.get(row["action_type"], row["action_type"]),
                    "summary": row["summary_text"],
                    "payload": payload,
                    "status": row["status"],
                    "resolution_note": row["resolution_note"] or "",
                    "created_at_text": display_timestamp(row["created_at"]),
                    "updated_at_text": display_timestamp(row["updated_at"]),
                }
            )
    payloads = []
    for row in ordered_messages:
        try:
            metadata = json.loads(row["metadata_json"] or "{}")
        except json.JSONDecodeError:
            metadata = {}
        payloads.append(
            {
                "id": row["id"],
                "role": row["role"],
                "content_text": row["content_text"],
                "metadata": metadata,
                "created_at_text": display_timestamp(row["created_at"]),
                "actions": actions_by_message.get(row["id"], []),
                "attachments": attachments_by_message.get(row["id"], []),
            }
        )
    return payloads


def save_gym_profile(username: str, form):
    user_id = ensure_app_user(username)
    goal = form.get("goal", "recomp").strip()
    if goal not in GYM_PROFILE_GOALS:
        goal = "recomp"
    first_name = form.get("first_name", "").strip()
    last_name = form.get("last_name", "").strip()
    height_cm = safe_float(form.get("height_cm"))
    age_years = normalize_age_years(form.get("age_years"))
    current_weight_kg = safe_float(form.get("current_weight_kg"))
    gender = normalize_gender(form.get("gender"))
    preferred_minutes = safe_int(form.get("preferred_session_minutes"), DEFAULT_GYM_PREFERRED_MINUTES) or DEFAULT_GYM_PREFERRED_MINUTES
    training_days = safe_int(form.get("training_days_per_week"), DEFAULT_GYM_DAYS_PER_WEEK) or DEFAULT_GYM_DAYS_PER_WEEK
    preferred_meals_per_day = normalize_meals_per_day(form.get("preferred_meals_per_day"), DEFAULT_DIET_MEALS_PER_DAY)
    notes = form.get("notes", "").strip()
    desired_weight_kg = suggested_weight_targets(
        {
            "height_cm": height_cm,
            "age_years": age_years,
            "current_weight_kg": current_weight_kg,
            "gender": gender,
            "goal": goal,
        }
    ).get("recommended_weight_kg")
    with gym_user_connection() as connection:
        user_row = connection.execute("SELECT username, first_name, last_name FROM users WHERE id = ?", (user_id,)).fetchone()
        resolved_first_name = first_name or (user_row["first_name"] or "")
        resolved_last_name = last_name or (user_row["last_name"] or "")
        display_name = compose_display_name(
            resolved_first_name,
            resolved_last_name,
            user_row["username"],
        )
        connection.execute(
            """
            UPDATE users
            SET first_name = ?, last_name = ?, display_name = ?, updated_at = ?
            WHERE id = ?
            """,
            (resolved_first_name, resolved_last_name, display_name, now_iso(), user_id),
        )
        connection.execute(
            """
            UPDATE gym_profiles
            SET height_cm = ?, age_years = ?, current_weight_kg = ?, desired_weight_kg = ?, gender = ?, goal = ?,
                preferred_session_minutes = ?, training_days_per_week = ?, preferred_meals_per_day = ?, notes = ?, updated_at = ?
            WHERE user_id = ?
            """,
            (height_cm, age_years, current_weight_kg, desired_weight_kg, gender, goal, preferred_minutes, training_days, preferred_meals_per_day, notes, now_iso(), user_id),
        )


def insert_body_measurement(
    connection: sqlite3.Connection,
    user_id: int,
    measured_on: str,
    weight_kg: float | None = None,
    arm_cm: float | None = None,
    chest_cm: float | None = None,
    waist_cm: float | None = None,
    thigh_cm: float | None = None,
    calf_cm: float | None = None,
    note: str = "",
    source: str = "manual",
):
    connection.execute(
        """
        INSERT INTO body_measurements (
            user_id, measured_on, weight_kg, arm_cm, chest_cm, waist_cm,
            thigh_cm, calf_cm, note, source, created_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (user_id, measured_on, weight_kg, arm_cm, chest_cm, waist_cm, thigh_cm, calf_cm, note, source, now_iso()),
    )


def save_health_measurement(username: str, form):
    user_id = ensure_app_user(username)
    measured_on = form.get("measured_on", "").strip() or datetime.now().date().isoformat()
    weight_kg = safe_float(form.get("weight_kg"))
    arm_cm = safe_float(form.get("arm_cm"))
    chest_cm = safe_float(form.get("chest_cm"))
    waist_cm = safe_float(form.get("waist_cm"))
    thigh_cm = safe_float(form.get("thigh_cm"))
    calf_cm = safe_float(form.get("calf_cm"))
    note = form.get("note", "").strip()
    if all(value is None for value in (weight_kg, arm_cm, chest_cm, waist_cm, thigh_cm, calf_cm)) and not note:
        raise ValueError("Enter at least one health measurement before saving.")
    with gym_user_connection() as connection:
        insert_body_measurement(
            connection,
            user_id,
            measured_on,
            weight_kg=weight_kg,
            arm_cm=arm_cm,
            chest_cm=chest_cm,
            waist_cm=waist_cm,
            thigh_cm=thigh_cm,
            calf_cm=calf_cm,
            note=note,
            source="manual",
        )
        if weight_kg is not None:
            sync_profile_current_weight_from_measurements(connection, user_id)


def save_current_exercise_weight(username: str, form):
    user_id = ensure_app_user(username)
    templates = {item["split_key"]: item for item in list_gym_templates()}
    split_key = normalize_split_key(form.get("split_key", ""))
    if split_key not in templates:
        split_key = next(iter(templates.keys()), "pull")
    action_intent = normalize_gym_exercise_status(form.get("intent", "done"))
    exercise_name = form.get("exercise_name", "").strip()
    if not exercise_name:
        raise ValueError("Exercise name is required.")
    performed_on = form.get("performed_on", "").strip() or datetime.now().date().isoformat()
    selected_day = normalize_gym_day(form.get("gym_day"))
    today_iso = datetime.now().date().isoformat()
    weight_kg = safe_float(form.get("weight_kg"))
    if weight_kg is not None and weight_kg < 0:
        raise ValueError("Enter a numeric weight in kg before saving.")
    profile = load_gym_profile(username)
    duration_minutes = profile.get("preferred_session_minutes") or DEFAULT_GYM_PREFERRED_MINUTES
    timestamp = now_iso()
    with gym_user_connection() as connection:
        session_id = ensure_gym_session(
            connection,
            user_id,
            performed_on,
            split_key,
            duration_minutes,
            profile.get("current_weight_kg"),
        )
        current_row = connection.execute(
            """
            SELECT id, weight_kg, weight_text, status, completed
            FROM gym_session_exercises
            WHERE session_id = ? AND exercise_name = ?
            ORDER BY id DESC LIMIT 1
            """,
            (session_id, exercise_name),
        ).fetchone()
        current_status = normalize_gym_exercise_status(
            current_row["status"] if current_row and current_row["status"] is not None else ("done" if current_row and current_row["completed"] else "pending")
        )
        next_status = "pending" if current_status == action_intent else action_intent
        stored_weight_kg = weight_kg
        if stored_weight_kg is None and current_row:
            stored_weight_kg = current_row["weight_kg"] if current_row["weight_kg"] is not None else parse_weight_kg_from_text(current_row["weight_text"])
        if next_status == "done" and stored_weight_kg is None:
            raise ValueError("Enter a numeric weight in kg before marking an exercise done.")
        stored_weight_text = format_weight_kg(stored_weight_kg) if stored_weight_kg is not None else ""
        completed_flag = 1 if next_status == "done" else 0
        if current_row:
            connection.execute(
                """
                UPDATE gym_session_exercises
                SET weight_kg = ?, weight_text = ?, completed = ?, status = ?, updated_at = ?
                WHERE id = ?
                """,
                (stored_weight_kg, stored_weight_text, completed_flag, next_status, timestamp, current_row["id"]),
            )
        else:
            connection.execute(
                """
                INSERT INTO gym_session_exercises (
                    session_id, exercise_name, weight_kg, weight_text, completed, status, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (session_id, exercise_name, stored_weight_kg, stored_weight_text, completed_flag, next_status, timestamp),
            )
        expected_exercise_names = planned_gym_exercise_names_for_session(
            connection,
            user_id,
            split_key,
            today_iso if selected_day == "today" else performed_on,
            templates,
        )
        session_rows = connection.execute(
            """
            SELECT exercise_name, weight_kg, weight_text, status, completed
            FROM gym_session_exercises
            WHERE session_id = ?
            ORDER BY id ASC
            """,
            (session_id,),
        ).fetchall()
        session_entries = {}
        for row in session_rows:
            session_entries[row["exercise_name"]] = {
                "weight_kg": row["weight_kg"],
                "weight_text": row["weight_text"],
                "status": normalize_gym_exercise_status(
                    row["status"] if row["status"] is not None else ("done" if row["completed"] else "pending")
                ),
            }
        progress = gym_session_progress(expected_exercise_names, session_entries)
        final_performed_on = today_iso if progress["is_complete"] else performed_on
        connection.execute(
            """
            UPDATE gym_sessions
            SET performed_on = ?, duration_minutes = ?, body_weight_kg = ?
            WHERE id = ?
            """,
            (final_performed_on, duration_minutes, profile.get("current_weight_kg"), session_id),
        )
        if next_status == "done" and stored_weight_kg is not None:
            connection.execute(
                """
                INSERT INTO user_exercise_weights (user_id, exercise_name, weight_kg, weight_text, updated_at)
                VALUES (?, ?, ?, ?, ?)
                ON CONFLICT(user_id, exercise_name) DO UPDATE
                SET weight_kg = excluded.weight_kg, weight_text = excluded.weight_text, updated_at = excluded.updated_at
                """,
                (user_id, exercise_name, stored_weight_kg, stored_weight_text, timestamp),
            )
    message_suffix = f" Session saved under {final_performed_on}." if progress["is_complete"] else ""
    if next_status == "done":
        return {"status": "done", "message": f"Marked {exercise_name} done.{message_suffix}", "performed_on": final_performed_on}
    if next_status == "skipped":
        return {"status": "skipped", "message": f"Skipped {exercise_name} for this session.{message_suffix}", "performed_on": final_performed_on}
    return {"status": "pending", "message": f"Returned {exercise_name} to pending.", "performed_on": final_performed_on}


def bmi_for_profile(profile: dict):
    height_cm = profile.get("height_cm")
    weight_kg = profile.get("current_weight_kg")
    if not height_cm or not weight_kg:
        return None
    height_m = height_cm / 100.0
    if height_m <= 0:
        return None
    return round(weight_kg / (height_m * height_m), 1)


def suggested_weight_targets(profile: dict):
    height_cm = profile.get("height_cm")
    current_weight = profile.get("current_weight_kg")
    gender_config = gender_profile_config(profile.get("gender"))
    if not height_cm:
        return {
            "range_text": "Add height to get a suggested weight range.",
            "recommended_weight_kg": None,
            "note": "Height is still missing. Add gender too if you want a more personalized target.",
        }
    height_m = height_cm / 100.0
    healthy_min = round(gender_config["target_bmi_min"] * height_m * height_m, 1)
    healthy_max = round(gender_config["target_bmi_max"] * height_m * height_m, 1)
    goal = profile.get("goal", "recomp")
    if current_weight is None:
        recommended = round((healthy_min + healthy_max) / 2.0, 1)
    elif goal == "fat_loss":
        recommended = max(healthy_min, round(min(current_weight - 4.0, healthy_max), 1))
    elif goal == "muscle_gain":
        recommended = round(min(max(current_weight + 2.0, healthy_min), healthy_max + 4.0), 1)
    elif goal == "strength":
        recommended = round(min(max(current_weight + 1.0, healthy_min), healthy_max + 2.0), 1)
    else:
        recommended = round(min(max(current_weight, healthy_min), healthy_max), 1)
    if gender_config["key"] == "unspecified":
        note = (
            f"Suggested using height and your {GYM_PROFILE_GOALS.get(goal, goal).lower()} goal. "
            "Add gender in Health to personalize the reference range further."
        )
    else:
        note = (
            f"Suggested using height, {gender_config['label'].lower()} reference ranges, "
            f"and your {GYM_PROFILE_GOALS.get(goal, goal).lower()} goal."
        )
    return {
        "range_text": f"{healthy_min:.1f}-{healthy_max:.1f} kg reference range",
        "recommended_weight_kg": recommended,
        "note": note,
    }


def list_diet_meal_templates():
    ensure_gym_databases()
    with gym_knowledge_connection() as connection:
        rows = connection.execute(
            """
            SELECT template_key, slot_key, title, goal_bias, summary, prep_text,
                   base_calories, protein_g, carbs_g, fat_g, ingredients_json, sort_order
            FROM diet_meal_templates
            ORDER BY slot_key, sort_order, template_key
            """
        ).fetchall()
        nutrient_rows = connection.execute(
            """
            SELECT template_key, nutrient_code, amount
            FROM diet_meal_template_nutrients
            ORDER BY template_key, nutrient_code
            """
        ).fetchall()
    nutrients_by_template: dict[str, dict[str, float]] = {}
    for row in nutrient_rows:
        nutrients_by_template.setdefault(row["template_key"], {})[row["nutrient_code"]] = round(float(row["amount"] or 0), 1)
    templates = []
    for row in rows:
        try:
            ingredients = json.loads(row["ingredients_json"] or "[]")
        except json.JSONDecodeError:
            ingredients = []
        nutrient_map = nutrients_by_template.get(row["template_key"], {})
        templates.append(
            {
                **dict(row),
                "ingredients": ingredients,
                "nutrients": nutrient_map,
            }
        )
    return templates


def list_diet_nutrient_definitions():
    ensure_gym_databases()
    with gym_knowledge_connection() as connection:
        rows = connection.execute(
            """
            SELECT nutrient_code, label, unit, category, sort_order, description
            FROM nutrient_definitions
            ORDER BY sort_order, nutrient_code
            """
        ).fetchall()
    return [
        {
            "code": row["nutrient_code"],
            "label": row["label"],
            "unit": row["unit"],
            "category": row["category"],
            "sort_order": row["sort_order"],
            "description": row["description"] or "",
        }
        for row in rows
    ]


def normalize_diet_food_text(value: str | None):
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower()).strip()


def diet_food_key_from_name(value: str | None):
    slug = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")
    return slug or "food"


def estimate_diet_food_calories(nutrients: dict | None):
    nutrient_map = nutrients or {}
    protein = float(nutrient_map.get("protein_g") or 0.0)
    carbs = float(nutrient_map.get("carbs_g") or 0.0)
    fat = float(nutrient_map.get("fat_g") or 0.0)
    if protein <= 0 and carbs <= 0 and fat <= 0:
        return None
    return round(protein * 4.0 + carbs * 4.0 + fat * 9.0, 1)


def normalize_diet_food_nutrients(value):
    source = value if isinstance(value, dict) else {}
    normalized = {}
    for nutrient_code in DIET_TRACKED_NUTRIENT_CODES:
        amount = safe_float(source.get(nutrient_code))
        if amount is None or amount < 0:
            continue
        normalized[nutrient_code] = round(float(amount), 1)
    return normalized


def _hydrate_diet_food_library_item(connection: sqlite3.Connection, food_row: sqlite3.Row | None):
    if not food_row:
        return None
    nutrient_rows = connection.execute(
        """
        SELECT nutrient_code, amount
        FROM diet_food_library_nutrients
        WHERE food_key = ?
        ORDER BY nutrient_code
        """,
        (food_row["food_key"],),
    ).fetchall()
    alias_rows = connection.execute(
        """
        SELECT alias_normalized
        FROM diet_food_library_aliases
        WHERE food_key = ?
        ORDER BY alias_normalized
        """,
        (food_row["food_key"],),
    ).fetchall()
    nutrients = {
        row["nutrient_code"]: round(float(row["amount"] or 0.0), 1)
        for row in nutrient_rows
    }
    return {
        "key": food_row["food_key"],
        "label": food_row["food_name"],
        "serving_text": food_row["serving_text"],
        "calories": round(float(food_row["calories"] or 0.0), 1),
        "nutrients": nutrients,
        "aliases": [row["alias_normalized"] for row in alias_rows],
        "source_label": food_row["source_label"] or "",
        "source_kind": food_row["source_kind"] or "",
        "normalized_name": food_row["normalized_name"],
    }


def _resolve_diet_custom_food_in_connection(connection: sqlite3.Connection, food_name: str | None):
    normalized = normalize_diet_food_text(food_name)
    if not normalized:
        return None
    row = connection.execute(
        """
        SELECT library.*
        FROM diet_food_library_aliases AS aliases
        JOIN diet_food_library AS library ON library.food_key = aliases.food_key
        WHERE aliases.alias_normalized = ?
        LIMIT 1
        """,
        (normalized,),
    ).fetchone()
    if not row:
        row = connection.execute(
            """
            SELECT *
            FROM diet_food_library
            WHERE normalized_name = ?
            LIMIT 1
            """,
            (normalized,),
        ).fetchone()
    if not row:
        alias_rows = connection.execute(
            """
            SELECT aliases.alias_normalized, library.*
            FROM diet_food_library_aliases AS aliases
            JOIN diet_food_library AS library ON library.food_key = aliases.food_key
            ORDER BY LENGTH(aliases.alias_normalized) DESC, aliases.alias_normalized
            """
        ).fetchall()
        normalized_words = {word for word in normalized.split() if word}
        for alias_row in alias_rows:
            alias = alias_row["alias_normalized"]
            alias_words = {word for word in str(alias or "").split() if word}
            if not alias:
                continue
            if alias_words and alias_words == normalized_words:
                row = alias_row
                break
            if len(alias_words) >= 2 and len(normalized_words) >= 2 and (alias in normalized or normalized in alias):
                row = alias_row
                break
    return _hydrate_diet_food_library_item(connection, row)


def _next_diet_food_key(connection: sqlite3.Connection, base_key: str):
    candidate = base_key or "food"
    suffix = 2
    while connection.execute(
        "SELECT 1 FROM diet_food_library WHERE food_key = ? LIMIT 1",
        (candidate,),
    ).fetchone():
        candidate = f"{base_key}_{suffix}"
        suffix += 1
    return candidate


def _normalize_diet_food_library_item(item: dict | None):
    if not isinstance(item, dict):
        return None
    food_name = truncate_text(item.get("food_name") or item.get("label"), 120)
    if not food_name:
        return None
    serving_text = truncate_text(item.get("serving_text"), 80) or f"1 serving {food_name}"
    nutrients = normalize_diet_food_nutrients(item.get("nutrients") or item)
    calories = safe_float(item.get("calories"))
    if calories is None:
        calories = estimate_diet_food_calories(nutrients)
    if calories is None:
        return None
    aliases = []
    raw_aliases = item.get("aliases") if isinstance(item.get("aliases"), (list, tuple, set)) else []
    for alias in [food_name, *raw_aliases]:
        normalized_alias = normalize_diet_food_text(alias)
        if normalized_alias and normalized_alias not in aliases:
            aliases.append(normalized_alias)
    return {
        "food_name": food_name,
        "normalized_name": normalize_diet_food_text(food_name),
        "serving_text": serving_text,
        "calories": round(float(calories), 1),
        "nutrients": nutrients,
        "aliases": aliases,
    }


def upsert_diet_food_library_item(
    item: dict,
    *,
    connection: sqlite3.Connection | None = None,
    source_label: str = "Custom food library",
    source_kind: str = "assistant",
):
    normalized_item = _normalize_diet_food_library_item(item)
    if not normalized_item:
        raise ValueError("Diet food item is incomplete.")
    owns_connection = connection is None
    if owns_connection:
        ensure_gym_databases()
        connection = gym_knowledge_connection()
    assert connection is not None
    try:
        existing = _resolve_diet_custom_food_in_connection(connection, normalized_item["food_name"])
        timestamp = now_iso()
        food_key = existing["key"] if existing else _next_diet_food_key(connection, diet_food_key_from_name(normalized_item["food_name"]))
        connection.execute(
            """
            INSERT INTO diet_food_library (
                food_key, normalized_name, food_name, serving_text, calories,
                source_label, source_kind, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(food_key) DO UPDATE SET
                normalized_name = excluded.normalized_name,
                food_name = excluded.food_name,
                serving_text = excluded.serving_text,
                calories = excluded.calories,
                source_label = excluded.source_label,
                source_kind = excluded.source_kind,
                updated_at = excluded.updated_at
            """,
            (
                food_key,
                normalized_item["normalized_name"],
                normalized_item["food_name"],
                normalized_item["serving_text"],
                normalized_item["calories"],
                truncate_text(source_label, 80),
                truncate_text(source_kind, 32) or "assistant",
                timestamp,
                timestamp,
            ),
        )
        connection.execute(
            "DELETE FROM diet_food_library_nutrients WHERE food_key = ?",
            (food_key,),
        )
        connection.executemany(
            """
            INSERT INTO diet_food_library_nutrients (food_key, nutrient_code, amount)
            VALUES (?, ?, ?)
            """,
            [
                (food_key, nutrient_code, amount)
                for nutrient_code, amount in normalized_item["nutrients"].items()
            ],
        )
        connection.executemany(
            """
            INSERT INTO diet_food_library_aliases (alias_normalized, food_key)
            VALUES (?, ?)
            ON CONFLICT(alias_normalized) DO UPDATE SET food_key = excluded.food_key
            """,
            [
                (alias, food_key)
                for alias in normalized_item["aliases"]
            ],
        )
        resolved = _resolve_diet_custom_food_in_connection(connection, normalized_item["food_name"])
        if owns_connection:
            connection.commit()
        return resolved
    finally:
        if owns_connection:
            connection.close()


def list_diet_food_library_items(limit: int | None = None):
    ensure_gym_databases()
    query = """
        SELECT *
        FROM diet_food_library
        ORDER BY food_name
    """
    params: tuple = ()
    if limit is not None:
        query = f"{query} LIMIT ?"
        params = (max(1, int(limit)),)
    with gym_knowledge_connection() as connection:
        rows = connection.execute(query, params).fetchall()
        return [_hydrate_diet_food_library_item(connection, row) for row in rows]


def list_diet_messages():
    ensure_gym_databases()
    with gym_knowledge_connection() as connection:
        rows = connection.execute(
            "SELECT message_text FROM diet_messages ORDER BY sort_order, id"
        ).fetchall()
    return [row["message_text"] for row in rows]


def diet_slot_group(slot_key: str) -> str:
    return "snack" if slot_key.startswith("snack") else slot_key


def normalize_diet_slot_key(value: str | None) -> str | None:
    slug = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")
    aliases = {
        "breakfast": "breakfast",
        "morning": "breakfast",
        "lunch": "lunch",
        "dinner": "dinner",
        "snack": "snack",
        "snack_1": "snack",
        "snack_2": "snack",
        "snack_am": "snack",
        "snack_pm": "snack",
    }
    return aliases.get(slug)


def diet_slot_label(slot_key: str | None) -> str:
    normalized = normalize_diet_slot_key(slot_key)
    labels = {
        "breakfast": "Breakfast",
        "lunch": "Lunch",
        "dinner": "Dinner",
        "snack": "Snack",
    }
    return labels.get(normalized or "", (normalized or "Meal").replace("_", " ").title())


def resolve_diet_custom_food(food_name: str | None):
    ensure_gym_databases()
    with gym_knowledge_connection() as connection:
        return _resolve_diet_custom_food_in_connection(connection, food_name)


def normalize_diet_meal_override_items(payload: dict | None):
    if not isinstance(payload, dict):
        return []
    raw_items = payload.get("items") if isinstance(payload.get("items"), list) else None
    if raw_items:
        candidates = raw_items
    else:
        candidates = [payload]
    normalized_items = []
    for raw_item in candidates[:8]:
        if not isinstance(raw_item, dict):
            continue
        food_name = truncate_text(raw_item.get("food_name"), 120)
        if not food_name:
            continue
        serving_count = min(8.0, max(0.25, safe_float(raw_item.get("servings")) or 1.0))
        serving_text = truncate_text(raw_item.get("serving_text"), 80)
        nutrients = normalize_diet_food_nutrients(raw_item.get("nutrients") or raw_item)
        calories = safe_float(raw_item.get("calories"))
        if calories is None:
            calories = estimate_diet_food_calories(nutrients)
        normalized_items.append(
            {
                "food_name": food_name,
                "servings": serving_count,
                "serving_text": serving_text,
                "calories": calories,
                "nutrients": nutrients,
            }
        )
    return normalized_items


def pluralize_diet_food_name(food_name: str, servings: float):
    current_name = str(food_name or "").strip()
    if abs(float(servings or 1.0) - 1.0) < 0.01:
        return current_name
    lower_name = current_name.lower()
    if lower_name.endswith("white"):
        return f"{current_name}s"
    if lower_name.endswith("y") and not lower_name.endswith(("ay", "ey", "iy", "oy", "uy")):
        return f"{current_name[:-1]}ies"
    if lower_name.endswith("s"):
        return current_name
    return f"{current_name}s"


def diet_override_items_summary_text(items: list[dict] | None):
    parts = []
    for item in items or []:
        food_name = str(item.get("food_name") or "").strip()
        if not food_name:
            continue
        servings = float(item.get("servings") or 1.0)
        servings_text = format_scaled_amount(servings, "")
        parts.append(f"{servings_text} {pluralize_diet_food_name(food_name, servings)}")
    return ", ".join(parts)


def build_diet_custom_food_item(item_payload: dict | None, *, connection: sqlite3.Connection | None = None):
    if not isinstance(item_payload, dict):
        return None
    food_name = truncate_text(item_payload.get("food_name"), 120)
    if not food_name:
        return None
    food = _resolve_diet_custom_food_in_connection(connection, food_name) if connection is not None else resolve_diet_custom_food(food_name)
    explicit_nutrients = normalize_diet_food_nutrients(item_payload.get("nutrients") or item_payload)
    explicit_calories = safe_float(item_payload.get("calories"))
    serving_count = min(8.0, max(0.25, safe_float(item_payload.get("servings")) or 1.0))
    if abs(serving_count - round(serving_count)) < 0.01:
        serving_count = float(int(round(serving_count)))
    if food:
        base_label = str(food.get("label") or food_name).strip()
        serving_text = truncate_text(item_payload.get("serving_text"), 80) or str(food.get("serving_text") or base_label).strip()
        base_nutrients = dict(food.get("nutrients") or {})
        if explicit_nutrients:
            base_nutrients.update(explicit_nutrients)
        per_serving_nutrients = base_nutrients
        per_serving_calories = explicit_calories if explicit_calories is not None else safe_float(food.get("calories"))
    else:
        base_label = food_name
        serving_text = truncate_text(item_payload.get("serving_text"), 80) or f"1 serving {base_label}"
        per_serving_nutrients = explicit_nutrients
        per_serving_calories = explicit_calories
    if per_serving_calories is None:
        per_serving_calories = estimate_diet_food_calories(per_serving_nutrients)
    if per_serving_calories is None:
        return None
    if abs(serving_count - 1.0) < 0.01:
        item_text = serving_text
    else:
        count_text = format_scaled_amount(serving_count, "")
        item_text = f"{count_text} {pluralize_diet_food_name(base_label, serving_count)}"
    nutrient_map = {
        nutrient_code: round(float(amount or 0.0) * serving_count, 1)
        for nutrient_code, amount in (per_serving_nutrients or {}).items()
    }
    return {
        "food_key": (food or {}).get("key"),
        "food_name": base_label,
        "serving_text": serving_text,
        "item_text": item_text,
        "servings": serving_count,
        "calories": round(float(per_serving_calories or 0.0) * serving_count, 1),
        "calories_per_serving": round(float(per_serving_calories or 0.0), 1),
        "per_serving_nutrients": {code: round(float(amount or 0.0), 1) for code, amount in (per_serving_nutrients or {}).items()},
        "nutrients": nutrient_map,
    }


def build_diet_custom_meal_shell(meal: dict):
    custom_meal = copy.deepcopy(meal)
    slot_key = normalize_diet_slot_key(custom_meal.get("slot_key"))
    slot_label_text = diet_slot_label(slot_key)
    custom_meal["template_key"] = f"custom_{slot_key or 'meal'}"
    custom_meal["meal_title"] = f"Custom {slot_label_text}"
    custom_meal["meal_summary"] = "Custom meal built around your requested items."
    custom_meal["prep_text"] = "Prepare this meal using your saved custom items."
    custom_meal["portion_multiplier"] = 1.0
    custom_meal["ingredients"] = []
    custom_meal["items"] = []
    custom_meal["calories"] = 0
    custom_meal["protein_g"] = 0
    custom_meal["carbs_g"] = 0
    custom_meal["fat_g"] = 0
    custom_meal["nutrients"] = {nutrient_code: 0.0 for nutrient_code in DIET_TRACKED_NUTRIENT_CODES}
    custom_meal["status"] = "pending"
    custom_meal["is_done"] = False
    custom_meal["is_skipped"] = False
    return custom_meal


def build_diet_food_estimation_prompt(items: list[dict]):
    request_payload = [
        {
            "food_name": str(item.get("food_name") or "").strip(),
            "serving_text": str(item.get("serving_text") or "").strip(),
        }
        for item in items
        if str(item.get("food_name") or "").strip()
    ]
    return (
        "You are estimating simple food nutrition for a private diet tracker.\n"
        "Return JSON only with this shape: "
        "{\"items\":[{\"food_name\":\"...\",\"serving_text\":\"...\",\"calories\":123,"
        "\"nutrients\":{\"protein_g\":0,\"carbs_g\":0,\"fat_g\":0,\"fiber_g\":0}}]}.\n"
        "Rules:\n"
        "- One entry per requested item.\n"
        "- Use practical single-serving estimates.\n"
        "- serving_text must describe one serving, for example '1 egg white' or '1 latte'.\n"
        "- calories and nutrients must be for one serving, not multiple servings.\n"
        "- Keep nutrient values numeric.\n"
        "- Do not add commentary or markdown.\n\n"
        f"Requested items:\n{json.dumps(request_payload, ensure_ascii=True)}"
    )


def estimate_diet_foods_with_agent(items: list[dict]):
    unresolved_items = [
        {
            "food_name": truncate_text(item.get("food_name"), 120),
            "serving_text": truncate_text(item.get("serving_text"), 80),
        }
        for item in items
        if str(item.get("food_name") or "").strip()
    ]
    if not unresolved_items or not agent_host_online():
        return {}
    response_text = remote_codex_exec(build_diet_food_estimation_prompt(unresolved_items), timeout=45)
    match = re.search(r"\{.*\}", str(response_text or "").strip(), flags=re.DOTALL)
    candidate = match.group(0) if match else str(response_text or "").strip()
    payload = json.loads(candidate)
    raw_items = payload.get("items") if isinstance(payload, dict) and isinstance(payload.get("items"), list) else []
    enriched = {}
    for raw_item in raw_items[: len(unresolved_items)]:
        if not isinstance(raw_item, dict):
            continue
        food_name = truncate_text(raw_item.get("food_name"), 120)
        if not food_name:
            continue
        nutrients = normalize_diet_food_nutrients(raw_item.get("nutrients") or raw_item)
        calories = safe_float(raw_item.get("calories"))
        if calories is None:
            calories = estimate_diet_food_calories(nutrients)
        if calories is None:
            continue
        enriched[normalize_diet_food_text(food_name)] = {
            "food_name": food_name,
            "serving_text": truncate_text(raw_item.get("serving_text"), 80),
            "calories": calories,
            "nutrients": nutrients,
        }
    return enriched


def assistant_image_analysis_enabled():
    return bool(OPENAI_API_KEY.strip())


def normalize_openai_vision_detail(value: str | None):
    current = str(value or "").strip().lower()
    return current if current in {"low", "high", "auto", "original"} else "high"


def extract_openai_response_text(payload: dict):
    if isinstance(payload.get("output_text"), str) and payload.get("output_text").strip():
        return payload["output_text"].strip()
    parts = []
    for item in payload.get("output", []):
        if not isinstance(item, dict):
            continue
        for content in item.get("content", []):
            if not isinstance(content, dict):
                continue
            text_value = content.get("text") or content.get("output_text")
            if isinstance(text_value, str) and text_value.strip():
                parts.append(text_value.strip())
    return "\n".join(parts).strip()


def build_diet_image_analysis_prompt(user_message: str, original_name: str):
    clean_message = truncate_text(user_message, 400)
    return (
        "You are analyzing a food photo for a private diet tracker.\n"
        "Return JSON only with keys summary, items, and notes.\n"
        "- summary: short plain-English summary of the meal.\n"
        "- items: array of visible food or drink estimates.\n"
        "- each item must have food_name, servings, serving_text, calories, and nutrients.\n"
        "- nutrients must include protein_g, carbs_g, fat_g, and fiber_g as numeric per-serving values.\n"
        "- servings must describe how many servings are visible in the photo.\n"
        "- serving_text must describe one serving, for example '1 boiled egg', '1 egg white', '150 g cooked rice', or '1 medium latte'.\n"
        "- prefer simple reusable food names that work well in a diet database.\n"
        "- estimate practical portions when the exact amount is unclear, and mention uncertainty in notes.\n"
        "- do not include markdown fences.\n\n"
        f"Original file name: {original_name}\n"
        f"User message context: {clean_message or 'No extra context supplied.'}"
    )


def parse_diet_image_analysis_payload(text: str):
    raw = str(text or "").strip()
    if not raw:
        return None
    match = re.search(r"\{.*\}", raw, flags=re.DOTALL)
    candidate = match.group(0) if match else raw
    try:
        payload = json.loads(candidate)
    except json.JSONDecodeError:
        return None
    if not isinstance(payload, dict):
        return None
    items = normalize_diet_meal_override_items({"items": payload.get("items")})
    summary = truncate_text(payload.get("summary"), 240)
    notes = []
    for note in payload.get("notes") if isinstance(payload.get("notes"), list) else []:
        note_text = truncate_text(note, 160)
        if note_text:
            notes.append(note_text)
    if not summary and items:
        summary = f"Meal photo suggests {', '.join(item['food_name'] for item in items[:3])}."
    if not summary and not items:
        return None
    return {
        "summary": summary or "Meal photo analyzed.",
        "items": items,
        "notes": notes[:4],
    }


def build_diet_image_analysis_text(payload: dict | None):
    if not isinstance(payload, dict):
        return ""
    lines = []
    summary = truncate_text(payload.get("summary"), 240)
    if summary:
        lines.append(f"Meal photo summary: {summary}")
    items = payload.get("items") if isinstance(payload.get("items"), list) else []
    if items:
        lines.append("Detected meal items:")
        for item in items[:6]:
            food_name = truncate_text(item.get("food_name"), 120)
            if not food_name:
                continue
            servings = float(item.get("servings") or 1.0)
            servings_text = format_scaled_amount(servings, "")
            serving_text = truncate_text(item.get("serving_text"), 80)
            calories = safe_float(item.get("calories"))
            nutrient_map = normalize_diet_food_nutrients(item.get("nutrients"))
            macros = []
            for code, label in (
                ("protein_g", "protein"),
                ("carbs_g", "carbs"),
                ("fat_g", "fat"),
                ("fiber_g", "fiber"),
            ):
                amount = safe_float(nutrient_map.get(code))
                if amount is None:
                    continue
                macros.append(f"{label} {amount:.1f} g")
            details = []
            if serving_text:
                details.append(f"per serving: {serving_text}")
            if calories is not None:
                details.append(f"{calories:.0f} kcal")
            if macros:
                details.append(", ".join(macros))
            suffix = f" ({'; '.join(details)})" if details else ""
            lines.append(f"- {servings_text} x {food_name}{suffix}")
    notes = payload.get("notes") if isinstance(payload.get("notes"), list) else []
    if notes:
        lines.append("Notes:")
        for note in notes[:4]:
            note_text = truncate_text(note, 160)
            if note_text:
                lines.append(f"- {note_text}")
    return "\n".join(lines).strip()


def analyze_diet_image_attachment(data: bytes, mime_type: str, user_message: str, original_name: str):
    if not assistant_image_analysis_enabled():
        return None
    mime = str(mime_type or "").strip().lower() or "image/jpeg"
    detail = normalize_openai_vision_detail(OPENAI_VISION_DETAIL)
    request_payload = {
        "model": OPENAI_VISION_MODEL,
        "input": [
            {
                "role": "user",
                "content": [
                    {"type": "input_text", "text": build_diet_image_analysis_prompt(user_message, original_name)},
                    {
                        "type": "input_image",
                        "image_url": f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}",
                        "detail": detail,
                    },
                ],
            }
        ],
    }
    body = json.dumps(request_payload).encode("utf-8")
    headers = {
        "Authorization": f"Bearer {OPENAI_API_KEY}",
        "Content-Type": "application/json",
        "X-Client-Request-Id": f"healthhub-meal-photo-{uuid.uuid4()}",
    }
    request_obj = urllib.request.Request(
        f"{OPENAI_API_BASE_URL.rstrip('/')}/responses",
        data=body,
        headers=headers,
        method="POST",
    )
    try:
        with urllib.request.urlopen(request_obj, timeout=max(15, OPENAI_VISION_TIMEOUT_SECONDS)) as response:
            raw_payload = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        error_body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Vision API request failed with HTTP {exc.code}: {truncate_text(error_body, 240)}") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"Vision API request failed: {exc.reason}") from exc
    response_text = extract_openai_response_text(raw_payload)
    return parse_diet_image_analysis_payload(response_text)


def enrich_assistant_image_attachments(prepared_attachments: list[dict], user_message: str):
    if not prepared_attachments:
        return
    image_items = [item for item in prepared_attachments if item.get("is_image")]
    if not image_items:
        return
    enabled = assistant_image_analysis_enabled()
    max_images = max(1, OPENAI_VISION_MAX_IMAGES_PER_MESSAGE)
    for index, item in enumerate(image_items):
        if not enabled:
            item["analysis_text"] = (
                "Image attached. Automatic meal-photo analysis is not configured on this Health Hub server yet."
            )
            item["preview_text"] = "Image attached. Meal-photo analysis is not configured yet."
            item["analysis_json"] = json.dumps({"summary": "Meal-photo analysis not configured.", "items": [], "notes": []}, sort_keys=True)
            continue
        if index >= max_images:
            item["analysis_text"] = f"Image attached. Only the first {max_images} image(s) in one coach message are analyzed automatically."
            item["preview_text"] = "Image attached. Analysis limit reached for this message."
            item["analysis_json"] = json.dumps({"summary": "Analysis limit reached for this message.", "items": [], "notes": []}, sort_keys=True)
            continue
        try:
            analysis_payload = analyze_diet_image_attachment(
                item.get("data") or b"",
                str(item.get("mime_type") or ""),
                user_message,
                str(item.get("original_name") or "meal-photo"),
            )
        except Exception:
            item["analysis_text"] = "Image attached. Automatic meal-photo analysis was unavailable for this image."
            item["preview_text"] = "Image attached. Meal-photo analysis unavailable right now."
            item["analysis_json"] = json.dumps({"summary": "Meal-photo analysis unavailable.", "items": [], "notes": []}, sort_keys=True)
            continue
        if not analysis_payload:
            item["analysis_text"] = "Image attached. The meal photo could not be understood clearly enough to build food items."
            item["preview_text"] = "Image attached. Meal-photo details were unclear."
            item["analysis_json"] = json.dumps({"summary": "Meal-photo details were unclear.", "items": [], "notes": []}, sort_keys=True)
            continue
        summary = truncate_text(analysis_payload.get("summary"), 180)
        item_names = [truncate_text(entry.get("food_name"), 40) for entry in analysis_payload.get("items", [])[:4] if entry.get("food_name")]
        item["analysis_text"] = build_diet_image_analysis_text(analysis_payload)
        item["analysis_json"] = json.dumps(analysis_payload, sort_keys=True)
        item["preview_text"] = (
            f"Meal photo: {', '.join(item_names)}"
            if item_names
            else summary or "Meal photo analyzed."
        )


def apply_diet_meal_customizations(meal: dict, meal_override_events: list[dict]):
    current_meal = copy.deepcopy(meal)
    applied_notes = []
    replaced_meal = False
    appended_custom_items = False
    for event in meal_override_events:
        mode = normalize_diet_meal_customization_mode(event.get("mode")) or "append"
        event_items = normalize_diet_meal_override_items(event)
        if mode == "replace":
            current_meal = build_diet_custom_meal_shell(current_meal)
            replaced_meal = True
        next_item_order = len(current_meal.get("items", [])) + 1
        with gym_knowledge_connection() as connection:
            for raw_item in event_items:
                item_spec = build_diet_custom_food_item(raw_item, connection=connection)
                if not item_spec:
                    continue
                appended_custom_items = True
                nutrient_map = item_spec.get("nutrients") or {}
                current_meal.setdefault("items", []).append(
                    {
                        "item_order": next_item_order,
                        "item_text": item_spec["item_text"],
                        "status": "pending",
                        "share_ratio": 0.0,
                        "calories": item_spec["calories"],
                        "nutrients": nutrient_map,
                    }
                )
                current_meal.setdefault("ingredients", []).append(item_spec["item_text"])
                current_meal["calories"] = int(round(float(current_meal.get("calories") or 0) + float(item_spec["calories"])))
                current_meal["protein_g"] = int(round(float(current_meal.get("protein_g") or 0) + float(nutrient_map.get("protein_g") or 0.0)))
                current_meal["carbs_g"] = int(round(float(current_meal.get("carbs_g") or 0) + float(nutrient_map.get("carbs_g") or 0.0)))
                current_meal["fat_g"] = int(round(float(current_meal.get("fat_g") or 0) + float(nutrient_map.get("fat_g") or 0.0)))
                meal_nutrients = dict(current_meal.get("nutrients") or {})
                for nutrient_code in DIET_TRACKED_NUTRIENT_CODES:
                    meal_nutrients[nutrient_code] = round(float(meal_nutrients.get(nutrient_code) or 0.0) + float(nutrient_map.get(nutrient_code) or 0.0), 1)
                current_meal["nutrients"] = meal_nutrients
                next_item_order += 1
    if replaced_meal:
        current_meal["meal_summary"] = "Custom meal built around your requested items."
        current_meal["prep_text"] = "Prepare this meal using your saved custom items."
    elif appended_custom_items:
        current_meal["meal_summary"] = "Updated with your requested items."
    return current_meal, applied_notes


def normalize_diet_meal_summary_text(template_key: str | None, meal_summary: str | None):
    template_text = str(template_key or "").strip().lower()
    summary_text = str(meal_summary or "").strip()
    if template_text.startswith("custom_"):
        return "Custom meal built around your requested items."
    if "Includes " in summary_text or "Rebuilt as a custom meal" in summary_text:
        return "Updated with your requested items."
    return summary_text


def diet_meal_time_window(slot_key: str, meal_count: int | None = None) -> str:
    normalized_slot = str(slot_key or "").strip().lower()
    count_key = normalize_meals_per_day(meal_count, DEFAULT_DIET_MEALS_PER_DAY) if meal_count is not None else None
    if count_key in DIET_MEAL_TIME_WINDOWS:
        text = DIET_MEAL_TIME_WINDOWS[count_key].get(normalized_slot)
        if text:
            return text
    slot_group = diet_slot_group(normalized_slot)
    return DIET_MEAL_TIME_WINDOWS["default"].get(normalized_slot) or DIET_MEAL_TIME_WINDOWS["default"].get(slot_group) or ""


def ingredient_share_score(ingredient_text: str) -> float:
    text = str(ingredient_text or "").strip().lower()
    if not text:
        return 1.0
    match = re.match(r"(?P<amount>\d+(?:\.\d+)?)\s*(?P<unit>[A-Za-z]+)?", text)
    amount = safe_float(match.group("amount")) if match else None
    amount = amount if amount is not None and amount > 0 else 1.0
    unit = (match.group("unit") if match else "") or ""
    unit = unit.strip().lower()
    if unit in {"g", "ml"}:
        base = amount
    elif unit in {"slice", "slices"}:
        base = amount * 35.0
    else:
        if "egg" in text:
            base = amount * 60.0
        elif "banana" in text:
            base = amount * 120.0
        elif "apple" in text:
            base = amount * 180.0
        elif "berries" in text:
            base = amount * 90.0
        else:
            base = amount * 80.0
    return max(base, 1.0)


def distribute_total_by_scores(total_value: float, scores: list[float], decimals: int = 1) -> list[float]:
    if total_value <= 0 or not scores:
        return [0.0 for _ in scores]
    normalized_scores = [max(float(score or 0), 0.0) for score in scores]
    total_score = sum(normalized_scores)
    if total_score <= 0:
        normalized_scores = [1.0 for _ in scores]
        total_score = float(len(scores))
    raw_values = [(float(total_value) * score) / total_score for score in normalized_scores]
    rounded_values = [round(value, decimals) for value in raw_values]
    difference = round(float(total_value) - sum(rounded_values), decimals)
    if rounded_values:
        rounded_values[-1] = round(rounded_values[-1] + difference, decimals)
    return rounded_values


def build_diet_meal_item_specs(ingredient_lines: list[str], calories: float, nutrient_map: dict | None = None):
    items = [str(item or "").strip() for item in ingredient_lines if str(item or "").strip()]
    if not items:
        return []
    scores = [ingredient_share_score(item) for item in items]
    share_denominator = sum(scores) or float(len(scores))
    allocated_calories = distribute_total_by_scores(float(calories or 0), scores, decimals=1)
    allocated_nutrients = {
        nutrient_code: distribute_total_by_scores(float(amount or 0), scores, decimals=1)
        for nutrient_code, amount in (nutrient_map or {}).items()
    }
    payloads = []
    for index, item_text in enumerate(items, start=1):
        nutrient_values = {
            nutrient_code: nutrient_list[index - 1]
            for nutrient_code, nutrient_list in allocated_nutrients.items()
        }
        payloads.append(
            {
                "item_order": index,
                "item_text": item_text,
                "status": "pending",
                "share_ratio": round(scores[index - 1] / share_denominator, 4) if share_denominator else 0.0,
                "calories": allocated_calories[index - 1],
                "nutrients": nutrient_values,
            }
        )
    return payloads


def round_to_step(value: float, step: float | None):
    if not step or step <= 0:
        return value
    return round(round(value / step) * step, 2)


def format_scaled_amount(quantity: float, unit: str):
    if quantity <= 0:
        quantity = 0
    if unit in {"g", "ml"}:
        rounded = int(round(quantity))
        return str(rounded)
    if abs(quantity - round(quantity)) < 0.05:
        return str(int(round(quantity)))
    return f"{quantity:.1f}".rstrip("0").rstrip(".")


def scale_ingredient_lines(ingredients: list[dict], multiplier: float):
    lines = []
    for item in ingredients:
        scaled_quantity = float(item.get("qty", 0) or 0) * multiplier
        scaled_quantity = round_to_step(scaled_quantity, item.get("step"))
        amount_text = format_scaled_amount(scaled_quantity, str(item.get("unit", "")))
        unit = str(item.get("unit", "")).strip()
        label_parts = [amount_text]
        if unit:
            label_parts.append(unit)
        ingredient_name = str(item.get("item", "")).strip()
        if ingredient_name:
            label_parts.append(ingredient_name)
        lines.append(" ".join(part for part in label_parts if part))
    return lines


def scale_nutrient_map(nutrients: dict | None, multiplier: float):
    scaled = {}
    for nutrient_code in DIET_TRACKED_NUTRIENT_CODES:
        amount = safe_float((nutrients or {}).get(nutrient_code))
        if amount is None:
            continue
        scaled[nutrient_code] = round(float(amount) * multiplier, 1)
    return scaled


def diet_bias_order(profile: dict, recommended_target: dict, diet_preference_keys: list[str] | None = None):
    preference_keys = effective_diet_preference_keys(diet_preference_keys)
    if "higher_carb" in preference_keys:
        return ["gain", "balanced", "lean"]
    if "lower_carb" in preference_keys or "lighter_day" in preference_keys or "high_protein" in preference_keys:
        return ["lean", "balanced", "gain"]
    current_weight = profile.get("current_weight_kg")
    desired_weight = profile.get("desired_weight_kg") or recommended_target.get("recommended_weight_kg")
    goal = profile.get("goal", "recomp")
    if goal == "fat_loss" or (current_weight is not None and desired_weight is not None and desired_weight < current_weight - 0.5):
        return ["lean", "balanced", "gain"]
    if goal in {"muscle_gain", "strength"} or (current_weight is not None and desired_weight is not None and desired_weight > current_weight + 0.5):
        return ["gain", "balanced", "lean"]
    return ["balanced", "lean", "gain"]


def estimate_diet_targets(profile: dict, recommended_target: dict, diet_preference_keys: list[str] | None = None):
    preference_keys = effective_diet_preference_keys(diet_preference_keys)
    current_weight = profile.get("current_weight_kg") or recommended_target.get("recommended_weight_kg") or 75.0
    desired_weight = profile.get("desired_weight_kg") or recommended_target.get("recommended_weight_kg") or current_weight
    training_days = profile.get("training_days_per_week") or DEFAULT_GYM_DAYS_PER_WEEK
    goal = profile.get("goal", "recomp")
    gender_config = gender_profile_config(profile.get("gender"))
    age_years = safe_float(profile.get("age_years"))
    height_cm = safe_float(profile.get("height_cm"))
    delta_weight = desired_weight - current_weight
    predicted_bmr = predicted_profile_bmr_kcal_day(
        {
            **profile,
            "current_weight_kg": current_weight,
        }
    )
    if predicted_bmr is not None:
        maintenance_calories = int(round((predicted_bmr * diet_activity_multiplier(training_days)) / 10.0) * 10)
    else:
        maintenance_calories = int(
            round(
                (current_weight * gender_config["maintenance_factor"] + max(training_days - 3, 0) * gender_config["training_bonus"]) / 10.0
            )
            * 10
        )
    if goal == "fat_loss":
        calorie_adjustment = -750
    elif goal == "recomp":
        calorie_adjustment = -650 if delta_weight <= -0.5 else -600
    elif goal == "muscle_gain":
        calorie_adjustment = 220
    elif goal == "strength":
        calorie_adjustment = 140
    else:
        calorie_adjustment = -600 if delta_weight < -0.5 else 0
    if delta_weight <= -2.0:
        calorie_adjustment -= 80
    elif delta_weight >= 2.0:
        calorie_adjustment += 80
    if "lighter_day" in preference_keys:
        calorie_adjustment -= 150
    target_calories = maintenance_calories + calorie_adjustment
    deficit_mode = calorie_adjustment < 0
    target_protein = int(round(max(140.0, current_weight * (2.0 if deficit_mode else 1.8))))
    target_fat = int(round(max(45.0, current_weight * (0.6 if deficit_mode else 0.75))))
    minimum_carbs = 110 if deficit_mode else 130
    minimum_calories = target_protein * 4 + target_fat * 9 + minimum_carbs * 4
    if target_calories < minimum_calories:
        target_calories = int(round(minimum_calories / 10.0) * 10)
    target_calories = int(round(target_calories / 10.0) * 10)
    target_carbs = int(round(max(float(minimum_carbs), (target_calories - target_protein * 4 - target_fat * 9) / 4.0)))
    if "high_protein" in preference_keys:
        target_protein = int(round(max(target_protein, current_weight * 2.2)))
    if "lower_carb" in preference_keys:
        target_carbs = int(round(max(90.0, target_carbs - 35.0)))
        target_fat = int(round(max(target_fat, target_fat + 10.0)))
    elif "higher_carb" in preference_keys:
        target_carbs = int(round(max(target_carbs + 30.0, target_carbs)))
        target_fat = int(round(max(40.0, target_fat - 8.0)))
    minimum_calories = target_protein * 4 + target_fat * 9 + target_carbs * 4
    if target_calories < minimum_calories:
        target_calories = int(round(minimum_calories / 10.0) * 10)
    adjustment_amount = abs(maintenance_calories - target_calories)
    if target_calories < maintenance_calories:
        adjustment_label = "Calorie Deficit"
        adjustment_text = f"{adjustment_amount} kcal below maintenance"
    elif target_calories > maintenance_calories:
        adjustment_label = "Calorie Surplus"
        adjustment_text = f"{adjustment_amount} kcal above maintenance"
    else:
        adjustment_label = "Calorie Target"
        adjustment_text = "At maintenance"
    if delta_weight < -0.5:
        direction_text = f"aiming to move about {abs(delta_weight):.1f} kg down toward your current target"
    elif delta_weight > 0.5:
        direction_text = f"aiming to move about {abs(delta_weight):.1f} kg up toward your current target"
    else:
        direction_text = "aiming to move down gradually while keeping the plan sustainable" if deficit_mode else "aiming to hold close to your current target weight"
    if predicted_bmr is not None:
        gender_note = f"using {gender_config['label'].lower()} calorie assumptions with age-aware maintenance math"
    elif gender_config["key"] == "unspecified":
        gender_note = "using a neutral profile because gender is not set"
    else:
        missing_fields = []
        if age_years is None:
            missing_fields.append("age")
        if height_cm is None:
            missing_fields.append("height")
        if missing_fields:
            missing_text = " and ".join(missing_fields)
            gender_note = f"using {gender_config['label'].lower()} calorie assumptions with fallback maintenance math because {missing_text} is missing"
        else:
            gender_note = f"using {gender_config['label'].lower()} calorie assumptions"
    preference_note = ""
    if preference_keys:
        preference_note = f" Active focus: {', '.join(diet_preference_labels(preference_keys)).lower()}."
    note = f"This plan uses {adjustment_text.lower()}, keeps prep simple, and is calculated {gender_note} while protein stays high.{preference_note}"
    return {
        "current_weight_kg": round(float(current_weight), 1),
        "desired_weight_kg": round(float(desired_weight), 1),
        "maintenance_calories": maintenance_calories,
        "target_calories": target_calories,
        "target_protein_g": target_protein,
        "target_carbs_g": target_carbs,
        "target_fat_g": target_fat,
        "adjustment_label": adjustment_label,
        "adjustment_text": adjustment_text,
        "daily_deficit_kcal": max(0, maintenance_calories - target_calories),
        "direction_text": direction_text,
        "gender_label": gender_config["label"],
        "note": note,
    }


def build_diet_nutrient_targets(profile: dict, target_metrics: dict):
    current_weight = safe_float(profile.get("current_weight_kg")) or target_metrics.get("current_weight_kg") or 75.0
    height_cm = safe_float(profile.get("height_cm")) or 175.0
    fiber_target = max(
        25.0,
        round((target_metrics.get("target_calories", 2000) * 0.014) + max(0.0, (height_cm - 170.0) * 0.02) + max(0.0, (current_weight - 75.0) * 0.03), 1),
    )
    return {
        "protein_g": round(float(target_metrics.get("target_protein_g") or 0), 1),
        "carbs_g": round(float(target_metrics.get("target_carbs_g") or 0), 1),
        "fat_g": round(float(target_metrics.get("target_fat_g") or 0), 1),
        "fiber_g": fiber_target,
    }


def build_diet_nutrient_cards(nutrient_definitions: list[dict], consumed: dict, targets: dict):
    cards = []
    for definition in nutrient_definitions:
        code = definition["code"]
        target_value = safe_float(targets.get(code))
        if target_value is None or target_value <= 0:
            continue
        consumed_value = safe_float(consumed.get(code)) or 0.0
        percent = max(0.0, min(100.0, round((consumed_value / target_value) * 100.0, 1)))
        cards.append(
            {
                "code": code,
                "label": definition["label"],
                "unit": definition["unit"],
                "current_value": round(consumed_value, 1),
                "target_value": round(target_value, 1),
                "current_text": f"{consumed_value:.1f} {definition['unit']}",
                "target_text": f"{target_value:.1f} {definition['unit']}",
                "percent": percent,
                "percent_text": f"{percent:.0f}%",
                "description": definition.get("description") or "",
            }
        )
    return cards


def build_diet_calorie_card(plan: dict):
    target_value = safe_float(plan.get("target_calories"))
    if target_value is None or target_value <= 0:
        return None
    consumed_value = safe_float((plan.get("progress") or {}).get("consumed_calories"))
    if consumed_value is None:
        consumed_value = safe_float(plan.get("consumed_calories")) or 0.0
    percent = max(0.0, min(100.0, round((consumed_value / target_value) * 100.0, 1)))
    return {
        "label": "Calories",
        "description": "Only ingredient items marked done count toward today's intake.",
        "current_value": round(consumed_value, 1),
        "target_value": round(target_value, 1),
        "current_text": f"{consumed_value:.0f} kcal",
        "target_text": f"{target_value:.0f} kcal",
        "percent": percent,
        "percent_text": f"{percent:.0f}%",
    }


def select_diet_message(messages: list[str], plan_date: str):
    if not messages:
        messages = DIET_FALLBACK_MESSAGES
    day_index = datetime.fromisoformat(plan_date).date().toordinal()
    return messages[day_index % len(messages)]


def build_diet_signature(
    profile: dict,
    recommended_target: dict,
    plan_date: str,
    *,
    diet_preference_keys: list[str] | None = None,
    meals_per_day_override: int | None = None,
    meal_override_events: list[dict] | None = None,
):
    payload = {
        "planner_version": DIET_PLANNER_VERSION,
        "plan_date": plan_date,
        "goal": profile.get("goal"),
        "gender": normalize_gender(profile.get("gender")),
        "current_weight_kg": round(profile.get("current_weight_kg"), 1) if profile.get("current_weight_kg") is not None else None,
        "desired_weight_kg": round(profile.get("desired_weight_kg"), 1) if profile.get("desired_weight_kg") is not None else (round(recommended_target.get("recommended_weight_kg"), 1) if recommended_target.get("recommended_weight_kg") is not None else None),
        "training_days_per_week": profile.get("training_days_per_week"),
        "preferred_meals_per_day": meals_per_day_override or normalize_meals_per_day(profile.get("preferred_meals_per_day"), DEFAULT_DIET_MEALS_PER_DAY),
        "diet_preferences": effective_diet_preference_keys(diet_preference_keys),
        "meal_overrides": [
            {
                "slot_key": normalize_diet_slot_key(item.get("slot_key")),
                "mode": normalize_diet_meal_customization_mode(item.get("mode")) or "append",
                "items": [
                    {
                        "food_name": str(entry.get("food_name") or "").strip(),
                        "servings": round(float(safe_float(entry.get("servings")) or 1.0), 2),
                        "serving_text": str(entry.get("serving_text") or "").strip(),
                        "calories": safe_float(entry.get("calories")),
                        "nutrients": normalize_diet_food_nutrients(entry.get("nutrients") or entry),
                    }
                    for entry in normalize_diet_meal_override_items(item)
                    if str(entry.get("food_name") or "").strip()
                ],
            }
            for item in (meal_override_events or [])
            if normalize_diet_slot_key(item.get("slot_key")) and normalize_diet_meal_override_items(item)
        ],
    }
    return json.dumps(payload, sort_keys=True)


def choose_diet_template(
    candidates: list[dict],
    bias_order: list[str],
    seed_value: int,
    used_keys: set[str],
    slot_target_calories: int | None = None,
    *,
    easy_prep: bool = False,
):
    if not candidates:
        return None
    order_map = {name: index for index, name in enumerate(bias_order)}
    pool = [item for item in candidates if item.get("template_key") not in used_keys] or list(candidates)
    ranked = sorted(
        pool,
        key=lambda item: (
            order_map.get(item.get("goal_bias"), len(order_map)),
            abs((item.get("base_calories") or 0) - (slot_target_calories or item.get("base_calories") or 0)),
            len(item.get("ingredients", [])) if easy_prep else 0,
            len(str(item.get("prep_text") or "")) if easy_prep else 0,
            item.get("sort_order", 0),
            item.get("template_key", ""),
        ),
    )
    selection_pool = ranked[: min(2, len(ranked))]
    return selection_pool[seed_value % len(selection_pool)]


def generate_daily_diet_plan(
    profile: dict,
    recommended_target: dict,
    plan_date: str,
    *,
    diet_preference_keys: list[str] | None = None,
    meals_per_day_override: int | None = None,
    meal_override_events: list[dict] | None = None,
    override_notes: list[str] | None = None,
):
    preference_keys = effective_diet_preference_keys(diet_preference_keys)
    target_metrics = estimate_diet_targets(profile, recommended_target, preference_keys)
    nutrient_targets = build_diet_nutrient_targets(profile, target_metrics)
    meal_count = meals_per_day_override or normalize_meals_per_day(profile.get("preferred_meals_per_day"), DEFAULT_DIET_MEALS_PER_DAY)
    slot_sequence = DIET_SLOT_SEQUENCES.get(meal_count, DIET_SLOT_SEQUENCES[DEFAULT_DIET_MEALS_PER_DAY])
    templates = list_diet_meal_templates()
    messages = list_diet_messages()
    bias_order = diet_bias_order(profile, recommended_target, preference_keys)
    date_seed = datetime.fromisoformat(plan_date).date().toordinal()
    used_keys: set[str] = set()
    meal_override_events = list(meal_override_events or [])
    meals = []
    total_calories = 0
    total_protein = 0
    total_carbs = 0
    total_fat = 0
    total_nutrients = {code: 0.0 for code in DIET_TRACKED_NUTRIENT_CODES}
    for meal_order, slot in enumerate(slot_sequence, start=1):
        slot_key = slot["key"]
        slot_group = diet_slot_group(slot_key)
        slot_candidates = [item for item in templates if item.get("slot_key") == slot_group]
        slot_target_calories = max(220, int(round(target_metrics["target_calories"] * slot["share"])))
        template = choose_diet_template(
            slot_candidates,
            bias_order,
            date_seed + meal_order,
            used_keys,
            slot_target_calories=slot_target_calories,
            easy_prep="easy_prep" in preference_keys,
        )
        if template is None:
            continue
        used_keys.add(template["template_key"])
        portion_multiplier = max(0.55, min(1.15, slot_target_calories / max(template["base_calories"], 1)))
        calories = int(round(template["base_calories"] * portion_multiplier))
        protein_g = int(round(template["protein_g"] * portion_multiplier))
        carbs_g = int(round(template["carbs_g"] * portion_multiplier))
        fat_g = int(round(template["fat_g"] * portion_multiplier))
        nutrient_map = scale_nutrient_map(template.get("nutrients"), portion_multiplier)
        ingredient_lines = scale_ingredient_lines(template.get("ingredients", []), portion_multiplier)
        meal_payload = {
            "meal_order": meal_order,
            "slot_key": slot_key,
            "meal_label": slot["label"],
            "meal_time_text": diet_meal_time_window(slot_key, meal_count),
            "template_key": template["template_key"],
            "meal_title": template["title"],
            "meal_summary": template["summary"],
            "prep_text": template["prep_text"],
            "portion_multiplier": round(portion_multiplier, 2),
            "calories": calories,
            "protein_g": protein_g,
            "carbs_g": carbs_g,
            "fat_g": fat_g,
            "nutrients": nutrient_map,
            "ingredients": ingredient_lines,
            "items": build_diet_meal_item_specs(ingredient_lines, calories, nutrient_map),
            "status": "pending",
            "is_done": False,
            "is_skipped": False,
        }
        slot_override_events = [
            event
            for event in meal_override_events
            if normalize_diet_slot_key(event.get("slot_key")) == diet_slot_group(slot_key)
        ]
        if slot_override_events:
            meal_payload, _ = apply_diet_meal_customizations(meal_payload, slot_override_events)
        total_calories += int(round(float(meal_payload.get("calories") or 0)))
        total_protein += int(round(float(meal_payload.get("protein_g") or 0)))
        total_carbs += int(round(float(meal_payload.get("carbs_g") or 0)))
        total_fat += int(round(float(meal_payload.get("fat_g") or 0)))
        for nutrient_code, amount in (meal_payload.get("nutrients") or {}).items():
            total_nutrients[nutrient_code] = round(total_nutrients.get(nutrient_code, 0.0) + float(amount or 0.0), 1)
        meals.append(meal_payload)
    message_text = select_diet_message(messages, plan_date)
    rationale = (
        f"{target_metrics['note']} {meal_count} simple meal"
        f"{'' if meal_count == 1 else 's'} keep the day practical instead of fancy."
    )
    if override_notes:
        rationale = f"{rationale} {' '.join(override_notes)}"
    return {
        "plan_date": plan_date,
        "meal_count": meal_count,
        "message_text": message_text,
        "rationale_text": rationale,
        "maintenance_calories": target_metrics["maintenance_calories"],
        "target_calories": target_metrics["target_calories"],
        "target_protein_g": target_metrics["target_protein_g"],
        "target_carbs_g": target_metrics["target_carbs_g"],
        "target_fat_g": target_metrics["target_fat_g"],
        "adjustment_label": target_metrics["adjustment_label"],
        "adjustment_text": target_metrics["adjustment_text"],
        "daily_deficit_kcal": target_metrics["daily_deficit_kcal"],
        "current_weight_kg": target_metrics["current_weight_kg"],
        "desired_weight_kg": target_metrics["desired_weight_kg"],
        "goal": profile.get("goal", "recomp"),
        "total_calories": total_calories,
        "total_protein_g": total_protein,
        "total_carbs_g": total_carbs,
        "total_fat_g": total_fat,
        "nutrient_targets": nutrient_targets,
        "total_nutrients": total_nutrients,
        "active_preference_keys": preference_keys,
        "active_preference_labels": diet_preference_labels(diet_preference_keys),
        "override_notes": list(override_notes or []),
        "meals": meals,
    }


def sync_diet_meal_status_from_items(connection: sqlite3.Connection, meal_id: int):
    item_rows = connection.execute(
        """
        SELECT status
        FROM diet_daily_meal_items
        WHERE meal_id = ?
        ORDER BY item_order, id
        """,
        (meal_id,),
    ).fetchall()
    summary = summarize_diet_item_statuses([{"status": row["status"]} for row in item_rows])
    connection.execute(
        "UPDATE diet_daily_meals SET status = ? WHERE id = ?",
        (summary["status"], meal_id),
    )
    return summary


def ensure_diet_meal_items_for_meal(
    connection: sqlite3.Connection,
    meal_row: sqlite3.Row,
    nutrient_map: dict[str, float] | None = None,
):
    existing_count = connection.execute(
        "SELECT COUNT(*) AS count FROM diet_daily_meal_items WHERE meal_id = ?",
        (meal_row["id"],),
    ).fetchone()["count"]
    if existing_count:
        return
    try:
        ingredient_lines = json.loads(meal_row["ingredients_json"] or "[]")
    except json.JSONDecodeError:
        ingredient_lines = []
    item_specs = build_diet_meal_item_specs(ingredient_lines, float(meal_row["calories"] or 0), nutrient_map or {})
    if not item_specs:
        return
    connection.executemany(
        """
        INSERT INTO diet_daily_meal_items (
            meal_id, item_order, item_text, status, share_ratio, calories
        ) VALUES (?, ?, ?, ?, ?, ?)
        """,
        [
            (
                meal_row["id"],
                item["item_order"],
                item["item_text"],
                normalize_diet_meal_status(item.get("status")),
                item["share_ratio"],
                item["calories"],
            )
            for item in item_specs
        ],
    )
    inserted_rows = connection.execute(
        """
        SELECT id, item_order
        FROM diet_daily_meal_items
        WHERE meal_id = ?
        ORDER BY item_order, id
        """,
        (meal_row["id"],),
    ).fetchall()
    item_id_by_order = {row["item_order"]: row["id"] for row in inserted_rows}
    connection.executemany(
        """
        INSERT INTO diet_daily_meal_item_nutrients (
            meal_item_id, nutrient_code, amount
        ) VALUES (?, ?, ?)
        """,
        [
            (item_id_by_order[item["item_order"]], nutrient_code, amount)
            for item in item_specs
            if item["item_order"] in item_id_by_order
            for nutrient_code, amount in (item.get("nutrients") or {}).items()
        ],
    )
    sync_diet_meal_status_from_items(connection, meal_row["id"])


def normalize_diet_item_text_key(item_text: str | None):
    return re.sub(r"\s+", " ", str(item_text or "").strip().lower())


def capture_diet_item_status_snapshot(connection: sqlite3.Connection, plan_id: int):
    rows = connection.execute(
        """
        SELECT
            meals.slot_key,
            items.item_text,
            items.status
        FROM diet_daily_meal_items AS items
        JOIN diet_daily_meals AS meals ON meals.id = items.meal_id
        WHERE meals.plan_id = ?
        ORDER BY meals.meal_order, items.item_order, items.id
        """,
        (plan_id,),
    ).fetchall()
    snapshot: dict[tuple[str, str], list[str]] = {}
    for row in rows:
        key = (
            normalize_diet_slot_key(row["slot_key"]) or str(row["slot_key"] or "").strip().lower(),
            normalize_diet_item_text_key(row["item_text"]),
        )
        snapshot.setdefault(key, []).append(normalize_diet_meal_status(row["status"]))
    return snapshot


def restore_diet_item_status_snapshot(connection: sqlite3.Connection, plan_id: int, snapshot: dict):
    if not snapshot:
        return
    rows = connection.execute(
        """
        SELECT
            items.id,
            items.meal_id,
            items.item_text,
            items.status,
            meals.slot_key
        FROM diet_daily_meal_items AS items
        JOIN diet_daily_meals AS meals ON meals.id = items.meal_id
        WHERE meals.plan_id = ?
        ORDER BY meals.meal_order, items.item_order, items.id
        """,
        (plan_id,),
    ).fetchall()
    working_snapshot = {
        key: list(values)
        for key, values in snapshot.items()
    }
    updates = []
    touched_meal_ids = set()
    for row in rows:
        key = (
            normalize_diet_slot_key(row["slot_key"]) or str(row["slot_key"] or "").strip().lower(),
            normalize_diet_item_text_key(row["item_text"]),
        )
        status_list = working_snapshot.get(key)
        if not status_list:
            continue
        restored_status = normalize_diet_meal_status(status_list.pop(0))
        if restored_status == normalize_diet_meal_status(row["status"]):
            continue
        updates.append((restored_status, row["id"]))
        touched_meal_ids.add(row["meal_id"])
    if updates:
        connection.executemany(
            """
            UPDATE diet_daily_meal_items
            SET status = ?
            WHERE id = ?
            """,
            updates,
        )
    for meal_id in touched_meal_ids:
        sync_diet_meal_status_from_items(connection, meal_id)


def hydrate_diet_plan(
    plan_row: sqlite3.Row,
    meal_rows: list[sqlite3.Row],
    meal_nutrient_rows: list[sqlite3.Row] | None = None,
    target_rows: list[sqlite3.Row] | None = None,
    meal_item_rows: list[sqlite3.Row] | None = None,
    meal_item_nutrient_rows: list[sqlite3.Row] | None = None,
):
    meals = []
    total_calories = 0
    total_protein = 0
    total_carbs = 0
    total_fat = 0
    consumed_calories = 0
    total_nutrients = {code: 0.0 for code in DIET_TRACKED_NUTRIENT_CODES}
    consumed_nutrients = {code: 0.0 for code in DIET_TRACKED_NUTRIENT_CODES}
    nutrient_rows_by_meal: dict[int, dict[str, float]] = {}
    for row in meal_nutrient_rows or []:
        nutrient_rows_by_meal.setdefault(row["meal_id"], {})[row["nutrient_code"]] = round(float(row["amount"] or 0), 1)
    item_nutrients_by_item: dict[int, dict[str, float]] = {}
    for row in meal_item_nutrient_rows or []:
        item_nutrients_by_item.setdefault(row["meal_item_id"], {})[row["nutrient_code"]] = round(float(row["amount"] or 0), 1)
    meal_items_by_meal: dict[int, list[dict]] = {}
    for row in meal_item_rows or []:
        item_nutrient_map = item_nutrients_by_item.get(row["id"], {})
        current_status = normalize_diet_meal_status(row["status"])
        meal_items_by_meal.setdefault(row["meal_id"], []).append(
            {
                "id": row["id"],
                "item_order": row["item_order"],
                "item_text": row["item_text"],
                "status": current_status,
                "is_done": current_status == "done",
                "is_skipped": current_status == "skipped",
                "calories": round(float(row["calories"] or 0), 1),
                "calories_text": f"{float(row['calories'] or 0):.0f} kcal",
                "share_ratio": round(float(row["share_ratio"] or 0), 4),
                "nutrients": item_nutrient_map,
            }
        )
    nutrient_targets = {
        row["nutrient_code"]: round(float(row["target_amount"] or 0), 1)
        for row in (target_rows or [])
    }
    for row in meal_rows:
        try:
            ingredients = json.loads(row["ingredients_json"] or "[]")
        except json.JSONDecodeError:
            ingredients = []
        nutrient_map = nutrient_rows_by_meal.get(row["id"], {})
        item_rows = sorted(meal_items_by_meal.get(row["id"], []), key=lambda item: (item["item_order"], item["id"]))
        item_summary = summarize_diet_item_statuses(item_rows)
        meal_consumed_calories = round(
            sum(item["calories"] for item in item_rows if normalize_diet_meal_status(item.get("status")) == "done"),
            1,
        )
        meal_consumed_nutrients = {code: 0.0 for code in DIET_TRACKED_NUTRIENT_CODES}
        for item in item_rows:
            if normalize_diet_meal_status(item.get("status")) != "done":
                continue
            for nutrient_code, amount in (item.get("nutrients") or {}).items():
                meal_consumed_nutrients[nutrient_code] = round(meal_consumed_nutrients.get(nutrient_code, 0.0) + amount, 1)
        consumed_calories = round(consumed_calories + meal_consumed_calories, 1)
        for nutrient_code, amount in meal_consumed_nutrients.items():
            consumed_nutrients[nutrient_code] = round(consumed_nutrients.get(nutrient_code, 0.0) + amount, 1)
        meals.append(
            {
                "id": row["id"],
                "meal_order": row["meal_order"],
                "slot_key": row["slot_key"],
                "meal_label": row["meal_label"],
                "meal_time_text": diet_meal_time_window(row["slot_key"], plan_row["meal_count"]),
                "template_key": row["template_key"],
                "meal_title": row["meal_title"],
                "meal_summary": normalize_diet_meal_summary_text(row["template_key"], row["meal_summary"]),
                "prep_text": row["prep_text"],
                "portion_multiplier": row["portion_multiplier"],
                "calories": row["calories"],
                "protein_g": row["protein_g"],
                "carbs_g": row["carbs_g"],
                "fat_g": row["fat_g"],
                "nutrients": nutrient_map,
                "ingredients": ingredients,
                "items": item_rows,
                "item_summary": item_summary,
                "consumed_calories": meal_consumed_calories,
                "consumed_nutrients": meal_consumed_nutrients,
                "status": item_summary["status"],
                "is_done": item_summary["is_done"],
                "is_skipped": item_summary["is_skipped"],
                "is_partial": item_summary["is_partial"],
                "is_pending": item_summary["is_pending"],
            }
        )
        total_calories += row["calories"]
        total_protein += row["protein_g"]
        total_carbs += row["carbs_g"]
        total_fat += row["fat_g"]
        for nutrient_code, amount in nutrient_map.items():
            total_nutrients[nutrient_code] = round(total_nutrients.get(nutrient_code, 0.0) + amount, 1)
    return {
        "plan_date": plan_row["plan_date"],
        "meal_count": plan_row["meal_count"],
        "message_text": plan_row["message_text"],
        "rationale_text": plan_row["rationale_text"],
        "target_calories": plan_row["target_calories"],
        "target_protein_g": plan_row["target_protein_g"],
        "target_carbs_g": plan_row["target_carbs_g"],
        "target_fat_g": plan_row["target_fat_g"],
        "current_weight_kg": plan_row["current_weight_kg"],
        "desired_weight_kg": plan_row["desired_weight_kg"],
        "goal": plan_row["goal"],
        "total_calories": total_calories,
        "consumed_calories": consumed_calories,
        "total_protein_g": total_protein,
        "total_carbs_g": total_carbs,
        "total_fat_g": total_fat,
        "total_nutrients": total_nutrients,
        "consumed_nutrients": consumed_nutrients,
        "nutrient_targets": nutrient_targets,
        "meals": meals,
    }


def diet_plan_progress(meals: list[dict]):
    total_count = len(meals)
    done_count = sum(1 for meal in meals if normalize_diet_meal_status(meal.get("status")) == "done")
    skipped_count = sum(1 for meal in meals if normalize_diet_meal_status(meal.get("status")) == "skipped")
    partial_count = sum(1 for meal in meals if normalize_diet_meal_status(meal.get("status")) == "partial")
    resolved_count = done_count + skipped_count
    consumed_calories = round(
        sum(safe_float(meal.get("consumed_calories")) or 0 for meal in meals),
        1,
    )
    return {
        "total_count": total_count,
        "done_count": done_count,
        "skipped_count": skipped_count,
        "partial_count": partial_count,
        "resolved_count": resolved_count,
        "pending_count": max(0, total_count - resolved_count - partial_count),
        "is_complete": total_count > 0 and resolved_count >= total_count,
        "consumed_calories": consumed_calories,
    }


def diet_calorie_history(username: str, limit: int = 21):
    user_id = ensure_app_user(username)
    with gym_user_connection() as connection:
        rows = connection.execute(
            """
            SELECT
                p.plan_date,
                COALESCE(
                    (
                        SELECT SUM(
                            CASE
                                WHEN EXISTS (
                                    SELECT 1 FROM diet_daily_meal_items AS items_exist
                                    WHERE items_exist.meal_id = meals.id
                                )
                                THEN COALESCE(
                                    (
                                        SELECT SUM(items_done.calories)
                                        FROM diet_daily_meal_items AS items_done
                                        WHERE items_done.meal_id = meals.id
                                          AND LOWER(COALESCE(items_done.status, 'pending')) = 'done'
                                    ),
                                    0
                                )
                                WHEN LOWER(COALESCE(meals.status, 'pending')) = 'done'
                                THEN meals.calories
                                ELSE 0
                            END
                        )
                        FROM diet_daily_meals AS meals
                        WHERE meals.plan_id = p.id
                    ),
                    0
                ) AS consumed_calories
            FROM diet_daily_plans p
            WHERE p.user_id = ?
            ORDER BY p.plan_date DESC, p.id DESC
            LIMIT ?
            """,
            (user_id, max(1, safe_int(limit, 21) or 21)),
        ).fetchall()
    ordered_rows = list(reversed(rows))
    return [
        {
            "date": row["plan_date"],
            "value": float(row["consumed_calories"] or 0),
        }
        for row in ordered_rows
    ]


def diet_history_series(username: str, limit: int = 400):
    user_id = ensure_app_user(username)
    with gym_user_connection() as connection:
        rows = connection.execute(
            """
            SELECT
                p.plan_date,
                p.target_calories,
                COALESCE(
                    (
                        SELECT SUM(
                            CASE
                                WHEN EXISTS (
                                    SELECT 1 FROM diet_daily_meal_items AS items_exist
                                    WHERE items_exist.meal_id = meals.id
                                )
                                THEN COALESCE(
                                    (
                                        SELECT SUM(items_done.calories)
                                        FROM diet_daily_meal_items AS items_done
                                        WHERE items_done.meal_id = meals.id
                                          AND LOWER(COALESCE(items_done.status, 'pending')) = 'done'
                                    ),
                                    0
                                )
                                WHEN LOWER(COALESCE(meals.status, 'pending')) = 'done'
                                THEN meals.calories
                                ELSE 0
                            END
                        )
                        FROM diet_daily_meals AS meals
                        WHERE meals.plan_id = p.id
                    ),
                    0
                ) AS consumed_calories,
                COALESCE(
                    (
                        SELECT COUNT(*)
                        FROM diet_daily_meals AS meals
                        WHERE meals.plan_id = p.id
                          AND LOWER(COALESCE(meals.status, 'pending')) IN ('done', 'skipped')
                    ),
                    0
                ) AS resolved_count,
                COALESCE(
                    (
                        SELECT COUNT(*)
                        FROM diet_daily_meals AS meals
                        WHERE meals.plan_id = p.id
                    ),
                    0
                ) AS meal_count
            FROM diet_daily_plans p
            WHERE p.user_id = ?
            ORDER BY p.plan_date DESC, p.id DESC
            LIMIT ?
            """,
            (user_id, max(7, safe_int(limit, 400) or 400)),
        ).fetchall()
    ordered_rows = list(reversed(rows))
    return [
        {
            "date": row["plan_date"],
            "value": float(row["consumed_calories"] or 0),
            "target": float(row["target_calories"] or 0),
            "resolved_count": int(row["resolved_count"] or 0),
            "meal_count": int(row["meal_count"] or 0),
        }
        for row in ordered_rows
    ]


def ensure_daily_diet_plan(username: str, profile: dict, recommended_target: dict, plan_date: str | None = None):
    user_id = ensure_app_user(username)
    current_plan_date = plan_date or datetime.now().date().isoformat()
    with gym_user_connection() as connection:
        existing_item_status_snapshot = {}
        override_state = collect_diet_override_state_for_user(connection, user_id, current_plan_date)
        signature = build_diet_signature(
            profile,
            recommended_target,
            current_plan_date,
            diet_preference_keys=override_state["active_preference_keys"],
            meals_per_day_override=override_state["today_meals_per_day"],
            meal_override_events=override_state.get("meal_events"),
        )
        plan_row = connection.execute(
            """
            SELECT * FROM diet_daily_plans
            WHERE user_id = ? AND plan_date = ?
            ORDER BY id DESC LIMIT 1
            """,
            (user_id, current_plan_date),
        ).fetchone()
        if plan_row and plan_row["source_signature"] == signature:
            meal_rows = connection.execute(
                """
                SELECT * FROM diet_daily_meals
                WHERE plan_id = ?
                ORDER BY meal_order, id
                """,
                (plan_row["id"],),
            ).fetchall()
            if meal_rows:
                meal_nutrient_rows = connection.execute(
                    """
                    SELECT *
                    FROM diet_daily_meal_nutrients
                    WHERE meal_id IN (
                        SELECT id FROM diet_daily_meals WHERE plan_id = ?
                    )
                    ORDER BY meal_id, nutrient_code
                    """,
                    (plan_row["id"],),
                ).fetchall()
                nutrient_map_by_meal = {}
                for nutrient_row in meal_nutrient_rows:
                    nutrient_map_by_meal.setdefault(nutrient_row["meal_id"], {})[nutrient_row["nutrient_code"]] = round(float(nutrient_row["amount"] or 0), 1)
                for meal_row in meal_rows:
                    ensure_diet_meal_items_for_meal(connection, meal_row, nutrient_map_by_meal.get(meal_row["id"], {}))
                meal_item_rows = connection.execute(
                    """
                    SELECT *
                    FROM diet_daily_meal_items
                    WHERE meal_id IN (
                        SELECT id FROM diet_daily_meals WHERE plan_id = ?
                    )
                    ORDER BY meal_id, item_order, id
                    """,
                    (plan_row["id"],),
                ).fetchall()
                meal_item_nutrient_rows = connection.execute(
                    """
                    SELECT *
                    FROM diet_daily_meal_item_nutrients
                    WHERE meal_item_id IN (
                        SELECT id FROM diet_daily_meal_items
                        WHERE meal_id IN (SELECT id FROM diet_daily_meals WHERE plan_id = ?)
                    )
                    ORDER BY meal_item_id, nutrient_code
                    """,
                    (plan_row["id"],),
                ).fetchall()
                target_rows = connection.execute(
                    """
                    SELECT *
                    FROM diet_daily_plan_targets
                    WHERE plan_id = ?
                    ORDER BY nutrient_code
                    """,
                    (plan_row["id"],),
                ).fetchall()
                hydrated = hydrate_diet_plan(plan_row, meal_rows, meal_nutrient_rows, target_rows, meal_item_rows, meal_item_nutrient_rows)
                hydrated["active_preference_keys"] = list(override_state["active_preference_keys"])
                hydrated["active_preference_labels"] = list(override_state["active_preference_labels"])
                hydrated["override_notes"] = list(override_state["notes"])
                return hydrated
        if plan_row:
            existing_item_status_snapshot = capture_diet_item_status_snapshot(connection, plan_row["id"])
        generated = generate_daily_diet_plan(
            profile,
            recommended_target,
            current_plan_date,
            diet_preference_keys=override_state["active_preference_keys"],
            meals_per_day_override=override_state["today_meals_per_day"],
            meal_override_events=override_state.get("meal_events"),
            override_notes=override_state["notes"],
        )
        connection.execute(
            "DELETE FROM diet_daily_plans WHERE user_id = ? AND plan_date = ?",
            (user_id, current_plan_date),
        )
        cursor = connection.execute(
            """
            INSERT INTO diet_daily_plans (
                user_id, plan_date, meal_count, target_calories, target_protein_g,
                target_carbs_g, target_fat_g, current_weight_kg, desired_weight_kg,
                goal, message_text, rationale_text, source_signature, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                user_id,
                current_plan_date,
                generated["meal_count"],
                generated["target_calories"],
                generated["target_protein_g"],
                generated["target_carbs_g"],
                generated["target_fat_g"],
                generated["current_weight_kg"],
                generated["desired_weight_kg"],
                generated["goal"],
                generated["message_text"],
                generated["rationale_text"],
                signature,
                now_iso(),
            ),
        )
        plan_id = cursor.lastrowid
        connection.executemany(
            """
            INSERT INTO diet_daily_meals (
                plan_id, meal_order, slot_key, meal_label, template_key, meal_title,
                meal_summary, prep_text, portion_multiplier, calories, protein_g, carbs_g,
                fat_g, ingredients_json, status
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    plan_id,
                    meal["meal_order"],
                    meal["slot_key"],
                    meal["meal_label"],
                    meal["template_key"],
                    meal["meal_title"],
                    meal["meal_summary"],
                    meal["prep_text"],
                    meal["portion_multiplier"],
                    meal["calories"],
                    meal["protein_g"],
                    meal["carbs_g"],
                    meal["fat_g"],
                    json.dumps(meal["ingredients"]),
                    normalize_diet_meal_status(meal.get("status")),
                )
                for meal in generated["meals"]
            ],
        )
        inserted_meals = connection.execute(
            """
            SELECT id, meal_order
            FROM diet_daily_meals
            WHERE plan_id = ?
            ORDER BY meal_order, id
            """,
            (plan_id,),
        ).fetchall()
        meal_id_by_order = {row["meal_order"]: row["id"] for row in inserted_meals}
        connection.executemany(
            """
            INSERT INTO diet_daily_meal_nutrients (
                meal_id, nutrient_code, amount
            ) VALUES (?, ?, ?)
            """,
            [
                (meal_id_by_order[meal["meal_order"]], nutrient_code, amount)
                for meal in generated["meals"]
                if meal["meal_order"] in meal_id_by_order
                for nutrient_code, amount in (meal.get("nutrients") or {}).items()
            ],
        )
        connection.executemany(
            """
            INSERT INTO diet_daily_meal_items (
                meal_id, item_order, item_text, status, share_ratio, calories
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    meal_id_by_order[meal["meal_order"]],
                    item["item_order"],
                    item["item_text"],
                    normalize_diet_meal_status(item.get("status")),
                    item["share_ratio"],
                    item["calories"],
                )
                for meal in generated["meals"]
                if meal["meal_order"] in meal_id_by_order
                for item in (meal.get("items") or [])
            ],
        )
        inserted_items = connection.execute(
            """
            SELECT id, meal_id, item_order
            FROM diet_daily_meal_items
            WHERE meal_id IN (
                SELECT id FROM diet_daily_meals WHERE plan_id = ?
            )
            ORDER BY meal_id, item_order, id
            """,
            (plan_id,),
        ).fetchall()
        item_id_by_key = {(row["meal_id"], row["item_order"]): row["id"] for row in inserted_items}
        connection.executemany(
            """
            INSERT INTO diet_daily_meal_item_nutrients (
                meal_item_id, nutrient_code, amount
            ) VALUES (?, ?, ?)
            """,
            [
                (
                    item_id_by_key[(meal_id_by_order[meal["meal_order"]], item["item_order"])],
                    nutrient_code,
                    amount,
                )
                for meal in generated["meals"]
                if meal["meal_order"] in meal_id_by_order
                for item in (meal.get("items") or [])
                if (meal_id_by_order[meal["meal_order"]], item["item_order"]) in item_id_by_key
                for nutrient_code, amount in (item.get("nutrients") or {}).items()
            ],
        )
        for meal_order, meal_id in meal_id_by_order.items():
            sync_diet_meal_status_from_items(connection, meal_id)
        restore_diet_item_status_snapshot(connection, plan_id, existing_item_status_snapshot)
        connection.executemany(
            """
            INSERT INTO diet_daily_plan_targets (
                plan_id, nutrient_code, target_amount
            ) VALUES (?, ?, ?)
            """,
            [
                (plan_id, nutrient_code, amount)
                for nutrient_code, amount in (generated.get("nutrient_targets") or {}).items()
            ],
        )
        stored_plan = connection.execute(
            "SELECT * FROM diet_daily_plans WHERE id = ?",
            (plan_id,),
        ).fetchone()
        stored_meals = connection.execute(
            """
            SELECT * FROM diet_daily_meals
            WHERE plan_id = ?
            ORDER BY meal_order, id
            """,
            (plan_id,),
        ).fetchall()
        stored_meal_nutrients = connection.execute(
            """
            SELECT *
            FROM diet_daily_meal_nutrients
            WHERE meal_id IN (
                SELECT id FROM diet_daily_meals WHERE plan_id = ?
            )
            ORDER BY meal_id, nutrient_code
            """,
            (plan_id,),
        ).fetchall()
        stored_item_rows = connection.execute(
            """
            SELECT *
            FROM diet_daily_meal_items
            WHERE meal_id IN (
                SELECT id FROM diet_daily_meals WHERE plan_id = ?
            )
            ORDER BY meal_id, item_order, id
            """,
            (plan_id,),
        ).fetchall()
        stored_item_nutrients = connection.execute(
            """
            SELECT *
            FROM diet_daily_meal_item_nutrients
            WHERE meal_item_id IN (
                SELECT id FROM diet_daily_meal_items
                WHERE meal_id IN (SELECT id FROM diet_daily_meals WHERE plan_id = ?)
            )
            ORDER BY meal_item_id, nutrient_code
            """,
            (plan_id,),
        ).fetchall()
        stored_targets = connection.execute(
            """
            SELECT *
            FROM diet_daily_plan_targets
            WHERE plan_id = ?
            ORDER BY nutrient_code
            """,
            (plan_id,),
        ).fetchall()
        hydrated = hydrate_diet_plan(stored_plan, stored_meals, stored_meal_nutrients, stored_targets, stored_item_rows, stored_item_nutrients)
        hydrated["active_preference_keys"] = list(override_state["active_preference_keys"])
        hydrated["active_preference_labels"] = list(override_state["active_preference_labels"])
        hydrated["override_notes"] = list(override_state["notes"])
        return hydrated


def save_diet_meal_status(username: str, form):
    user_id = ensure_app_user(username)
    meal_id = safe_int(form.get("meal_id"))
    if not meal_id:
        raise ValueError("Meal selection is missing. Reload the Diet tab and try again.")
    action_intent = normalize_diet_meal_status(form.get("intent", "done"))
    with gym_user_connection() as connection:
        current_row = connection.execute(
            """
            SELECT m.id, m.meal_label, m.meal_title, m.status, p.plan_date
            FROM diet_daily_meals m
            JOIN diet_daily_plans p ON p.id = m.plan_id
            WHERE m.id = ? AND p.user_id = ?
            LIMIT 1
            """,
            (meal_id, user_id),
        ).fetchone()
        if not current_row:
            raise ValueError("Meal entry was not found. Reload the Diet tab and try again.")
        item_rows = connection.execute(
            """
            SELECT id, status
            FROM diet_daily_meal_items
            WHERE meal_id = ?
            ORDER BY item_order, id
            """,
            (current_row["id"],),
        ).fetchall()
        if not item_rows:
            raise ValueError("Meal items were not found. Reload the Diet tab and try again.")
        current_summary = summarize_diet_item_statuses([{"status": row["status"]} for row in item_rows])
        next_status = "pending" if current_summary["status"] == action_intent else action_intent
        connection.execute(
            """
            UPDATE diet_daily_meal_items
            SET status = ?
            WHERE meal_id = ?
            """,
            (next_status, current_row["id"]),
        )
        sync_diet_meal_status_from_items(connection, current_row["id"])
    meal_name = current_row["meal_label"] or current_row["meal_title"] or "Meal"
    message_map = {
        "done": f"All items in {meal_name} marked done.",
        "skipped": f"All items in {meal_name} skipped.",
        "pending": f"All items in {meal_name} reset to pending.",
    }
    return {
        "status": next_status,
        "message": message_map.get(next_status, f"{meal_name} updated."),
        "updated_at": now_iso(),
    }


def save_diet_meal_item_status(username: str, form):
    user_id = ensure_app_user(username)
    meal_item_id = safe_int(form.get("meal_item_id"))
    if not meal_item_id:
        raise ValueError("Meal item selection is missing. Reload the Diet tab and try again.")
    action_intent = normalize_diet_meal_status(form.get("intent", "done"))
    with gym_user_connection() as connection:
        current_row = connection.execute(
            """
            SELECT
                items.id,
                items.meal_id,
                items.item_text,
                items.status,
                meals.meal_label,
                meals.meal_title
            FROM diet_daily_meal_items AS items
            JOIN diet_daily_meals AS meals ON meals.id = items.meal_id
            JOIN diet_daily_plans AS plans ON plans.id = meals.plan_id
            WHERE items.id = ? AND plans.user_id = ?
            LIMIT 1
            """,
            (meal_item_id, user_id),
        ).fetchone()
        if not current_row:
            raise ValueError("Meal item was not found. Reload the Diet tab and try again.")
        current_status = normalize_diet_meal_status(current_row["status"])
        next_status = "pending" if current_status == action_intent else action_intent
        connection.execute(
            """
            UPDATE diet_daily_meal_items
            SET status = ?
            WHERE id = ?
            """,
            (next_status, current_row["id"]),
        )
        meal_summary = sync_diet_meal_status_from_items(connection, current_row["meal_id"])
    item_name = current_row["item_text"] or current_row["meal_label"] or current_row["meal_title"] or "Meal item"
    meal_name = current_row["meal_label"] or current_row["meal_title"] or "Meal"
    meal_status_text = {
        "done": "complete",
        "skipped": "skipped",
        "pending": "pending",
        "partial": "partly done",
    }.get(meal_summary["status"], meal_summary["status"])
    if action_intent == "done":
        item_message = f"{item_name} checked." if next_status == "done" else f"{item_name} unchecked."
    else:
        status_text = {
            "done": "done",
            "skipped": "skipped",
            "pending": "pending",
        }.get(next_status, next_status)
        item_message = f"{item_name} marked {status_text}."
    return {
        "status": next_status,
        "message": f"{item_message} {meal_name} is now {meal_status_text}.",
        "updated_at": now_iso(),
    }


def build_agent_ui_state(agent_key: str, available: bool, insight: dict | None):
    has_cached = insight is not None
    if available and has_cached:
        status_label = "Ready"
        status_class = "running"
    elif available:
        status_label = "Available"
        status_class = "partial"
    elif has_cached:
        status_label = "Cached"
        status_class = "partial"
    else:
        status_label = "Offline"
        status_class = "stopped"
    return {
        "available": available,
        "status_label": status_label,
        "status_class": status_class,
        "button_label": "Refresh AI Notes" if has_cached else "Generate AI Notes",
        "help_text": "Codex runs on MSI only when you request it. The built-in planner stays active even if MSI is offline.",
        "offline_text": "MSI is offline, so the built-in planner is active and the AI coach is paused.",
        "insight": insight,
        "source_label": (insight or {}).get("source_label", "Codex on MSI"),
        "updated_at_text": (insight or {}).get("updated_at_text", "Not generated yet"),
    }


def diet_agent_context_key(plan: dict):
    return str(plan.get("plan_date") or datetime.now().date().isoformat())


def _legacy_build_diet_agent_prompt(profile: dict, plan: dict):
    meal_items = []
    for meal in plan.get("meals", []):
        meal_items.append(
            {
                "slot": meal.get("meal_label"),
                "title": meal.get("meal_title"),
                "status": meal.get("status"),
                "calories": meal.get("calories"),
                "protein_g": meal.get("protein_g"),
                "carbs_g": meal.get("carbs_g"),
                "fat_g": meal.get("fat_g"),
                "ingredients": meal.get("ingredients", []),
            }
        )
    payload = {
        "profile": {
            "display_name": profile.get("display_name"),
            "goal": profile.get("goal_label"),
            "gender": profile.get("gender_label"),
            "current_weight_kg": profile.get("current_weight_kg"),
            "target_weight_kg": plan.get("desired_weight_kg"),
            "training_days_per_week": profile.get("training_days_per_week"),
            "meals_per_day": profile.get("preferred_meals_per_day"),
        },
        "today_plan": {
            "date": plan.get("plan_date"),
            "target_calories": plan.get("target_calories"),
            "planned_calories": plan.get("total_calories"),
            "consumed_calories": plan.get("progress", {}).get("consumed_calories"),
            "daily_deficit_kcal": plan.get("daily_deficit_kcal"),
            "progress_text": plan.get("progress_text"),
            "rationale_text": plan.get("rationale_text"),
            "meals": meal_items,
        },
    }
    return (
        "You are the Diet AI Coach for Health Hub.\n"
        "Return JSON only with keys headline, bullets, watchout.\n"
        "Rules:\n"
        "- headline: one short sentence\n"
        "- bullets: exactly 3 short strings\n"
        "- watchout: one short sentence\n"
        "- stay practical, simple, and meal-prep focused\n"
        "- use only the provided data\n"
        "- do not ask questions\n"
        "- do not include markdown fences or extra keys\n\n"
        f"Context JSON:\n{json.dumps(payload, ensure_ascii=True)}"
    )


def build_diet_state(username: str, agent_available: bool = False):
    profile = load_gym_profile(username)
    recommended_target = suggested_weight_targets(profile)
    plan = ensure_daily_diet_plan(username, profile, recommended_target)
    current_targets = estimate_diet_targets(profile, recommended_target, plan.get("active_preference_keys"))
    progress = diet_plan_progress(plan.get("meals", []))
    history_points = diet_history_series(username)
    intake_points = [{"date": item["date"], "value": item["value"]} for item in history_points]
    latest_intake_value = intake_points[-1]["value"] if intake_points else None
    nutrient_definitions = list_diet_nutrient_definitions()
    nutrient_cards = build_diet_nutrient_cards(
        nutrient_definitions,
        plan.get("consumed_nutrients", {}),
        plan.get("nutrient_targets", {}),
    )
    calorie_card = build_diet_calorie_card(plan)
    if progress["is_complete"]:
        status_label = "Finished"
        status_class = "running"
    elif progress["partial_count"] > 0:
        status_label = f"{progress['partial_count']} partial"
        status_class = "partial"
    elif progress["done_count"] > 0 or progress["skipped_count"] > 0:
        status_label = f"{progress['done_count']} complete, {progress['skipped_count']} skipped"
        status_class = "partial"
    else:
        status_label = "Open"
        status_class = "partial"
    progress_text = (
        f"{progress['done_count']} complete, {progress['partial_count']} partial, {progress['pending_count']} pending, {progress['skipped_count']} skipped"
        if progress["total_count"]
        else "No meals yet"
    )
    plan.update(
        {
            "maintenance_calories": current_targets["maintenance_calories"],
            "target_calories": current_targets["target_calories"],
            "target_protein_g": current_targets["target_protein_g"],
            "target_carbs_g": current_targets["target_carbs_g"],
            "target_fat_g": current_targets["target_fat_g"],
            "adjustment_label": current_targets["adjustment_label"],
            "adjustment_text": current_targets["adjustment_text"],
            "daily_deficit_kcal": current_targets["daily_deficit_kcal"],
            "progress": progress,
            "progress_text": progress_text,
            "status_label": status_label,
            "status_class": status_class,
        }
    )
    context_key = diet_agent_context_key(plan)
    insight = load_coach_insight(username, "diet", context_key)
    return {
        "profile": profile,
        "recommended_target": recommended_target,
        "plan": plan,
        "agent": {
            **build_agent_ui_state("diet", agent_available, insight),
            "context_key": context_key,
        },
        "intake_chart": {
            "label": "Calories Taken",
            "unit": "kcal",
            "points": intake_points,
            "point_count": len(intake_points),
            "latest_value_text": f"{latest_intake_value:.0f} kcal" if latest_intake_value is not None else "No data yet",
        },
        "today_label": datetime.now().strftime("%A, %d %b %Y"),
        "meals_per_day": normalize_meals_per_day(profile.get("preferred_meals_per_day"), DEFAULT_DIET_MEALS_PER_DAY),
        "needs_profile_setup": not profile.get("current_weight_kg"),
        "status_label": GYM_PROFILE_GOALS.get(profile.get("goal", "recomp"), "Diet plan"),
        "active_preference_labels": plan.get("active_preference_labels", []),
        "calorie_card": calorie_card,
        "nutrient_cards": nutrient_cards,
        "history_panel": {
            "label": "Intake history",
            "unit": "kcal",
            "points": history_points,
            "point_count": len(history_points),
            "latest_value_text": f"{latest_intake_value:.0f} kcal" if latest_intake_value is not None else "No data yet",
            "default_range": "month",
        },
    }


def build_gym_plan(
    username: str,
    profile: dict,
    sessions: list[dict],
    templates: list[dict],
    rules: dict,
    override_state: dict | None = None,
    *,
    selected_session: dict | None = None,
):
    override_state = override_state or {}
    template_map = {item["split_key"]: item for item in templates}
    today = datetime.now().date()
    today_iso = today.isoformat()
    preferred_minutes = override_state.get("today_session_minutes") or profile.get("preferred_session_minutes") or DEFAULT_GYM_PREFERRED_MINUTES
    plan_split_key = override_state.get("today_split") or override_state.get("long_term_split") or rules.get("default_start_split", "pull")
    performed_on = today.isoformat()
    rationale = []
    override_notes = list(override_state.get("notes") or [])
    heading = "Current session"
    progress = {
        "total_count": 0,
        "done_count": 0,
        "skipped_count": 0,
        "resolved_count": 0,
        "is_complete": False,
        "has_activity": False,
    }
    force_today_session = bool(override_state.get("force_today_session"))
    if selected_session:
        plan_split_key = normalize_split_key(selected_session.get("split_key"))
        performed_on = selected_session.get("performed_on") or performed_on
        preferred_minutes = safe_int(selected_session.get("duration_minutes"), preferred_minutes) or preferred_minutes
        heading = "Session history"
        rationale.append(
            f"Viewing saved {split_label(plan_split_key)} session from {performed_on}."
        )
        rationale.append("You can switch back to Today at any time from the history strip.")
        force_today_session = False
    elif sessions:
        last_session = sessions[0]
        last_date = datetime.fromisoformat(last_session["performed_on"]).date()
        days_since = (today - last_date).days
        last_split = last_session["split_key"]
        last_template = template_map.get(last_split)
        last_exercise_names = [row["exercise_name"] for row in (last_template or {}).get("exercises", [])[:6]]
        last_entries = current_session_logged_weights(username, last_split, last_session["performed_on"])
        last_progress = gym_session_progress(last_exercise_names, last_entries)
        rationale.append(f"Last logged session: {split_label(last_split)} on {last_session['performed_on']}.")
        if last_progress["has_activity"] and not last_progress["is_complete"] and not force_today_session:
            plan_split_key = last_split
            performed_on = last_session["performed_on"]
            progress = last_progress
            rationale.append("This session is still open. Mark every movement done or skipped before the day is finished.")
        elif days_since <= 0 and last_progress["is_complete"] and not override_state.get("today_split"):
            if override_state.get("long_term_split"):
                plan_split_key = override_state["long_term_split"]
                rationale.append("Today's previous session is finished. The regular coach-selected split is ready if you want another session today.")
            elif last_split == "legs":
                plan_split_key = rules.get("legs_next_split", "pull")
                rationale.append("Today's previous session is finished. The planner now opens the next recovery-friendly split for today.")
            else:
                plan_split_key = template_map.get(last_split, {}).get("next_split_key") or rules.get("default_start_split", "pull")
                rationale.append("Today's previous session is finished. The next split is ready if you want another session today.")
            performed_on = today_iso
            progress = {
                "total_count": 0,
                "done_count": 0,
                "skipped_count": 0,
                "resolved_count": 0,
                "is_complete": False,
                "has_activity": False,
            }
        elif override_state.get("today_split"):
            plan_split_key = override_state["today_split"]
            performed_on = today_iso
            rationale.append("A coach change is forcing a custom split for today.")
        elif days_since >= safe_int(rules.get("reentry_after_days"), 4):
            plan_split_key = rules.get("reentry_split", "full_body")
            rationale.append("Training gap is long enough that a reentry session is safer than jumping back into full intensity.")
        elif override_state.get("long_term_split"):
            plan_split_key = override_state["long_term_split"]
            rationale.append("The regular plan currently uses the coach-selected split.")
        elif last_split == "legs":
            plan_split_key = rules.get("legs_next_split", "pull")
            rationale.append("After a leg-heavy day, the plan moves to pull work to keep recovery balanced.")
        else:
            plan_split_key = template_map.get(last_split, {}).get("next_split_key") or rules.get("default_start_split", "pull")
            rationale.append("The planner advances to the next split in the current rotation.")
    else:
        if override_state.get("today_split"):
            plan_split_key = override_state["today_split"]
            rationale.append("No workout history is stored yet, so today's coach-selected split is used.")
        elif override_state.get("long_term_split"):
            plan_split_key = override_state["long_term_split"]
            rationale.append("No workout history is stored yet, so the regular coach-selected split is used.")
        else:
            rationale.append("No workout history is stored yet, so the planner starts from the default split.")

    template = template_map.get(plan_split_key) or next(iter(template_map.values()), None)
    if not template:
        return {
            "heading": heading,
            "title": "No program template loaded",
            "duration_minutes": preferred_minutes,
            "duration_text": f"{preferred_minutes} min",
            "rationale": rationale,
            "exercise_rows": [],
            "equipment_summary": "No exercise templates are available yet.",
            "source_name": "",
            "performed_on": performed_on,
            "body_weight_kg": safe_float(selected_session.get("body_weight_kg")) if selected_session else safe_float(profile.get("current_weight_kg")),
            "progress": progress,
            "status_label": "No plan",
            "status_class": "stopped",
            "progress_text": "No movements loaded",
            "override_notes": override_notes,
        }

    equipment = sorted({row["equipment_type"] for row in template["exercises"] if row.get("equipment_type")})
    exercise_rows = []
    for row in template["exercises"][:6]:
        exercise_rows.append(
            {
                **row,
                "suggested_weight_kg": parse_weight_kg_from_text(row.get("suggested_weight_text")),
            }
        )
    override_notes.extend(apply_gym_exercise_override_events(exercise_rows, override_state.get("long_term_events", []), note_prefix="Regular plan"))
    override_notes.extend(apply_gym_exercise_override_events(exercise_rows, override_state.get("today_events", []), note_prefix="Today only"))
    if not progress["total_count"]:
        progress = gym_session_progress([row["exercise_name"] for row in exercise_rows], current_session_logged_weights(username, template["split_key"], performed_on))
    rationale.append(f"Session length is capped at about {preferred_minutes} minutes.")
    progress_text = f"{progress['resolved_count']}/{progress['total_count']} resolved"
    if progress["is_complete"]:
        status_label = "Finished"
        status_class = "running"
    elif progress["resolved_count"] > 0:
        status_label = progress_text
        status_class = "partial"
    else:
        status_label = "Open"
        status_class = "partial"
    return {
        "heading": heading,
        "title": f"{template['title']} day",
        "split_key": template["split_key"],
        "performed_on": performed_on,
        "body_weight_kg": safe_float(selected_session.get("body_weight_kg")) if selected_session else safe_float(profile.get("current_weight_kg")),
        "duration_minutes": preferred_minutes,
        "duration_text": f"{preferred_minutes} min",
        "rationale": rationale,
        "exercise_rows": exercise_rows,
        "equipment_summary": ", ".join(item.replace("_", " ").title() for item in equipment) or "Mixed equipment",
        "source_name": template["source_name"],
        "progress": progress,
        "progress_text": progress_text,
        "status_label": status_label,
        "status_class": status_class,
        "override_notes": override_notes,
    }


def gym_agent_context_key(plan: dict):
    split_key = plan.get("split_key") or "current"
    performed_on = plan.get("performed_on") or datetime.now().date().isoformat()
    return f"{split_key}:{performed_on}"


def _legacy_build_gym_agent_prompt(profile: dict, plan: dict, sessions: list[dict]):
    recent_sessions = [
        {
            "performed_on": item.get("performed_on"),
            "split_label": item.get("split_label"),
            "duration_minutes": item.get("duration_minutes"),
            "body_weight_kg": item.get("body_weight_kg"),
        }
        for item in sessions[:3]
    ]
    exercise_items = []
    for item in plan.get("exercise_rows", []):
        exercise_items.append(
            {
                "exercise_name": item.get("exercise_name"),
                "sets_text": item.get("sets_text"),
                "reps_text": item.get("reps_text"),
                "rest_text": item.get("rest_text"),
                "equipment_type": item.get("equipment_type"),
                "suggested_weight_kg": item.get("suggested_weight_kg"),
                "current_weight_kg": safe_float(item.get("input_weight_kg")),
                "status": item.get("status"),
                "latest_history": item.get("history_latest_text"),
            }
        )
    payload = {
        "profile": {
            "display_name": profile.get("display_name"),
            "goal": profile.get("goal_label"),
            "gender": profile.get("gender_label"),
            "current_weight_kg": profile.get("current_weight_kg"),
            "target_weight_kg": profile.get("desired_weight_kg"),
            "preferred_session_minutes": profile.get("preferred_session_minutes"),
            "training_days_per_week": profile.get("training_days_per_week"),
        },
        "current_session": {
            "title": plan.get("title"),
            "split_key": plan.get("split_key"),
            "performed_on": plan.get("performed_on"),
            "progress_text": plan.get("progress_text"),
            "duration_text": plan.get("duration_text"),
            "rationale": plan.get("rationale", []),
            "exercises": exercise_items,
        },
        "recent_sessions": recent_sessions,
    }
    return (
        "You are the Gym AI Coach for Health Hub.\n"
        "Return JSON only with keys headline, bullets, watchout.\n"
        "Rules:\n"
        "- headline: one short sentence\n"
        "- bullets: exactly 3 short strings\n"
        "- watchout: one short sentence\n"
        "- focus on the current session only\n"
        "- be practical about progression, technique focus, and fatigue management\n"
        "- use only the provided data\n"
        "- do not ask questions\n"
        "- do not include markdown fences or extra keys\n\n"
        f"Context JSON:\n{json.dumps(payload, ensure_ascii=True)}"
    )


def build_gym_state(username: str, agent_available: bool = False, selected_day: str | None = None):
    profile = load_gym_profile(username)
    templates = list_gym_templates()
    rules = list_gym_rules()
    sessions = recent_gym_sessions(username, limit=24)
    user_id = ensure_app_user(username)
    normalized_selected_day = normalize_gym_day(selected_day)
    today_iso = datetime.now().date().isoformat()
    selected_session = None
    with gym_user_connection() as connection:
        if normalized_selected_day != "today":
            selected_session = load_latest_gym_session_for_day(connection, user_id, normalized_selected_day)
            if not selected_session:
                normalized_selected_day = "today"
        effective_plan_date = (selected_session or {}).get("performed_on") or today_iso
        override_state = collect_gym_override_state_for_user(connection, user_id, effective_plan_date)
    plan = build_gym_plan(
        username,
        profile,
        sessions,
        templates,
        rules,
        override_state=override_state,
        selected_session=selected_session,
    )
    history_map = exercise_weight_history(username, [item["exercise_name"] for item in plan.get("exercise_rows", [])], limit_per_exercise=10)
    current_logged = current_session_logged_weights(username, plan.get("split_key", ""), plan.get("performed_on", datetime.now().date().isoformat()))
    exercise_rows = []
    for exercise in plan.get("exercise_rows", []):
        media = resolve_exercise_media(exercise["exercise_name"])
        current_entry = current_logged.get(exercise["exercise_name"]) or {}
        input_value = format_weight_kg(current_entry.get("weight_kg")) or format_weight_kg(exercise.get("suggested_weight_kg"))
        current_status = normalize_gym_exercise_status(current_entry.get("status"))
        history_points = history_map.get(exercise["exercise_name"], [])
        latest_history_value = history_points[-1]["value"] if history_points else None
        exercise_rows.append(
            {
                **exercise,
                "input_weight_kg": input_value,
                "status": current_status,
                "is_saved": current_status == "done",
                "is_skipped": current_status == "skipped",
                "history_points": history_points,
                "history_point_count": len(history_points),
                "history_latest_text": f"{latest_history_value:.1f} kg" if latest_history_value is not None else "No trend yet",
                "image_urls": media.get("image_urls", []),
                "instructions": media.get("instructions", []),
                "image_source_name": media.get("name", ""),
                "image_source_url": media.get("source_url", ""),
            }
        )
    plan["exercise_rows"] = exercise_rows
    plan["burn_summary"] = estimate_gym_session_burn(
        profile,
        exercise_rows,
        plan.get("duration_minutes"),
        body_weight_kg=plan.get("body_weight_kg"),
    )
    profile_missing_core = not profile.get("height_cm") or not profile.get("current_weight_kg")
    planner_item = {
        "value": "today",
        "kicker": "Planner",
        "label": "Next Session",
        "meta": plan.get("title") or "Current planner",
        "is_active": normalized_selected_day == "today",
        "is_planner": True,
    }
    history_day_items = []
    seen_days = set()
    for item in sessions:
        performed_on = item.get("performed_on") or ""
        if not performed_on or performed_on in seen_days:
            continue
        seen_days.add(performed_on)
        date_value = datetime.fromisoformat(performed_on)
        history_day_items.append(
            {
                "value": performed_on,
                "kicker": "Today" if performed_on == today_iso else date_value.strftime("%a"),
                "label": date_value.strftime("%d %b"),
                "meta": item.get("split_label") or split_label(item.get("split_key") or ""),
                "is_active": normalized_selected_day == performed_on,
                "is_planner": False,
            }
        )
        if len(history_day_items) >= 7:
            break
    return {
        "profile": profile,
        "goal_options": [{"value": key, "label": label} for key, label in GYM_PROFILE_GOALS.items()],
        "gender_options": [{"value": key, "label": label} for key, label in HEALTH_GENDER_OPTIONS.items()],
        "recommended_target": suggested_weight_targets(profile),
        "bmi": bmi_for_profile(profile),
        "plan": plan,
        "session_nav": {
            "selected_day": normalized_selected_day,
            "selected_label": "Next session" if normalized_selected_day == "today" else datetime.fromisoformat(normalized_selected_day).strftime("%d %b %Y"),
            "planner_item": planner_item,
            "items": history_day_items,
            "history_available": bool(history_day_items),
            "history_count": len(history_day_items),
        },
        "recent_sessions": sessions,
        "needs_profile_setup": profile_missing_core,
        "override_notes": plan.get("override_notes", []),
    }


def build_health_state(username: str):
    profile = load_gym_profile(username)
    recommended_target = suggested_weight_targets(profile)
    measurements = recent_body_measurements(username, limit=48)
    latest_measurement = measurements[0] if measurements else None
    ordered_measurements = sorted(
        measurements,
        key=lambda item: ((item.get("measured_on") or ""), safe_int(item.get("id"), 0) or 0),
    )
    weight_field = next((field for field in HEALTH_MEASUREMENT_FIELDS if field["key"] == "weight_kg"), None)
    charts = [build_body_measurement_chart(field, ordered_measurements) for field in HEALTH_MEASUREMENT_FIELDS if field["key"] != "weight_kg"]
    weight_chart = build_body_measurement_chart(weight_field, ordered_measurements) if weight_field else None
    return {
        "profile": profile,
        "goal_options": [{"value": key, "label": label} for key, label in GYM_PROFILE_GOALS.items()],
        "gender_options": [{"value": key, "label": label} for key, label in HEALTH_GENDER_OPTIONS.items()],
        "meal_count_options": [{"value": value, "label": f"{value} meals"} for value in DIET_MEAL_COUNT_OPTIONS],
        "recommended_target": recommended_target,
        "bmi": bmi_for_profile(profile),
        "measurements": measurements,
        "latest_measurement": latest_measurement,
        "charts": charts,
        "weight_chart": weight_chart,
        "today_iso": datetime.now().date().isoformat(),
        "needs_profile_setup": not profile.get("height_cm") or not profile.get("current_weight_kg"),
    }


def build_assistant_prompt(username: str, user_message: str):
    health = build_health_state(username)
    diet = build_diet_state(username, agent_available=False)
    gym = build_gym_state(username, agent_available=False)
    templates = list_gym_templates()
    diet_foods = list_diet_food_library_items(limit=24)
    history = load_assistant_messages(username, limit=8, include_attachment_text=True)
    context = {
        "user": {
            "display_name": health["profile"]["display_name"],
            "height_cm": health["profile"]["height_cm"],
            "age_years": health["profile"]["age_years"],
            "current_weight_kg": health["profile"]["current_weight_kg"],
            "desired_weight_kg": health["profile"]["desired_weight_kg"],
            "goal": health["profile"]["goal"],
            "goal_label": health["profile"]["goal_label"],
            "gender": health["profile"]["gender"],
            "meals_per_day": health["profile"]["preferred_meals_per_day"],
            "training_days_per_week": health["profile"]["training_days_per_week"],
            "preferred_session_minutes": health["profile"]["preferred_session_minutes"],
            "notes": truncate_text(health["profile"].get("notes"), 240),
        },
        "latest_health_checkin": health.get("latest_measurement"),
        "health_trends": {
            "bmi": health.get("bmi"),
            "recommended_target_weight_kg": health["recommended_target"].get("recommended_weight_kg"),
        },
        "diet_today": {
            "plan_date": diet["plan"]["plan_date"],
            "target_calories": diet["plan"]["target_calories"],
            "planned_calories": diet["plan"]["total_calories"],
            "consumed_calories": diet["plan"]["progress"]["consumed_calories"],
            "meal_progress": diet["plan"]["progress_text"],
            "active_preferences": diet["plan"].get("active_preference_labels", []),
            "override_notes": diet["plan"].get("override_notes", []),
            "meals": [
                {
                    "label": item["meal_label"],
                    "title": item["meal_title"],
                    "status": item["status"],
                    "calories": item["calories"],
                }
                for item in diet["plan"]["meals"]
            ],
        },
        "diet_food_library": [
            {
                "food_name": item.get("label"),
                "serving_text": item.get("serving_text"),
                "calories": item.get("calories"),
                "nutrients": item.get("nutrients", {}),
            }
            for item in diet_foods
        ],
        "gym_today": {
            "title": gym["plan"]["title"],
            "split_key": gym["plan"].get("split_key"),
            "performed_on": gym["plan"].get("performed_on"),
            "progress_text": gym["plan"]["progress_text"],
            "override_notes": gym["plan"].get("override_notes", []),
            "exercises": [
                {
                    "name": item["exercise_name"],
                    "status": item["status"],
                    "suggested_weight_kg": item.get("suggested_weight_kg"),
                    "current_weight_kg": safe_float(item.get("input_weight_kg")),
                }
                for item in gym["plan"]["exercise_rows"]
            ],
        },
        "gym_program": [
            {
                "split_key": item.get("split_key"),
                "title": item.get("title"),
                "exercises": [row.get("exercise_name") for row in item.get("exercises", [])[:6]],
            }
            for item in templates[:6]
        ],
        "recent_chat": [
            {
                "role": item["role"],
                "content": item["content_text"],
                "attachments": [
                    {
                        "name": attachment["original_name"],
                        "kind": attachment["kind_label"],
                        "size": attachment["size_text"],
                        "is_image": bool(attachment.get("is_image")),
                        "preview_text": truncate_text(attachment.get("preview_text", ""), 180),
                        "analysis_text": truncate_text(attachment.get("analysis_text", ""), 800),
                        "analysis_summary": truncate_text(attachment.get("analysis_summary", ""), 240),
                        "analysis_items": [
                            {
                                "food_name": truncate_text(entry.get("food_name"), 120),
                                "servings": safe_float(entry.get("servings")),
                                "serving_text": truncate_text(entry.get("serving_text"), 80),
                                "calories": safe_float(entry.get("calories")),
                                "nutrients": normalize_diet_food_nutrients(entry.get("nutrients")),
                            }
                            for entry in attachment.get("analysis_items", [])[:6]
                            if isinstance(entry, dict)
                        ],
                    }
                    for attachment in item.get("attachments", [])[:2]
                ],
            }
            for item in history[-6:]
        ],
    }
    return build_assistant_prompt_from_context(context, user_message)


def generate_assistant_reply(username: str, user_message: str, attachments: list[dict] | None = None):
    prepared_attachments = attachments or []
    clean_message = truncate_text(user_message, 1200)
    if not clean_message and not prepared_attachments:
        raise ValueError("Enter a message for the coach.")
    if not clean_message and prepared_attachments:
        clean_message = "Please review the attached file(s) and help based on them."
    if prepared_attachments:
        enrich_assistant_image_attachments(prepared_attachments, clean_message)
    if not agent_host_online():
        raise RuntimeError("MSI is offline. The coach is unavailable right now, but your saved plans and data are still here.")
    user_message_id = save_assistant_message(
        username,
        "user",
        clean_message,
        metadata={"attachment_count": len(prepared_attachments)},
    )
    if prepared_attachments:
        save_assistant_attachments(username, user_message_id, prepared_attachments)
    response_text = remote_codex_exec(build_assistant_prompt(username, clean_message))
    parsed = parse_assistant_response_text(response_text)
    assistant_message_id = save_assistant_message(
        username,
        "assistant",
        parsed["reply"],
        metadata={"source_label": "Codex on MSI"},
    )
    save_assistant_actions(username, assistant_message_id, parsed["actions"])
    return {
        "user_message_id": user_message_id,
        "assistant_message_id": assistant_message_id,
        "reply": parsed["reply"],
        "actions": parsed["actions"],
    }


def assistant_action_refresh_tabs(action_type: str):
    mapping = {
        "record_weight_checkin": ["health", "diet", "gym"],
        "record_health_measurement": ["health", "diet", "gym"],
        "update_height_cm": ["health", "diet", "gym"],
        "update_age_years": ["health", "diet", "gym"],
        "update_target_weight": ["health", "diet", "gym"],
        "update_gender": ["health", "diet", "gym"],
        "update_goal": ["health", "diet", "gym"],
        "update_training_days_per_week": ["health", "diet", "gym"],
        "update_session_minutes": ["health", "gym"],
        "update_meals_per_day": ["health", "diet"],
        "customize_diet_meal": ["diet"],
        "set_diet_preferences": ["diet"],
        "set_gym_split": ["gym"],
        "replace_gym_exercise": ["gym"],
        "update_gym_exercise_scheme": ["gym"],
        "update_gym_exercise_load": ["gym"],
        "refresh_diet_plan": ["diet"],
    }
    return mapping.get(str(action_type or "").strip(), [])


def apply_assistant_action(username: str, action_id: int, decision: str):
    user_id = ensure_app_user(username)
    with gym_user_connection() as connection:
        row = connection.execute(
            """
            SELECT *
            FROM assistant_actions
            WHERE id = ? AND user_id = ?
            ORDER BY id DESC LIMIT 1
            """,
            (action_id, user_id),
        ).fetchone()
        if not row:
            raise ValueError("The requested action was not found.")
        if row["status"] != "pending":
            raise ValueError("This action has already been resolved.")
        action_type = row["action_type"]
        try:
            payload = json.loads(row["payload_json"] or "{}")
        except json.JSONDecodeError:
            payload = {}
        timestamp = now_iso()
        if decision == "cancel":
            connection.execute(
                """
                UPDATE assistant_actions
                SET status = 'cancelled', resolution_note = ?, updated_at = ?
                WHERE id = ?
                """,
                ("Cancelled by user.", timestamp, row["id"]),
            )
            return {"message": "Action cancelled.", "refresh_tabs": []}
        if decision != "confirm":
            raise ValueError("Unknown action decision.")

        resolution_note = ""
        refresh_tabs = assistant_action_refresh_tabs(action_type)
        if action_type == "record_weight_checkin":
            weight_kg = safe_float(payload.get("weight_kg"))
            measured_on = str(payload.get("measured_on") or datetime.now().date().isoformat()).strip()
            note = truncate_text(payload.get("note"), 200)
            connection.execute(
                """
                INSERT INTO body_measurements (
                    user_id, measured_on, weight_kg, arm_cm, chest_cm, waist_cm, thigh_cm, calf_cm, note, source, created_at
                ) VALUES (?, ?, ?, NULL, NULL, NULL, NULL, NULL, ?, ?, ?)
                """,
                (user_id, measured_on, weight_kg, note, "assistant", timestamp),
            )
            sync_profile_current_weight_from_measurements(connection, user_id)
            clear_coach_insights(connection, user_id, agent_keys=("diet", "gym"))
            resolution_note = f"Recorded weight {weight_kg:.1f} kg for {measured_on}."
        elif action_type == "record_health_measurement":
            measured_on = str(payload.get("measured_on") or datetime.now().date().isoformat()).strip()
            weight_kg = safe_float(payload.get("weight_kg"))
            arm_cm = safe_float(payload.get("arm_cm"))
            chest_cm = safe_float(payload.get("chest_cm"))
            waist_cm = safe_float(payload.get("waist_cm"))
            thigh_cm = safe_float(payload.get("thigh_cm"))
            calf_cm = safe_float(payload.get("calf_cm"))
            note = truncate_text(payload.get("note"), 200)
            connection.execute(
                """
                INSERT INTO body_measurements (
                    user_id, measured_on, weight_kg, arm_cm, chest_cm, waist_cm, thigh_cm, calf_cm, note, source, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (user_id, measured_on, weight_kg, arm_cm, chest_cm, waist_cm, thigh_cm, calf_cm, note, "assistant", timestamp),
            )
            if weight_kg is not None:
                sync_profile_current_weight_from_measurements(connection, user_id)
            clear_coach_insights(connection, user_id, agent_keys=("diet", "gym"))
            summary_text = measurement_summary_text(payload)
            resolution_note = f"Recorded health check-in for {measured_on}{': ' + summary_text if summary_text else ''}."
        elif action_type == "update_height_cm":
            height_cm = safe_float(payload.get("height_cm"))
            if height_cm is None or height_cm < 100 or height_cm > 260:
                raise ValueError("The requested height is not supported.")
            connection.execute(
                "UPDATE gym_profiles SET height_cm = ?, updated_at = ? WHERE user_id = ?",
                (height_cm, timestamp, user_id),
            )
            clear_coach_insights(connection, user_id, agent_keys=("diet", "gym"))
            resolution_note = f"Height changed to {height_cm:.1f} cm."
        elif action_type == "update_age_years":
            age_years = normalize_age_years(payload.get("age_years"))
            if age_years is None:
                raise ValueError("The requested age is not supported.")
            connection.execute(
                "UPDATE gym_profiles SET age_years = ?, updated_at = ? WHERE user_id = ?",
                (age_years, timestamp, user_id),
            )
            clear_coach_insights(connection, user_id, agent_keys=("diet", "gym"))
            resolution_note = f"Age changed to {age_years:.0f} years."
        elif action_type == "update_target_weight":
            target_weight_kg = safe_float(payload.get("target_weight_kg"))
            if target_weight_kg is None or target_weight_kg <= 0 or target_weight_kg > 400:
                raise ValueError("The requested target weight is not supported.")
            connection.execute(
                "UPDATE gym_profiles SET desired_weight_kg = ?, updated_at = ? WHERE user_id = ?",
                (target_weight_kg, timestamp, user_id),
            )
            clear_coach_insights(connection, user_id, agent_keys=("diet", "gym"))
            resolution_note = f"Target weight changed to {target_weight_kg:.1f} kg."
        elif action_type == "update_gender":
            gender = normalize_gender(payload.get("gender"))
            if gender not in HEALTH_GENDER_OPTIONS:
                raise ValueError("The requested gender is not supported.")
            connection.execute(
                "UPDATE gym_profiles SET gender = ?, updated_at = ? WHERE user_id = ?",
                (gender, timestamp, user_id),
            )
            clear_coach_insights(connection, user_id, agent_keys=("diet", "gym"))
            resolution_note = f"Gender changed to {HEALTH_GENDER_OPTIONS[gender]}."
        elif action_type == "update_goal":
            goal = str(payload.get("goal") or "").strip()
            if goal not in GYM_PROFILE_GOALS:
                raise ValueError("The requested goal is not supported.")
            connection.execute(
                "UPDATE gym_profiles SET goal = ?, updated_at = ? WHERE user_id = ?",
                (goal, timestamp, user_id),
            )
            clear_coach_insights(connection, user_id, agent_keys=("diet", "gym"))
            resolution_note = f"Goal changed to {GYM_PROFILE_GOALS[goal]}."
        elif action_type == "update_training_days_per_week":
            training_days = safe_int(payload.get("training_days_per_week"), DEFAULT_GYM_DAYS_PER_WEEK) or DEFAULT_GYM_DAYS_PER_WEEK
            training_days = max(1, min(7, training_days))
            connection.execute(
                "UPDATE gym_profiles SET training_days_per_week = ?, updated_at = ? WHERE user_id = ?",
                (training_days, timestamp, user_id),
            )
            clear_coach_insights(connection, user_id, agent_keys=("diet", "gym"))
            resolution_note = f"Training days per week changed to {training_days}."
        elif action_type == "update_session_minutes":
            scope = normalize_coach_scope(payload.get("scope")) or "long_term"
            session_minutes = safe_int(payload.get("preferred_session_minutes"), DEFAULT_GYM_PREFERRED_MINUTES) or DEFAULT_GYM_PREFERRED_MINUTES
            session_minutes = max(20, min(180, session_minutes))
            if scope == "today":
                effective_date = str(payload.get("effective_date") or datetime.now().date().isoformat()).strip()
                append_plan_override(
                    connection,
                    user_id,
                    "gym",
                    "today",
                    "update_session_minutes",
                    {"preferred_session_minutes": session_minutes},
                    effective_date=effective_date,
                )
            else:
                connection.execute(
                    "UPDATE gym_profiles SET preferred_session_minutes = ?, updated_at = ? WHERE user_id = ?",
                    (session_minutes, timestamp, user_id),
                )
            clear_coach_insights(connection, user_id, agent_keys=("gym",))
            resolution_note = (
                f"Session length changed to {session_minutes} minutes for today."
                if scope == "today"
                else f"Session length changed to {session_minutes} minutes."
            )
        elif action_type == "update_meals_per_day":
            scope = normalize_coach_scope(payload.get("scope")) or "long_term"
            meals_per_day = normalize_meals_per_day(payload.get("meals_per_day"), DEFAULT_DIET_MEALS_PER_DAY)
            if scope == "today":
                plan_date = str(payload.get("effective_date") or datetime.now().date().isoformat()).strip()
                append_plan_override(
                    connection,
                    user_id,
                    "diet",
                    "today",
                    "update_meals_per_day",
                    {"meals_per_day": meals_per_day},
                    effective_date=plan_date,
                )
            else:
                connection.execute(
                    "UPDATE gym_profiles SET preferred_meals_per_day = ?, updated_at = ? WHERE user_id = ?",
                    (meals_per_day, timestamp, user_id),
                )
            clear_coach_insights(connection, user_id, agent_keys=("diet",))
            resolution_note = (
                f"Meals per day changed to {meals_per_day} for today."
                if scope == "today"
                else f"Meals per day changed to {meals_per_day}."
            )
        elif action_type == "customize_diet_meal":
            scope = normalize_coach_scope(payload.get("scope")) or "today"
            slot_key = normalize_diet_slot_key(payload.get("slot_key"))
            mode = normalize_diet_meal_customization_mode(payload.get("mode")) or "append"
            requested_items = normalize_diet_meal_override_items(payload)
            if not slot_key or not requested_items:
                raise ValueError("The requested meal customization is not supported.")
            effective_date = str(payload.get("effective_date") or datetime.now().date().isoformat()).strip() if scope == "today" else None
            unresolved_items = []
            prepared_items = []
            for requested_item in requested_items:
                item_spec = build_diet_custom_food_item(requested_item)
                if item_spec:
                    prepared_items.append((requested_item, item_spec))
                    continue
                unresolved_items.append(requested_item)
            if unresolved_items:
                estimated_items = estimate_diet_foods_with_agent(unresolved_items)
                for requested_item in unresolved_items:
                    normalized_food_key = normalize_diet_food_text(requested_item.get("food_name"))
                    estimated = estimated_items.get(normalized_food_key)
                    if estimated:
                        merged_item = {
                            **requested_item,
                            "food_name": estimated.get("food_name") or requested_item.get("food_name"),
                            "serving_text": estimated.get("serving_text") or requested_item.get("serving_text"),
                            "calories": estimated.get("calories"),
                            "nutrients": estimated.get("nutrients") or requested_item.get("nutrients"),
                        }
                        item_spec = build_diet_custom_food_item(merged_item)
                        if item_spec:
                            prepared_items.append((merged_item, item_spec))
                            continue
                    prepared_items.append((requested_item, None))
            stored_items = []
            failed_items = []
            for requested_item, item_spec in prepared_items:
                if not item_spec:
                    failed_items.append(str(requested_item.get("food_name") or "").strip())
                    continue
                raw_food_name = str(requested_item.get("food_name") or "").strip()
                food_record = upsert_diet_food_library_item(
                    {
                        "food_name": item_spec["food_name"],
                        "serving_text": item_spec["serving_text"],
                        "calories": item_spec["calories_per_serving"],
                        "nutrients": item_spec["per_serving_nutrients"],
                        "aliases": [raw_food_name] if raw_food_name else [],
                    },
                    source_label="Health Hub Coach",
                    source_kind="assistant_generated" if (requested_item.get("calories") is not None or requested_item.get("serving_text") or requested_item.get("nutrients")) else "assistant",
                )
                stored_items.append(
                    {
                        "food_name": (food_record or {}).get("label") or item_spec["food_name"],
                        "serving_text": (food_record or {}).get("serving_text") or item_spec["serving_text"],
                        "servings": item_spec["servings"],
                        "calories": (food_record or {}).get("calories", item_spec["calories_per_serving"]),
                        "nutrients": (food_record or {}).get("nutrients") or item_spec["per_serving_nutrients"],
                    }
                )
            if failed_items:
                missing_text = ", ".join(item for item in failed_items if item)
                raise ValueError(f"Could not estimate nutrition for: {missing_text}. Try again and I will rebuild the meal with exact items.")
            if not stored_items:
                raise ValueError("The requested meal customization is incomplete.")
            append_plan_override(
                connection,
                user_id,
                "diet",
                scope,
                "customize_diet_meal",
                {
                    "slot_key": slot_key,
                    "mode": mode,
                    "items": stored_items,
                },
                effective_date=effective_date,
            )
            clear_coach_insights(connection, user_id, agent_keys=("diet",), context_key=effective_date or datetime.now().date().isoformat())
            slot_label_text = diet_slot_label(slot_key)
            item_summary = diet_override_items_summary_text(stored_items)
            action_text = "now includes" if mode == "append" else "is now rebuilt around"
            resolution_note = (
                f"{slot_label_text} {action_text} {item_summary.lower()} for today."
                if scope == "today"
                else f"Regular {slot_label_text.lower()} {action_text} {item_summary.lower()}."
            )
        elif action_type == "set_diet_preferences":
            scope = normalize_coach_scope(payload.get("scope")) or "long_term"
            preference_keys = normalize_diet_preference_keys(payload.get("preference_keys"))
            if not preference_keys:
                raise ValueError("The requested diet preference is not supported.")
            effective_date = str(payload.get("effective_date") or datetime.now().date().isoformat()).strip() if scope == "today" else None
            append_plan_override(
                connection,
                user_id,
                "diet",
                scope,
                "set_diet_preferences",
                {"preference_keys": preference_keys},
                effective_date=effective_date,
            )
            clear_coach_insights(connection, user_id, agent_keys=("diet",))
            resolution_note = (
                f"Diet focus changed to {', '.join(diet_preference_labels(preference_keys))} for today."
                if scope == "today"
                else f"Regular diet focus changed to {', '.join(diet_preference_labels(preference_keys))}."
            )
        elif action_type == "set_gym_split":
            scope = normalize_coach_scope(payload.get("scope")) or "long_term"
            split_key = normalize_split_key(payload.get("split_key"))
            if not split_key:
                raise ValueError("The requested split is not supported.")
            effective_date = str(payload.get("effective_date") or datetime.now().date().isoformat()).strip() if scope == "today" else None
            append_plan_override(
                connection,
                user_id,
                "gym",
                scope,
                "set_gym_split",
                {"split_key": split_key},
                effective_date=effective_date,
            )
            clear_coach_insights(connection, user_id, agent_keys=("gym",))
            resolution_note = (
                f"Today's split changed to {split_label(split_key)}."
                if scope == "today"
                else f"Regular split changed to {split_label(split_key)}."
            )
        elif action_type == "replace_gym_exercise":
            scope = normalize_coach_scope(payload.get("scope")) or "long_term"
            split_key = normalize_split_key(payload.get("split_key"))
            exercise_name = str(payload.get("exercise_name") or "").strip()
            replacement_name = str(payload.get("replacement_exercise_name") or "").strip()
            if not split_key or not exercise_name or not replacement_name:
                raise ValueError("The requested exercise change is incomplete.")
            effective_date = str(payload.get("effective_date") or datetime.now().date().isoformat()).strip() if scope == "today" else None
            append_plan_override(
                connection,
                user_id,
                "gym",
                scope,
                "replace_gym_exercise",
                {
                    "split_key": split_key,
                    "exercise_name": exercise_name,
                    "replacement_exercise_name": replacement_name,
                },
                effective_date=effective_date,
            )
            clear_coach_insights(connection, user_id, agent_keys=("gym",))
            resolution_note = (
                f"Today's {split_label(split_key)} plan now uses {replacement_name} instead of {exercise_name}."
                if scope == "today"
                else f"Regular {split_label(split_key)} plan now uses {replacement_name} instead of {exercise_name}."
            )
        elif action_type == "update_gym_exercise_scheme":
            scope = normalize_coach_scope(payload.get("scope")) or "long_term"
            split_key = normalize_split_key(payload.get("split_key"))
            exercise_name = str(payload.get("exercise_name") or "").strip()
            scheme_payload = {
                "split_key": split_key,
                "exercise_name": exercise_name,
                "sets_text": truncate_text(payload.get("sets_text"), 32),
                "reps_text": truncate_text(payload.get("reps_text"), 32),
                "rest_text": truncate_text(payload.get("rest_text"), 32),
            }
            if not split_key or not exercise_name or not any(scheme_payload[key] for key in ("sets_text", "reps_text", "rest_text")):
                raise ValueError("The requested gym scheme change is incomplete.")
            effective_date = str(payload.get("effective_date") or datetime.now().date().isoformat()).strip() if scope == "today" else None
            append_plan_override(
                connection,
                user_id,
                "gym",
                scope,
                "update_gym_exercise_scheme",
                scheme_payload,
                effective_date=effective_date,
            )
            clear_coach_insights(connection, user_id, agent_keys=("gym",))
            resolution_note = (
                f"Updated {exercise_name} scheme for today."
                if scope == "today"
                else f"Updated {exercise_name} scheme in the regular plan."
            )
        elif action_type == "update_gym_exercise_load":
            scope = normalize_coach_scope(payload.get("scope")) or "long_term"
            split_key = normalize_split_key(payload.get("split_key"))
            exercise_name = str(payload.get("exercise_name") or "").strip()
            suggested_weight_kg = safe_float(payload.get("suggested_weight_kg"))
            if not split_key or not exercise_name or suggested_weight_kg is None or suggested_weight_kg < 0:
                raise ValueError("The requested gym load change is incomplete.")
            effective_date = str(payload.get("effective_date") or datetime.now().date().isoformat()).strip() if scope == "today" else None
            append_plan_override(
                connection,
                user_id,
                "gym",
                scope,
                "update_gym_exercise_load",
                {
                    "split_key": split_key,
                    "exercise_name": exercise_name,
                    "suggested_weight_kg": suggested_weight_kg,
                },
                effective_date=effective_date,
            )
            clear_coach_insights(connection, user_id, agent_keys=("gym",))
            resolution_note = (
                f"Set {exercise_name} to {suggested_weight_kg:.1f} kg for today."
                if scope == "today"
                else f"Set the regular suggested load for {exercise_name} to {suggested_weight_kg:.1f} kg."
            )
        elif action_type == "refresh_diet_plan":
            plan_date = str(payload.get("plan_date") or datetime.now().date().isoformat()).strip()
            connection.execute(
                "DELETE FROM diet_daily_plans WHERE user_id = ? AND plan_date = ?",
                (user_id, plan_date),
            )
            clear_coach_insights(connection, user_id, agent_keys=("diet",), context_key=plan_date)
            resolution_note = f"Diet plan for {plan_date} was cleared and updated immediately."
        else:
            raise ValueError("Unsupported action type.")

        connection.execute(
            """
            UPDATE assistant_actions
            SET status = 'applied', resolution_note = ?, updated_at = ?
            WHERE id = ?
            """,
            (resolution_note, timestamp, row["id"]),
        )
        return {"message": resolution_note, "refresh_tabs": refresh_tabs}


def build_assistant_state(username: str, agent_available: bool = False):
    messages = load_assistant_messages(username, limit=24)
    image_analysis_enabled = assistant_image_analysis_enabled()
    return {
        "agent_name": SUPERAGENT_NAME,
        "available": agent_available,
        "status_label": "Ready" if agent_available else "Offline",
        "status_class": "running" if agent_available else "stopped",
        "messages": messages,
        "empty_text": "Ask about body data, today's meals, attached meal photos, or changing your regular gym plan. If scope is unclear, the coach should ask whether you mean today only or the regular plan before proposing a change.",
        "offline_text": "MSI is offline, so the coach chat is paused. Your normal Health Hub data is still available.",
        "attachment_policy": {
            "accept_text": coach_attachment_accept_text(),
            "max_files": COACH_ATTACHMENT_MAX_FILES_PER_MESSAGE,
            "max_file_size_text": format_file_size(COACH_ATTACHMENT_MAX_FILE_BYTES),
            "max_total_size_text": format_file_size(COACH_ATTACHMENT_MAX_TOTAL_BYTES_PER_USER),
            "retention_days": coach_attachment_retention_days(),
            "image_analysis_enabled": image_analysis_enabled,
            "image_hint": (
                "Meal photos are analyzed into food guesses before the coach replies."
                if image_analysis_enabled
                else "Attach text, CSV, JSON, XLSX, or PDF files. Image uploads stay disabled until DASHBOARD_OPENAI_API_KEY is configured for Health Hub on server 106."
            ),
        },
    }


def empty_history():
    return {
        server_id: {metric_name: [] for metric_name in HISTORY_METRICS}
        for server_id in SERVERS
    }


def load_history():
    history = empty_history()
    if not HISTORY_FILE.exists():
        return history
    try:
        payload = json.loads(HISTORY_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return history
    for server_id in SERVERS:
        server_payload = payload.get("servers", {}).get(server_id, {})
        for metric_name in HISTORY_METRICS:
            points = []
            for item in server_payload.get(metric_name, []):
                if not isinstance(item, dict):
                    continue
                try:
                    timestamp = int(item.get("ts"))
                except (TypeError, ValueError):
                    continue
                value = item.get("value")
                if value is not None:
                    try:
                        value = float(value)
                    except (TypeError, ValueError):
                        continue
                points.append({"ts": timestamp, "value": value})
            history[server_id][metric_name] = points
    return history


def save_history(history):
    HISTORY_FILE.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "sample_seconds": SAMPLE_SECONDS,
        "history_seconds": HISTORY_SECONDS,
        "servers": history,
    }
    HISTORY_FILE.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def metric_value(snapshot: dict, metric_name: str):
    if not snapshot.get("reachable"):
        return None
    if metric_name == "cpu_percent":
        return snapshot.get("cpu_percent")
    if metric_name == "memory_percent":
        return snapshot.get("memory", {}).get("percent")
    if metric_name == "disk_percent":
        return snapshot.get("disk", {}).get("percent")
    if metric_name == "temperature_c":
        return snapshot.get("temperature_c")
    return None


def trim_history(history, now_ts: int):
    cutoff = now_ts - HISTORY_SECONDS
    for server_history in history.values():
        for metric_name in HISTORY_METRICS:
            server_history[metric_name] = [
                point for point in server_history.get(metric_name, [])
                if point.get("ts", 0) >= cutoff
            ]


def append_snapshot_history(history, server_id: str, snapshot: dict, timestamp: int):
    for metric_name in HISTORY_METRICS:
        history[server_id][metric_name].append(
            {
                "ts": timestamp,
                "value": metric_value(snapshot, metric_name),
            }
        )


def refresh_state_once():
    global HISTORY_CACHE
    snapshots = {
        "106": collect_local_snapshot(),
        "118": collect_remote_snapshot(),
    }
    updated_at = int(time.time())
    with CACHE_LOCK:
        if not HISTORY_CACHE:
            HISTORY_CACHE = load_history()
        for server_id, snapshot in snapshots.items():
            append_snapshot_history(HISTORY_CACHE, server_id, snapshot, updated_at)
        trim_history(HISTORY_CACHE, updated_at)
        STATE_CACHE["updated_at"] = updated_at
        STATE_CACHE["snapshots"] = snapshots
        save_history(HISTORY_CACHE)


def sampler_loop():
    while True:
        started = time.monotonic()
        try:
            refresh_state_once()
        except Exception:
            pass
        delay = max(5.0, SAMPLE_SECONDS - (time.monotonic() - started))
        time.sleep(delay)


def start_sampler():
    global SAMPLER_STARTED, HISTORY_CACHE
    if SAMPLER_STARTED:
        return
    HISTORY_CACHE = load_history()
    refresh_state_once()
    thread = threading.Thread(target=sampler_loop, name="healthhub-sampler", daemon=True)
    thread.start()
    SAMPLER_STARTED = True


def current_cached_state():
    global HISTORY_CACHE
    with CACHE_LOCK:
        if STATE_CACHE["snapshots"]:
            return (
                STATE_CACHE["updated_at"],
                copy.deepcopy(STATE_CACHE["snapshots"]),
                copy.deepcopy(HISTORY_CACHE or empty_history()),
            )
    if not HISTORY_CACHE:
        HISTORY_CACHE = load_history()
    refresh_state_once()
    with CACHE_LOCK:
        return (
            STATE_CACHE["updated_at"],
            copy.deepcopy(STATE_CACHE["snapshots"]),
            copy.deepcopy(HISTORY_CACHE or empty_history()),
        )


def format_percent_value(value):
    if value is None:
        return "n/a"
    return f"{value:.1f}%"


def format_celsius_value(value):
    if value is None:
        return "n/a"
    return f"{value:.1f} C"


def chart_scale(metric_name: str, values):
    if metric_name in {"cpu_percent", "memory_percent", "disk_percent"}:
        return 0.0, 100.0
    if not values:
        return 0.0, 1.0
    low = min(values)
    high = max(values)
    padding = max(2.0, (high - low) * 0.3)
    low = max(0.0, low - padding)
    high = high + padding
    if high - low < 1.0:
        high = low + 1.0
    return low, high


def graph_point(window_start: int, timestamp: int, value: float, scale_min: float, scale_max: float):
    inner_width = GRAPH_WIDTH - (GRAPH_PADDING * 2)
    inner_height = GRAPH_HEIGHT - (GRAPH_PADDING * 2)
    x = GRAPH_PADDING + ((timestamp - window_start) / max(HISTORY_SECONDS, 1)) * inner_width
    y_ratio = (value - scale_min) / max(scale_max - scale_min, 0.001)
    y = GRAPH_HEIGHT - GRAPH_PADDING - (y_ratio * inner_height)
    return round(x, 2), round(y, 2)


def line_path(points):
    if not points:
        return ""
    return "M " + " L ".join(f"{x} {y}" for x, y in points)


def area_path(points):
    if not points:
        return ""
    base_y = GRAPH_HEIGHT - GRAPH_PADDING
    head = f"M {points[0][0]} {base_y} L {points[0][0]} {points[0][1]}"
    body = " L ".join(f"{x} {y}" for x, y in points[1:])
    tail = f"L {points[-1][0]} {base_y} Z"
    return " ".join(part for part in [head, body, tail] if part)


def build_chart(metric_name: str, history_points, updated_at: int | None):
    window_end = updated_at or int(time.time())
    window_start = window_end - HISTORY_SECONDS
    numeric_points = []
    for item in history_points:
        timestamp = item.get("ts")
        value = item.get("value")
        if timestamp is None or value is None or timestamp < window_start:
            continue
        numeric_points.append((int(timestamp), float(value)))

    values = [value for _, value in numeric_points]
    scale_min, scale_max = chart_scale(metric_name, values)
    line_paths = []
    area_paths = []
    last_x = None
    last_y = None

    if numeric_points:
        segments = []
        current_segment = []
        for timestamp, value in numeric_points:
            if current_segment and timestamp - current_segment[-1][0] > GRAPH_GAP_SECONDS:
                segments.append(current_segment)
                current_segment = []
            current_segment.append((timestamp, value))
        if current_segment:
            segments.append(current_segment)

        for segment in segments:
            coords = [
                graph_point(window_start, timestamp, value, scale_min, scale_max)
                for timestamp, value in segment
            ]
            line_paths.append(line_path(coords))
            area_paths.append(area_path(coords))

        last_x, last_y = graph_point(
            window_start,
            numeric_points[-1][0],
            numeric_points[-1][1],
            scale_min,
            scale_max,
        )

    return {
        "width": GRAPH_WIDTH,
        "height": GRAPH_HEIGHT,
        "mid_y": round(GRAPH_HEIGHT / 2, 2),
        "line_paths": line_paths,
        "area_paths": area_paths,
        "has_data": bool(numeric_points),
        "last_x": last_x,
        "last_y": last_y,
        "sample_count": len(numeric_points),
        "window_label": "Last 10 min",
    }


def build_metric_tiles(snapshot: dict, server_history: dict, updated_at: int | None):
    memory = snapshot.get("memory", {})
    disk = snapshot.get("disk", {})
    metrics = [
        {
            "key": "cpu",
            "css_class": "cpu",
            "label": "CPU Load",
            "value_text": format_percent_value(snapshot.get("cpu_percent")),
            "detail_text": f"{snapshot.get('load_avg', ['n/a', 'n/a', 'n/a'])[0]} / {snapshot.get('load_avg', ['n/a', 'n/a', 'n/a'])[1]} / {snapshot.get('load_avg', ['n/a', 'n/a', 'n/a'])[2]} load avg",
            "chart": build_chart("cpu_percent", server_history.get("cpu_percent", []), updated_at),
        },
        {
            "key": "memory",
            "css_class": "memory",
            "label": "Memory",
            "value_text": format_percent_value(memory.get("percent")),
            "detail_text": f"{memory.get('used_gb', 'n/a')} GB / {memory.get('total_gb', 'n/a')} GB",
            "chart": build_chart("memory_percent", server_history.get("memory_percent", []), updated_at),
        },
        {
            "key": "disk",
            "css_class": "disk",
            "label": "Disk",
            "value_text": format_percent_value(disk.get("percent")),
            "detail_text": f"{disk.get('used_gb', 'n/a')} GB / {disk.get('total_gb', 'n/a')} GB",
            "chart": build_chart("disk_percent", server_history.get("disk_percent", []), updated_at),
        },
        {
            "key": "temperature",
            "css_class": "temperature",
            "label": "CPU Temp",
            "value_text": format_celsius_value(snapshot.get("temperature_c")),
            "detail_text": f"Uptime {uptime_label(snapshot.get('uptime_seconds', 0))}" if snapshot.get("uptime_seconds") is not None else "Uptime n/a",
            "chart": build_chart("temperature_c", server_history.get("temperature_c", []), updated_at),
        },
    ]
    return metrics


def _legacy_dashboard_username() -> str:
    return normalize_username(os.getenv("DASHBOARD_USERNAME", "sam")) or "sam"


def _legacy_dashboard_password() -> str:
    return os.getenv("DASHBOARD_PASSWORD", "")


def _legacy_dashboard_accounts() -> dict[str, dict]:
    primary_username = dashboard_username()
    accounts = {
        primary_username: {
            "password": dashboard_password(),
            "display_name": primary_username.title(),
        },
        "melina": {
            "password": os.getenv("DASHBOARD_MELINA_PASSWORD", "11001100"),
            "display_name": "Melina",
        },
        "nora": {
            "password": os.getenv("DASHBOARD_NORA_PASSWORD", "11001100"),
            "display_name": "Nora",
        },
    }
    return {normalize_username(username): details for username, details in accounts.items() if normalize_username(username)}


def _legacy_display_name_for_username(username: str | None):
    current_username = normalize_username(username)
    if not current_username:
        return ""
    account = dashboard_accounts().get(current_username)
    if account and account.get("display_name"):
        return str(account["display_name"]).strip()
    return current_username.title()


def _legacy_viewer_app_title(username: str | None):
    display_name = display_name_for_username(username)
    if not display_name:
        return APP_TITLE
    return f"{display_name} {APP_TITLE}"


def _legacy_authenticate_dashboard_user(username: str | None, password: str | None):
    current_username = normalize_username(username)
    if not current_username:
        return None
    account = dashboard_accounts().get(current_username)
    if not account:
        return None
    if not hmac.compare_digest(str(password or ""), str(account.get("password", ""))):
        return None
    return current_username


def _legacy_viewer_username() -> str:
    return normalize_username(session.get("username", ""))


def _legacy_dashboard_sections(server_access: bool):
    sections = [
        {"id": "diet", "label": "Diet"},
        {"id": "gym", "label": "Gym"},
        {"id": "health", "label": "Health"},
        {"id": "coach", "label": "Coach"},
    ]
    return sections


def _legacy_resolve_dashboard_tab(server_access: bool, preferred_tab: str | None = None, health_needs_profile_setup: bool = False):
    section_ids = {item["id"] for item in dashboard_sections()}
    requested = (preferred_tab or "").strip().lower()
    if requested in section_ids:
        return requested
    return "health" if health_needs_profile_setup else "diet"


def _legacy_login_required(view):
    @functools.wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("authenticated"):
            return redirect(url_for("login"))
        ensure_csrf_token()
        return view(*args, **kwargs)

    return wrapped


def _legacy_ensure_csrf_token() -> str:
    token = session.get("csrf_token")
    if not token:
        token = secrets.token_hex(24)
        session["csrf_token"] = token
    return token


def _legacy_verify_csrf() -> bool:
    return hmac.compare_digest(session.get("csrf_token", ""), request.form.get("csrf_token", ""))


def _legacy_request_wants_json() -> bool:
    requested_with = str(request.headers.get("X-Requested-With", "")).strip().lower()
    accept = str(request.headers.get("Accept", "")).strip().lower()
    return requested_with == "xmlhttprequest" or "application/json" in accept


def status_slug(status_text: str) -> str:
    lowered = status_text.lower()
    if lowered.startswith("up"):
        return "running"
    if "restarting" in lowered:
        return "restarting"
    if "created" in lowered:
        return "created"
    return "stopped"


def uptime_label(total_seconds: int) -> str:
    days, rem = divmod(int(total_seconds), 86400)
    hours, rem = divmod(rem, 3600)
    minutes, _ = divmod(rem, 60)
    parts = []
    if days:
        parts.append(f"{days}d")
    if hours or days:
        parts.append(f"{hours}h")
    parts.append(f"{minutes}m")
    return " ".join(parts)


def local_cpu_percent(delay: float = 0.15) -> float:
    def read_cpu():
        with open("/proc/stat", "r", encoding="utf-8") as handle:
            parts = handle.readline().split()[1:]
        values = [int(item) for item in parts]
        idle = values[3] + values[4]
        total = sum(values)
        return idle, total

    idle1, total1 = read_cpu()
    time.sleep(delay)
    idle2, total2 = read_cpu()
    total_delta = total2 - total1
    idle_delta = idle2 - idle1
    if total_delta <= 0:
        return 0.0
    return round(100.0 * (1.0 - (idle_delta / total_delta)), 1)


def local_temperature():
    probes = [
        Path("/sys/class/thermal/thermal_zone0/temp"),
        Path("/sys/class/thermal/thermal_zone1/temp"),
    ]
    for probe in probes:
        if probe.exists():
            raw = probe.read_text(encoding="utf-8").strip()
            try:
                value = float(raw)
            except ValueError:
                continue
            return round(value / 1000.0, 1) if value > 1000 else round(value, 1)
    return None


def local_memory():
    values = {}
    with open("/proc/meminfo", "r", encoding="utf-8") as handle:
        for line in handle:
            key, value = line.split(":", 1)
            values[key] = int(value.strip().split()[0])
    total = values.get("MemTotal", 0)
    available = values.get("MemAvailable", 0)
    used = max(total - available, 0)
    percent = round((used / total) * 100, 1) if total else 0.0
    return {
        "total_gb": round(total / 1024 / 1024, 2),
        "used_gb": round(used / 1024 / 1024, 2),
        "percent": percent,
    }


def parse_docker_lines(output: str):
    containers = []
    for line in output.splitlines():
        line = line.strip()
        if not line:
            continue
        try:
            payload = json.loads(line)
        except json.JSONDecodeError:
            continue
        containers.append(
            {
                "name": payload.get("Names", ""),
                "image": payload.get("Image", ""),
                "status_text": payload.get("Status", ""),
            }
        )
    return containers


def _legacy_run_local(command, timeout=20):
    return subprocess.run(command, capture_output=True, text=True, timeout=timeout, check=False)


def collect_local_snapshot():
    disk = shutil.disk_usage("/")
    docker = run_local(["sudo", "-n", "docker", "ps", "-a", "--format", "{{json .}}"], timeout=30)
    container_error = docker.stderr.strip() if docker.returncode else ""
    return {
        "reachable": True,
        "hostname": run_local(["hostname"]).stdout.strip() or "server-106",
        "cpu_percent": local_cpu_percent(),
        "load_avg": [round(value, 2) for value in os.getloadavg()],
        "memory": local_memory(),
        "disk": {
            "total_gb": round(disk.total / 1024 / 1024 / 1024, 1),
            "used_gb": round((disk.total - disk.free) / 1024 / 1024 / 1024, 1),
            "percent": round(((disk.total - disk.free) / disk.total) * 100, 1) if disk.total else 0.0,
        },
        "temperature_c": local_temperature(),
        "uptime_seconds": int(float(Path("/proc/uptime").read_text(encoding="utf-8").split()[0])),
        "containers": parse_docker_lines(docker.stdout),
        "container_error": container_error,
    }


def _legacy_remote_client():
    key_path = REMOTE118["key_path"]
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    client.connect(
        REMOTE118["host"],
        port=REMOTE118["port"],
        username=REMOTE118["user"],
        key_filename=key_path,
        look_for_keys=False,
        allow_agent=False,
        timeout=6,
        banner_timeout=6,
        auth_timeout=6,
    )
    return client


def _legacy_run_remote(command: str, timeout: int = 25):
    client = remote_client()
    try:
        stdin, stdout, stderr = client.exec_command(command, timeout=timeout)
        exit_code = stdout.channel.recv_exit_status()
        return {
            "returncode": exit_code,
            "stdout": stdout.read().decode("utf-8", errors="replace"),
            "stderr": stderr.read().decode("utf-8", errors="replace"),
        }
    finally:
        client.close()


def agent_host_online() -> bool:
    _, snapshots, _ = current_cached_state()
    return bool((snapshots.get("118") or {}).get("reachable"))


def _legacy_remote_codex_exec(prompt: str, timeout: int = AGENT_REMOTE_TIMEOUT_SECONDS):
    prompt_b64 = base64.b64encode(prompt.encode("utf-8")).decode("ascii")
    remote_script = f"""
import base64
import json
import pathlib
import subprocess
import tempfile

prompt = base64.b64decode({prompt_b64!r}).decode("utf-8")
output_file = tempfile.NamedTemporaryFile(prefix="codex-agent-", suffix=".txt", delete=False)
output_path = pathlib.Path(output_file.name)
output_file.close()
cmd = [
    "/usr/local/bin/codex",
    "exec",
    "--skip-git-repo-check",
    "--cd",
    "/home/sam",
    "--ephemeral",
    "-s",
    "read-only",
    "--color",
    "never",
    "-o",
    str(output_path),
    prompt,
]
try:
    completed = subprocess.run(cmd, capture_output=True, text=True, timeout={max(30, timeout)})
    content = ""
    if output_path.exists():
        content = output_path.read_text(encoding="utf-8", errors="replace")
    payload = {{
        "returncode": completed.returncode,
        "stdout": completed.stdout,
        "stderr": completed.stderr,
        "content": content,
    }}
except subprocess.TimeoutExpired as exc:
    payload = {{
        "returncode": 124,
        "stdout": exc.stdout or "",
        "stderr": exc.stderr or "Timed out while waiting for Codex.",
        "content": "",
    }}
finally:
    try:
        output_path.unlink(missing_ok=True)
    except Exception:
        pass
print(json.dumps(payload))
""".strip()
    result = run_remote("python3 - <<'PY'\n" + remote_script + "\nPY", timeout=max(timeout + 20, 60))
    if result["returncode"] != 0:
        raise RuntimeError(result["stderr"].strip() or "Remote Codex runner failed.")
    try:
        payload = json.loads(result["stdout"])
    except json.JSONDecodeError as exc:
        raise RuntimeError("Remote Codex response was not valid JSON.") from exc
    if payload.get("returncode") != 0:
        stderr_text = truncate_text(payload.get("stderr") or payload.get("stdout") or "Codex execution failed.", 260)
        raise RuntimeError(stderr_text)
    content = str(payload.get("content") or "").strip()
    if not content:
        raise RuntimeError("Codex returned an empty response.")
    return truncate_text(content, AGENT_OUTPUT_MAX_CHARS)


def _legacy_collect_remote_snapshot():
    command = "python3 - <<'PY'\n" + REMOTE_SNAPSHOT_SCRIPT + "\nPY"
    try:
        result = run_remote(command, timeout=30)
    except Exception as exc:
        return {
            "reachable": False,
            "error": str(exc),
            "containers": [],
        }
    if result["returncode"] != 0:
        return {
            "reachable": False,
            "error": result["stderr"].strip() or "Remote collection failed",
            "containers": [],
        }
    try:
        return json.loads(result["stdout"])
    except json.JSONDecodeError:
        return {
            "reachable": False,
            "error": "Remote data was not valid JSON",
            "containers": [],
        }


def enrich_services(server_id: str, snapshot: dict):
    container_map = {item["name"]: item for item in snapshot.get("containers", [])}
    services = []
    for service in SERVICES[server_id]:
        expected = service.get("containers", [])
        found = [container_map[name] for name in expected if name in container_map]
        running_count = sum(1 for item in found if status_slug(item["status_text"]) == "running")
        if not found:
            service_state = "stopped"
            service_text = "No active containers detected"
        elif running_count == len(expected) and len(found) == len(expected):
            service_state = "running"
            service_text = f"{running_count}/{len(expected)} containers running"
        elif running_count == 0:
            service_state = "stopped"
            service_text = f"{len(found)} containers present but not running"
        else:
            service_state = "partial"
            service_text = f"{running_count}/{len(expected)} containers running"
        services.append(
            {
                **service,
                "state": service_state,
                "state_text": service_text,
                "containers_found": found,
                "actions": [] if service["control_mode"] == "read_only" else ["start", "stop", "restart"],
            }
        )
    return services


def server_state(server_id: str, snapshot: dict, server_history: dict, updated_at: int | None):
    server = SERVERS[server_id].copy()
    server["snapshot"] = snapshot
    server["services"] = enrich_services(server_id, snapshot)
    server["metrics"] = build_metric_tiles(snapshot, server_history, updated_at)
    return server


def build_dashboard(username: str | None = None, preferred_tab: str | None = None, gym_day: str | None = None):
    updated_at, snapshots, history = current_cached_state()
    current_username = normalize_username(username if username is not None else viewer_username())
    viewer_display_name = display_name_for_username(current_username or dashboard_username())
    agent_available = bool((snapshots.get("118") or {}).get("reachable"))
    diet_state = build_diet_state(current_username or dashboard_username(), agent_available=agent_available)
    gym_state = build_gym_state(current_username or dashboard_username(), agent_available=agent_available, selected_day=gym_day)
    health_state = build_health_state(current_username or dashboard_username())
    assistant_state = build_assistant_state(current_username or dashboard_username(), agent_available=agent_available)
    return {
        "viewer_username": current_username,
        "viewer_display_name": viewer_display_name,
        "app_title": viewer_app_title(current_username or dashboard_username()),
        "sections": dashboard_sections(),
        "default_tab": resolve_dashboard_tab(preferred_tab, health_state.get("needs_profile_setup", False)),
        "diet": diet_state,
        "gym": gym_state,
        "health": health_state,
        "assistant": assistant_state,
        "updated_at": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(updated_at or time.time())),
        "refresh_seconds": REFRESH_SECONDS,
        "client_poll_seconds": CLIENT_POLL_SECONDS,
    }


def _legacy_compose_shell(path: str, action: str, sudo_prefix: str = "") -> str:
    safe_path = shlex.quote(path)
    if action == "start":
        body = "(docker-compose up -d || docker compose up -d)"
    elif action == "stop":
        body = "(docker-compose stop || docker compose stop)"
    elif action == "restart":
        body = "(docker-compose restart || docker compose restart)"
    else:
        raise ValueError("Unsupported compose action")
    if sudo_prefix:
        body = body.replace("docker-compose", f"{sudo_prefix}docker-compose").replace("docker compose", f"{sudo_prefix}docker compose")
    return f"cd {safe_path} && {body}"


def _legacy_container_shell(containers, action: str, sudo_prefix: str = "") -> str:
    safe_names = " ".join(shlex.quote(item) for item in containers)
    if action == "start":
        verb = "start"
    elif action == "stop":
        verb = "stop"
    elif action == "restart":
        verb = "restart"
    else:
        raise ValueError("Unsupported container action")
    return f"{sudo_prefix}docker {verb} {safe_names}"


def local_service_action(service: dict, action: str):
    if service["control_mode"] == "compose":
        command = compose_shell(service["path"], action, sudo_prefix="sudo -n ")
    else:
        command = container_shell(service["containers"], action, sudo_prefix="sudo -n ")
    return run_local(["bash", "-lc", command], timeout=180)


def remote_service_action(service: dict, action: str):
    if service["control_mode"] == "compose":
        command = compose_shell(service["path"], action)
    else:
        command = container_shell(service["containers"], action)
    result = run_remote(f"bash -lc {shlex.quote(command)}", timeout=180)
    completed = subprocess.CompletedProcess(
        args=command,
        returncode=result["returncode"],
        stdout=result["stdout"],
        stderr=result["stderr"],
    )
    return completed


def _legacy_remote_host_action(action: str):
    if action == "reboot":
        run_remote("bash -lc 'sudo -n systemctl reboot >/dev/null 2>&1 &'")
        return "Reboot command sent to server 118."
    raise ValueError("Unsupported host action")


def _legacy_local_host_action(action: str):
    if action != "reboot":
        raise ValueError("Unsupported host action")
    command = "nohup bash -lc 'sleep 3; sudo -n systemctl reboot' >/dev/null 2>&1 &"
    result = run_local(["bash", "-lc", command], timeout=10)
    if result.returncode != 0:
        message = (result.stderr or result.stdout or "Failed to queue reboot").strip()
        raise RuntimeError(message)
    return "Reboot command queued for server 106."


def generate_agent_insight(username: str, agent_key: str):
    current_agent = str(agent_key or "").strip().lower()
    if current_agent not in AGENT_DISPLAY_NAMES:
        raise ValueError("Unknown AI coach selected.")
    if not agent_host_online():
        raise RuntimeError("MSI is offline. The built-in planner still works, but the AI coach is unavailable right now.")

    if current_agent == "diet":
        state = build_diet_state(username, agent_available=True)
        context_key = state["agent"]["context_key"]
        prompt = build_diet_agent_prompt(state["profile"], state["plan"])
    else:
        state = build_gym_state(username, agent_available=True)
        context_key = state["agent"]["context_key"]
        prompt = build_gym_agent_prompt(state["profile"], state["plan"], state.get("recent_sessions", []))

    response_text = remote_codex_exec(prompt)
    payload = parse_agent_payload_text(current_agent, response_text)
    insight = save_coach_insight(username, current_agent, context_key, payload, source_label="Codex on MSI")
    return {
        "agent_key": current_agent,
        "context_key": context_key,
        "insight": insight,
        "message": f"{AGENT_DISPLAY_NAMES[current_agent]} updated.",
    }


def render_login_page(
    *,
    login_username: str = "",
    register_data: dict | None = None,
    active_form: str = "login",
    status_code: int = 200,
):
    defaults = {
        "first_name": "",
        "last_name": "",
        "username": "",
        "height_cm": "",
        "age_years": "",
        "current_weight_kg": "",
        "gender": "unspecified",
        "goal": "recomp",
        "preferred_meals_per_day": str(DEFAULT_DIET_MEALS_PER_DAY),
    }
    values = {**defaults, **{key: str(value or "") for key, value in (register_data or {}).items()}}
    allow_registration = host_allows_registration()
    current_host = (request.host or "").split(":", 1)[0] if request else ""
    response = render_template(
        "login.html",
        app_title=APP_TITLE,
        csrf_token=ensure_csrf_token(),
        login_username=login_username,
        active_auth_panel=active_form,
        register_data=values,
        allow_registration=allow_registration,
        public_host=current_host or "health.sam-mousavi.com",
        goal_options=[{"value": key, "label": label} for key, label in GYM_PROFILE_GOALS.items()],
        gender_options=[{"value": key, "label": label} for key, label in HEALTH_GENDER_OPTIONS.items()],
        meal_count_options=[{"value": value, "label": f"{value} meals"} for value in DIET_MEAL_COUNT_OPTIONS],
    )
    return response, status_code


def validate_registration_form(form) -> dict:
    first_name = form.get("first_name", "").strip()
    last_name = form.get("last_name", "").strip()
    username = normalize_username(form.get("username", ""))
    password = form.get("password", "")
    password_confirm = form.get("password_confirm", "")
    height_cm = safe_float(form.get("height_cm"))
    age_years = normalize_age_years(form.get("age_years"))
    current_weight_kg = safe_float(form.get("current_weight_kg"))
    gender = normalize_gender(form.get("gender"))
    goal = form.get("goal", "recomp").strip()
    preferred_meals_per_day = normalize_meals_per_day(form.get("preferred_meals_per_day"), DEFAULT_DIET_MEALS_PER_DAY)

    if not first_name:
        raise ValueError("First name is required.")
    if not last_name:
        raise ValueError("Family name is required.")
    if not REGISTRATION_USERNAME_PATTERN.fullmatch(str(form.get("username", "")).strip()):
        raise ValueError("Username must be 3-32 characters using letters, numbers, dots, underscores, or hyphens.")
    if len(password) < 8:
        raise ValueError("Password must be at least 8 characters.")
    if password != password_confirm:
        raise ValueError("Password confirmation does not match.")
    if height_cm is None or height_cm < 100 or height_cm > 250:
        raise ValueError("Enter a valid height in cm.")
    if age_years is None:
        raise ValueError("Enter a valid age in years.")
    if current_weight_kg is None or current_weight_kg < 30 or current_weight_kg > 300:
        raise ValueError("Enter a valid current weight in kg.")
    if goal not in GYM_PROFILE_GOALS:
        goal = "recomp"

    return {
        "first_name": first_name,
        "last_name": last_name,
        "username": username,
        "password": password,
        "height_cm": height_cm,
        "age_years": age_years,
        "current_weight_kg": current_weight_kg,
        "gender": gender,
        "goal": goal,
        "preferred_meals_per_day": preferred_meals_per_day,
    }


@app.route("/login", methods=["GET", "POST"])
def login():
    ensure_gym_databases()
    if session.get("authenticated") and request.method == "GET":
        return redirect(url_for("index"))
    if request.method == "POST":
        if not verify_csrf():
            flash("Security token mismatch. Reload the page and try again.", "error")
            return render_login_page(login_username=request.form.get("username", ""), active_form="login", status_code=400)
        username = request.form.get("username", "")
        password = request.form.get("password", "")
        authenticated_username = authenticate_dashboard_user(username, password)
        if authenticated_username:
            session.clear()
            session["authenticated"] = True
            session["username"] = authenticated_username
            ensure_app_user(authenticated_username)
            ensure_csrf_token()
            return redirect(url_for("index"))
        flash("Login failed. Check the Health Hub username and password.", "error")
        return render_login_page(login_username=username, active_form="login", status_code=401)
    return render_login_page()


@app.route("/register", methods=["POST"])
def register():
    ensure_gym_databases()
    if not host_allows_registration():
        flash("Registration is not available on this host.", "error")
        return render_login_page(active_form="login", status_code=403)
    if not verify_csrf():
        flash("Security token mismatch. Reload the page and try again.", "error")
        return render_login_page(register_data=request.form.to_dict(), active_form="register", status_code=400)

    register_values = request.form.to_dict()
    try:
        payload = validate_registration_form(request.form)
        authenticated_username = register_dashboard_user(
            payload["username"],
            payload["password"],
            payload["first_name"],
            payload["last_name"],
            role="member",
        )
    except ValueError as exc:
        flash(str(exc), "error")
        register_values.pop("password", None)
        register_values.pop("password_confirm", None)
        return render_login_page(register_data=register_values, active_form="register", status_code=400)

    save_gym_profile(authenticated_username, request.form)
    with gym_user_connection() as connection:
        user_row = connection.execute(
            "SELECT id FROM users WHERE username = ? LIMIT 1",
            (authenticated_username,),
        ).fetchone()
        if user_row and payload["current_weight_kg"] is not None:
            insert_body_measurement(
                connection,
                user_row["id"],
                datetime.now().date().isoformat(),
                weight_kg=payload["current_weight_kg"],
                note="Initial registration",
                source="signup",
            )
            sync_profile_current_weight_from_measurements(connection, user_row["id"])

    session.clear()
    session["authenticated"] = True
    session["username"] = authenticated_username
    ensure_csrf_token()
    flash("Account created. Welcome to Health Hub.", "success")
    return redirect(url_for("index"))


@app.route("/logout", methods=["POST"])
@login_required
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def index():
    state = build_dashboard(viewer_username(), request.args.get("tab"), request.args.get("gym_day"))
    return render_template(
        "dashboard.html",
        app_title=state["app_title"],
        state=state,
        csrf_token=ensure_csrf_token(),
    )


@app.route("/api/state")
@login_required
def api_state():
    return jsonify(build_dashboard(viewer_username(), gym_day=request.args.get("gym_day")))


@app.route("/gym/profile", methods=["POST"])
@login_required
def gym_profile():
    if not verify_csrf():
        flash("Security token mismatch. Reload the page and try again.", "error")
        return redirect(url_for("index"))
    try:
        save_gym_profile(viewer_username() or dashboard_username(), request.form)
    except ValueError as exc:
        flash(str(exc), "error")
        return redirect(url_for("index", tab=request.form.get("return_tab") or "health"))
    flash("Health profile saved.", "success")
    return redirect(url_for("index", tab=request.form.get("return_tab") or "health"))


@app.route("/gym/exercise", methods=["POST"])
@login_required
def gym_exercise():
    is_xhr = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    target_tab = request.form.get("return_tab") or "gym"
    redirect_kwargs = {"tab": target_tab}
    request_gym_day = normalize_gym_day(request.form.get("gym_day"))
    gym_day = request_gym_day
    if target_tab == "gym" and gym_day != "today":
        redirect_kwargs["gym_day"] = gym_day
    if not verify_csrf():
        if is_xhr:
            return jsonify({"ok": False, "message": "Security token mismatch. Reload the page and try again.", "refresh_tabs": []}), 400
        flash("Security token mismatch. Reload the page and try again.", "error")
        return redirect(url_for("index", **redirect_kwargs))
    try:
        result = save_current_exercise_weight(viewer_username() or dashboard_username(), request.form)
    except ValueError as exc:
        if is_xhr:
            return jsonify({"ok": False, "message": str(exc), "refresh_tabs": []}), 400
        flash(str(exc), "error")
        return redirect(url_for("index", **redirect_kwargs))
    result_day = normalize_gym_day((result or {}).get("performed_on"))
    if target_tab == "gym" and request_gym_day != "today" and result_day == datetime.now().date().isoformat():
        gym_day = "today"
        redirect_kwargs.pop("gym_day", None)
    elif target_tab == "gym" and request_gym_day != "today":
        gym_day = request_gym_day
    if is_xhr:
        return jsonify(
            {
                "ok": True,
                "message": (result or {}).get("message", f"Updated {request.form.get('exercise_name', 'exercise')}."),
                "refresh_tabs": ["gym"],
                "gym_day": gym_day,
            }
        )
    flash((result or {}).get("message", f"Updated {request.form.get('exercise_name', 'exercise')}."), "success")
    return redirect(url_for("index", **redirect_kwargs))


@app.route("/agent/run", methods=["POST"])
@login_required
def run_agent():
    if not verify_csrf():
        flash("Security token mismatch. Reload the page and try again.", "error")
        return redirect(url_for("index"))
    target_tab = request.form.get("return_tab") or "diet"
    redirect_kwargs = {"tab": target_tab}
    gym_day = normalize_gym_day(request.form.get("gym_day"))
    if target_tab == "gym" and gym_day != "today":
        redirect_kwargs["gym_day"] = gym_day
    try:
        result = generate_agent_insight(viewer_username() or dashboard_username(), request.form.get("agent_key", ""))
    except (ValueError, RuntimeError) as exc:
        flash(str(exc), "error")
        return redirect(url_for("index", **redirect_kwargs))
    except Exception:
        flash("The AI coach failed unexpectedly. The built-in planner is still available.", "error")
        return redirect(url_for("index", **redirect_kwargs))
    flash((result or {}).get("message", "AI coach updated."), "success")
    return redirect(url_for("index", **redirect_kwargs))


@app.route("/assistant/attachment/<int:attachment_id>", methods=["GET"])
@login_required
def assistant_attachment_download(attachment_id: int):
    current_username = viewer_username() or dashboard_username()
    prune_assistant_attachments(current_username)
    user_id = ensure_app_user(current_username)
    with gym_user_connection() as connection:
        row = connection.execute(
            """
            SELECT *
            FROM assistant_attachments
            WHERE id = ? AND user_id = ?
            LIMIT 1
            """,
            (attachment_id, user_id),
        ).fetchone()
    if not row:
        abort(404)
    file_path = APP_ROOT / str(row["storage_path"] or "")
    if not file_path.exists():
        with gym_user_connection() as connection:
            connection.execute("DELETE FROM assistant_attachments WHERE id = ?", (attachment_id,))
        abort(404)
    is_inline = request.args.get("inline", "").strip().lower() in {"1", "true", "yes"}
    allow_inline = is_inline and coach_attachment_is_image(row["file_ext"], row["mime_type"])
    return send_file(
        file_path,
        as_attachment=not allow_inline,
        download_name=row["original_name"],
        mimetype=row["mime_type"] or "application/octet-stream",
        max_age=0,
    )


@app.route("/assistant/chat", methods=["POST"])
@login_required
def assistant_chat():
    current_username = viewer_username() or dashboard_username()
    json_mode = request_wants_json()
    if not verify_csrf():
        if json_mode:
            return jsonify({
                "ok": False,
                "message": "Security token mismatch. Reload the page and try again.",
                "assistant": build_assistant_state(current_username, agent_available=agent_host_online()),
            }), 400
        flash("Security token mismatch. Reload the page and try again.", "error")
        return redirect(url_for("index"))
    try:
        attachments = prepare_uploaded_assistant_attachments(request.files.getlist("attachments"))
        result = generate_assistant_reply(current_username, request.form.get("message", ""), attachments=attachments)
    except (ValueError, RuntimeError) as exc:
        if json_mode:
            return jsonify({
                "ok": False,
                "message": str(exc),
                "assistant": build_assistant_state(current_username, agent_available=agent_host_online()),
            }), 400
        flash(str(exc), "error")
        return redirect(url_for("index", tab="coach"))
    except Exception:
        if json_mode:
            return jsonify({
                "ok": False,
                "message": "The coach failed unexpectedly. Try again in a moment.",
                "assistant": build_assistant_state(current_username, agent_available=agent_host_online()),
            }), 500
        flash("The coach failed unexpectedly. Try again in a moment.", "error")
        return redirect(url_for("index", tab="coach"))
    if json_mode:
        payload = {
            "ok": True,
            "message": "Coach reply ready.",
            "assistant": build_assistant_state(current_username, agent_available=agent_host_online()),
        }
        if result.get("actions"):
            payload["secondary_message"] = "The coach proposed a change. Review it and confirm before it is applied."
        return jsonify(payload)
    flash("Coach reply ready.", "success")
    if result.get("actions"):
        flash("The coach proposed a change. Review it and confirm before it is applied.", "success")
    return redirect(url_for("index", tab="coach"))


@app.route("/assistant/action", methods=["POST"])
@login_required
def assistant_action():
    current_username = viewer_username() or dashboard_username()
    json_mode = request_wants_json()
    if not verify_csrf():
        if json_mode:
            return jsonify({
                "ok": False,
                "message": "Security token mismatch. Reload the page and try again.",
                "assistant": build_assistant_state(current_username, agent_available=agent_host_online()),
            }), 400
        flash("Security token mismatch. Reload the page and try again.", "error")
        return redirect(url_for("index"))
    try:
        action_id = safe_int(request.form.get("action_id"))
        if not action_id:
            raise ValueError("Action selection is missing.")
        result = apply_assistant_action(current_username, action_id, request.form.get("decision", ""))
    except (ValueError, RuntimeError) as exc:
        if json_mode:
            return jsonify({
                "ok": False,
                "message": str(exc),
                "assistant": build_assistant_state(current_username, agent_available=agent_host_online()),
            }), 400
        flash(str(exc), "error")
        return redirect(url_for("index", tab="coach"))
    if json_mode:
        return jsonify({
            "ok": True,
            "message": result["message"],
            "assistant": build_assistant_state(current_username, agent_available=agent_host_online()),
            "refresh_tabs": result.get("refresh_tabs", []),
        })
    flash(result["message"], "success")
    return redirect(url_for("index", tab="coach"))


@app.route("/diet/meal", methods=["POST"])
@login_required
def diet_meal():
    if not verify_csrf():
        flash("Security token mismatch. Reload the page and try again.", "error")
        return redirect(url_for("index"))
    try:
        result = save_diet_meal_status(viewer_username() or dashboard_username(), request.form)
    except ValueError as exc:
        flash(str(exc), "error")
        return redirect(url_for("index", tab=request.form.get("return_tab") or "diet"))
    flash((result or {}).get("message", "Diet meal updated."), "success")
    return redirect(url_for("index", tab=request.form.get("return_tab") or "diet"))


@app.route("/diet/item", methods=["POST"])
@login_required
def diet_item():
    if not verify_csrf():
        if request_wants_json():
            return jsonify({"ok": False, "message": "Security token mismatch. Reload the page and try again.", "refresh_tabs": []}), 400
        flash("Security token mismatch. Reload the page and try again.", "error")
        return redirect(url_for("index"))
    try:
        result = save_diet_meal_item_status(viewer_username() or dashboard_username(), request.form)
    except ValueError as exc:
        if request_wants_json():
            return jsonify({"ok": False, "message": str(exc), "refresh_tabs": []}), 400
        flash(str(exc), "error")
        return redirect(url_for("index", tab=request.form.get("return_tab") or "diet"))
    if request_wants_json():
        return jsonify(
            {
                "ok": True,
                "message": (result or {}).get("message", "Diet item updated."),
                "refresh_tabs": ["diet"],
            }
        )
    flash((result or {}).get("message", "Diet item updated."), "success")
    return redirect(url_for("index", tab=request.form.get("return_tab") or "diet"))


@app.route("/health/measurement", methods=["POST"])
@login_required
def health_measurement():
    if not verify_csrf():
        flash("Security token mismatch. Reload the page and try again.", "error")
        return redirect(url_for("index"))
    try:
        save_health_measurement(viewer_username() or dashboard_username(), request.form)
    except ValueError as exc:
        flash(str(exc), "error")
        return redirect(url_for("index", tab=request.form.get("return_tab") or "health"))
    flash("Health check-in saved.", "success")
    return redirect(url_for("index", tab=request.form.get("return_tab") or "health"))


@app.route("/action", methods=["POST"])
@login_required
def action():
    abort(404)


@app.template_filter("percent")
def percent_filter(value):
    if value is None:
        return "n/a"
    return f"{value:.1f}%"


@app.template_filter("celsius")
def celsius_filter(value):
    if value is None:
        return "n/a"
    return f"{value:.1f} C"


@app.template_filter("uptime")
def uptime_filter(value):
    if value is None:
        return "n/a"
    return uptime_label(value)


if __name__ == "__main__":
    port = int(os.getenv("DASHBOARD_PORT", "8091"))
    start_sampler()
    serve(app, host="127.0.0.1", port=port)
